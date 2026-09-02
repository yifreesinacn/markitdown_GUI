import tkinter as tk
from tkinter import ttk, filedialog, messagebox, scrolledtext
import os
import threading
import subprocess
import sys
import webbrowser
import builtins
import logging
from logging.handlers import RotatingFileHandler
from pathlib import Path
import tempfile
import shutil
import glob
# 下一行为新增
# 暂未使用 import time
import json
import re

_MARKITDOWN_SENTINEL = object()  # 2026-03-29 13:41:12 修改：延迟导入markitdown，避免启动硬依赖崩溃
_MARKITDOWN = _MARKITDOWN_SENTINEL

def _get_markitdown():
    global _MARKITDOWN
    if _MARKITDOWN is not _MARKITDOWN_SENTINEL:
        return _MARKITDOWN
    try:
        import markitdown as module
    except Exception:
        module = None
    _MARKITDOWN = module
    return module

def _get_markitdown_version():
    module = _get_markitdown()
    if module is None:
        return "unknown"
    return getattr(module, "__version__", "unknown")

def _get_markitdown_version_display():  # 2026-04-01 18:10:00 修改：关于页版本显示仅保留数字与点（去除右侧字母后缀）
    version_text = str(_get_markitdown_version() or "")
    match = re.match(r'^\s*(\d+(?:\.\d+)*)', version_text)
    return match.group(1) if match else "unknown"

APP_VERSION = "0.1.7"          # 2026-09-02 修改：新增项目自身版本号常量（唯一权威来源，与MarkItDown引擎版本分离）
APP_VERSION_DATE = "20260902"  # 2026-09-02 修改：新增项目版本日期常量（统一YYYYMMDD格式，作为版本辅助信息）

_MARKITDOWN_CLI_TIMEOUT = 600  # 2026-03-29 13:41:12 修改：外部转换命令增加超时，防止卡死

DOCUMENT_OUTPUT_MODE_LABELS = {
    'smart': '智能转换（图文并存）',
    'layout': '保持版式（页面转图片）',
    'image': '保留图像（仅提取图片）',
    'text': '保留文字（仅提取文字）',
    'text_ocr': '保留文字（含识别文字）'
}
DOCUMENT_OUTPUT_MODE_KEYS = {label: key for key, label in DOCUMENT_OUTPUT_MODE_LABELS.items()}

# 2026-09-01 新增：PDF 引擎策略（仅作用于智能转换模式）
PDF_ENGINE_STRATEGY_LABELS = ['自动', '图文混排优先', 'pdf-inspector优先']
PDF_ENGINE_STRATEGY_DEFAULT = '自动'

def _create_module_logger():
    logger = logging.getLogger("anytomd.app")
    logger.setLevel(logging.INFO)
    for handler in list(logger.handlers):
        try:
            if getattr(getattr(handler, 'stream', None), 'closed', False):
                logger.removeHandler(handler)
                handler.close()
        except Exception:
            try:
                logger.removeHandler(handler)
            except Exception:
                pass
    if logger.handlers:
        return logger
    log_dir = os.path.join(os.path.expanduser('~'), '.anytomd_logs')
    os.makedirs(log_dir, exist_ok=True)
    log_file = os.path.join(log_dir, 'anytomd_app.log')
    file_handler = RotatingFileHandler(log_file, maxBytes=2 * 1024 * 1024, backupCount=3, encoding='utf-8')
    file_handler.setFormatter(logging.Formatter('%(asctime)s | %(levelname)s | %(message)s'))
    logger.addHandler(file_handler)
    logger.propagate = False
    return logger

_APP_LOGGER = _create_module_logger()
_BUILTIN_PRINT = builtins.print

def _rebuild_module_logger():
    global _APP_LOGGER
    logger = logging.getLogger("anytomd.app")
    for handler in list(logger.handlers):
        try:
            handler.close()
        except Exception:
            pass
        try:
            logger.removeHandler(handler)
        except Exception:
            pass
    _APP_LOGGER = _create_module_logger()
    return _APP_LOGGER

def print(*args, **kwargs):
    message = " ".join(str(item) for item in args)
    try:
        _APP_LOGGER.info(json.dumps({"event": "console", "message": message}, ensure_ascii=False))
    except Exception:
        try:
            _rebuild_module_logger().info(json.dumps({"event": "console", "message": message}, ensure_ascii=False))
        except Exception:
            pass
    return _BUILTIN_PRINT(*args, **kwargs)

ctk = None
try:
    import customtkinter as ctk
except Exception:
    try:
        import site
        user_site = site.getusersitepackages()
        if user_site and os.path.isdir(user_site) and user_site not in sys.path:
            sys.path.append(user_site)
        import customtkinter as ctk
    except Exception:
        ctk = None

class MarkItDownApp:
    def __init__(self, root):
        self.root = root
        self.use_ctk = ctk is not None and isinstance(root, ctk.CTk)
        self.logger = _APP_LOGGER
        self.root.title(f'AnyToMarkDown 文档转换工具 v{APP_VERSION} 版 【yifree {APP_VERSION_DATE}】')  # 2026-09-02 修改：标题栏版本号改为项目自身版本常量（与MarkItDown引擎版本分离）
        self.root.iconbitmap(default=self.resource_path('icon.ico') if hasattr(sys, '_MEIPASS') else None)
        
        # 设置主题和样式
        self.setup_styles()
        
        # 设置窗口大小和位置
        window_width = 1000
        screen_width = root.winfo_screenwidth()
        screen_height = root.winfo_screenheight()
        min_window_height = 660  # 2026-04-01 18:45:00 修改：主窗口最小高度略缩小（更协调）
        max_window_height = max(min_window_height, int(screen_height * 0.9))
        window_height = min(max_window_height, 720)  # 2026-04-01 18:45:00 修改：默认窗口高度略缩小（更协调）
        x = (screen_width - window_width) // 2
        y = (screen_height - window_height) // 4
        self.root.geometry(f'{window_width}x{window_height}+{x}+{y}')
        self.root.minsize(800, min_window_height)
        
        # 创建主框架
        if self.use_ctk:
            self.main_frame = ctk.CTkFrame(self.root)
            self.main_frame.pack(fill=tk.BOTH, expand=True, padx=10, pady=10)
        else:
            self.main_frame = ttk.Frame(self.root, padding="10") 
            self.main_frame.pack(fill=tk.BOTH, expand=True)
        
        # 创建标题和说明
        self.create_header()
        
        # 创建选项卡
        self.create_notebook()
        
        # 创建状态栏
        self.create_statusbar()
        self.root.update_idletasks()
        content_height = self.root.winfo_reqheight() + 16
        calibrated_height = max(min_window_height, min(content_height, max_window_height))
        if calibrated_height != window_height:
            y = max(0, (screen_height - calibrated_height) // 4)
            self.root.geometry(f'{window_width}x{calibrated_height}+{x}+{y}')
        
        # 初始化变量
        self.current_file = None
        self.output_folder = None
        self.conversion_in_progress = False
        self.batch_files = []
        # 临时文件和目录管理
        self.temp_files = []
        self.temp_dirs = []
        
        # 支持的文件类型
        self.file_types = [  # 2026-09-02 修改：为旧版 .ppt/.xls 补充类型（与 AnyDoc 兜底引擎能力一致，此前漏列导致文件选择框与“添加文件夹”均无法选到 .ppt/.xls）
            ('所有支持的文件', '*.pdf;*.pptx;*.ppt;*.docx;*.doc;*.xlsx;*.xls;*.png;*.jpg;*.jpeg;*.mp3;*.wav;*.html;*.htm;*.mhtml;*.mht;*.csv;*.json;*.xml;*.zip'),
            ('PDF 文档', '*.pdf'),
            ('PowerPoint 文档', '*.pptx;*.ppt'),
            ('Word 文档', '*.docx;*.doc'),  # 添加.doc支持
            ('Excel 文档', '*.xlsx;*.xls'),
            ('图片文件', '*.png;*.jpg;*.jpeg'),
            ('音频文件', '*.mp3;*.wav'),
            ('HTML 文件', '*.html;*.htm;*.mhtml;*.mht'),
            ('文本文件', '*.csv;*.json;*.xml'),
            ('压缩文件', '*.zip')
        ]
        
        # 初始化markitdown功能信息并进行初始检测
        self.markitdown_info = self.get_markitdown_info()
        self.run_initial_feature_check()
        
        # 统一更新状态栏信息
        version = self.markitdown_info.get('version', '未知')
        features = self.markitdown_info.get('features_text', '基本转换')
        self.markitdown_status_var.set(f'已集成 Microsoft MarkItDown v{version} ({features})')
        if self.use_ctk and hasattr(self, 'markitdown_status_label'):
            self.markitdown_status_label.configure(text=self.markitdown_status_var.get())
        
        print(f"MarkItDown 版本: {version}")
        if self.markitdown_info.get('all_params'):
            print(f"支持的参数: {', '.join(self.markitdown_info['all_params'])}")

    def resource_path(self, relative_path):
        """获取资源的绝对路径，适用于PyInstaller打包后的情况"""
        try:
            base_path = sys._MEIPASS
        except Exception:
            # 改进：使用脚本所在目录而不是当前工作目录
            base_path = os.path.dirname(os.path.abspath(__file__))
        return os.path.join(base_path, relative_path)

    def find_poppler_bin(self):
        bases = []
        if hasattr(sys, "_MEIPASS"):
            bases.append(sys._MEIPASS)
        # 改进：优先搜索脚本所在目录，确保在不同目录下运行都能找到
        script_dir = os.path.dirname(os.path.abspath(__file__))
        bases.append(script_dir)
        # 兼容性：保留当前目录搜索
        if os.path.abspath(".") != script_dir:
            bases.append(os.path.abspath("."))

        for base in bases:
            # 搜索 poppler 开头的目录
            for poppler_dir in sorted(glob.glob(os.path.join(base, "poppler*"))):
                # 检查 Windows 下典型的 Library/bin 路径
                poppler_bin = os.path.join(poppler_dir, "Library", "bin")
                if os.path.exists(os.path.join(poppler_bin, "pdftoppm.exe")):
                    return poppler_bin
                # 兼容：直接在 bin 下的情况
                poppler_bin_alt = os.path.join(poppler_dir, "bin")
                if os.path.exists(os.path.join(poppler_bin_alt, "pdftoppm.exe")):
                    return poppler_bin_alt
        return None

    def _is_path_within_dir(self, candidate_path, base_dir):
        try:
            candidate_norm = os.path.normcase(os.path.abspath(os.path.realpath(candidate_path)))
            base_norm = os.path.normcase(os.path.abspath(os.path.realpath(base_dir)))
            return os.path.commonpath([candidate_norm, base_norm]) == base_norm
        except Exception:
            return False

    def setup_styles(self):
        """设置应用程序的样式和主题"""
        if self.use_ctk:
            ctk.set_appearance_mode("System")
            ctk.set_default_color_theme("blue")
            return
        style = ttk.Style()
        
        # 尝试使用更现代的主题
        try:
            style.theme_use('vista')
        except tk.TclError:
            try:
                style.theme_use('xpnative')
            except tk.TclError:
                try:
                    style.theme_use('clam')
                except tk.TclError:
                    pass
        
        style.configure('TButton', font=('微软雅黑', 10), padding=(10, 7))  # 2026-04-01 17:20:00 修改：统一按钮字体与留白（更清爽）
        style.configure('Primary.TButton', font=('微软雅黑', 10, 'bold'), padding=(10, 7))  # 2026-04-01 17:20:00 修改：主按钮仅加粗，不强制配色
        
        # 自定义标签样式
        style.configure('TLabel', font=('微软雅黑', 10))  # 2026-04-01 17:20:00 修改：统一正文字号
        style.configure('Header.TLabel', font=('微软雅黑', 18, 'bold'), foreground='#000000')  # 2026-04-01 17:20:00 修改：标题更清晰
        style.configure('Subheader.TLabel', font=('微软雅黑', 11), foreground='#000000')  # 2026-04-01 17:20:00 修改：副标题更轻量
        style.configure('TLabelframe.Label', font=('微软雅黑', 10, 'bold'))  # 2026-04-01 17:20:00 修改：分组标题统一层级
        style.configure('TEntry', font=('微软雅黑', 10))  # 2026-04-01 17:20:00 修改：输入框字体统一
        style.configure('TCombobox', font=('微软雅黑', 10))  # 2026-04-01 17:20:00 修改：下拉框字体统一
        style.configure('Treeview', font=('微软雅黑', 10))  # 2026-04-01 17:20:00 修改：列表字体统一（如使用）
        
        # 自定义框架样式
        style.configure('Card.TFrame', relief='raised', borderwidth=1)  
        
        # 自定义进度条样式
        style.configure('TProgressbar', thickness=8)

    def _apply_notebook_menu_style(self, tab_names):
        if self.use_ctk:
            segmented = getattr(self.notebook, '_segmented_button', None)
            if segmented is None:
                return
            base_width = max(90, int((max(len(name) for name in tab_names) + 2) * 16))
            target_width = int(base_width * 1.26)
            total_width = 0
            buttons_dict = getattr(segmented, '_buttons_dict', {})
            for name in tab_names:
                button = buttons_dict.get(name)
                if button is not None:
                    height = button.cget('height') if str(button.cget('height')).isdigit() else 28
                    button_width = target_width
                    total_width += button_width
                    button.configure(width=button_width, height=max(26, int(height * 0.92)), font=('微软雅黑', 10, 'bold'), anchor='center')
            if total_width > 0:
                segmented.configure(dynamic_resizing=False, width=total_width)
            return
        style = ttk.Style()
        tab_width = max(9, int((max(len(name) for name in tab_names) + 2) * 1.2))
        style.configure('TNotebook', borderwidth=0)
        style.configure('TNotebook.Tab', font=('微软雅黑', 10, 'bold'), padding=(12, 10), anchor='center', justify='center', width=tab_width)  # 2026-04-01 17:20:00 修改：选项卡间距更接近 Win11 清爽风格

    def _ctk_group(self, parent, title):
        outer = ctk.CTkFrame(parent)
        title_label = ctk.CTkLabel(outer, text=title, font=('微软雅黑', 12, 'bold'))
        title_label.pack(anchor=tk.W, padx=12, pady=(10, 0))
        inner = ctk.CTkFrame(outer, fg_color="transparent")
        inner.pack(fill=tk.BOTH, expand=True, padx=12, pady=12)
        return outer, inner

    def create_header(self):
        """创建应用程序标题和说明"""
        if self.use_ctk:
            header_frame = ctk.CTkFrame(self.main_frame)
            header_frame.pack(fill=tk.X, pady=(0, 6))  # 2026-04-01 19:40:00 修改：小幅收紧标题区留白

            title_label = ctk.CTkLabel(
                header_frame,
                text='AnyToMarkDown 文档转换工具',
                font=('微软雅黑', 20, 'bold')
            )
            title_label.pack(anchor=tk.N, pady=(10, 0))

            description = "将各种文档格式转换为 Markdown 格式，支持 PDF、PowerPoint、Word、Excel 等多种格式。"
            desc_label = ctk.CTkLabel(header_frame, text=description, font=('微软雅黑', 12))
            desc_label.pack(anchor=tk.N, pady=(6, 8))  # 2026-04-01 19:40:00 修改：小幅收紧标题区留白
            return
        header_frame = ttk.Frame(self.main_frame)
        header_frame.pack(fill=tk.X, pady=(0, 6))  # 2026-04-01 19:40:00 修改：小幅收紧标题区留白
        
        # 标题 修改每项为一行，删除具体版本号
        title_label = ttk.Label(
            header_frame, 
            text='AnyToMarkDown 文档转换工具',  
            style='Header.TLabel'
            # 如需要，上可改为：f'AnyToMarkDown 文档转换工具 v{markitdown.__version__}'
        )
        title_label.pack(anchor=tk.N, pady=(0, 0))
        
        # 说明
        description = "将各种文档格式转换为 Markdown 格式，支持 PDF、PowerPoint、Word、Excel 等多种格式。"
        desc_label = ttk.Label(header_frame, text=description, style='Subheader.TLabel')
        desc_label.pack(anchor=tk.N, pady=(6, 8))  # 2026-04-01 19:40:00 修改：小幅收紧标题区留白

    def create_notebook(self):
        """创建选项卡界面"""
        tab_names = ['单文件转换', '批量转换', '设置', '关于']
        if self.use_ctk:
            self.notebook = ctk.CTkTabview(self.main_frame)
            self.notebook.pack(fill=tk.BOTH, expand=True, pady=(5, 0))

            self.single_frame = self.notebook.add('单文件转换')
            self.create_single_file_tab()

            self.batch_frame = self.notebook.add('批量转换')
            self.create_batch_file_tab()

            self.settings_frame = self.notebook.add('设置')
            self.create_settings_tab()

            self.about_frame = self.notebook.add('关于')
            self.create_about_tab()
            self._apply_notebook_menu_style(tab_names)
            return
        self.notebook = ttk.Notebook(self.main_frame)
        self.notebook.pack(fill=tk.BOTH, anchor=tk.N, expand=True, pady=(6, 0))  # 2026-04-01 19:25:00 修改：减少选项卡下方留白，使页面更紧凑
        self._apply_notebook_menu_style(tab_names)
        
        # 单文件转换选项卡
        self.single_frame = ttk.Frame(self.notebook, padding=8)  # 2026-04-01 19:25:00 修改：减少内容区内边距，减少底部留白
        self.notebook.add(self.single_frame, text='单文件转换')
        self.create_single_file_tab()
        
        # 批量转换选项卡
        self.batch_frame = ttk.Frame(self.notebook, padding=8)  # 2026-04-01 19:25:00 修改：减少内容区内边距，减少底部留白
        self.notebook.add(self.batch_frame, text='批量转换')
        self.create_batch_file_tab()
        
        # 设置选项卡
        self.settings_frame = ttk.Frame(self.notebook, padding=8)  # 2026-04-01 19:25:00 修改：减少内容区内边距，减少底部留白
        self.notebook.add(self.settings_frame, text='设置')
        self.create_settings_tab()
        
        # 关于选项卡
        self.about_frame = ttk.Frame(self.notebook, padding=8)  # 2026-04-01 19:25:00 修改：减少内容区内边距，减少底部留白
        self.notebook.add(self.about_frame, text='关于')
        self.create_about_tab()
        for i in range(self.notebook.index('end')):
            self.notebook.tab(i, compound='center', padding=(12, 10))  # 2026-04-01 17:20:00 修改：统一 tab padding，避免与 style 配置不一致

    def create_single_file_tab(self):
        """创建单文件转换选项卡内容"""
        if self.use_ctk:
            left_panel = ctk.CTkFrame(self.single_frame)
            left_panel.pack(side=tk.LEFT, fill=tk.BOTH, expand=True, padx=(0, 10), pady=(0, 0))

            file_group, file_frame = self._ctk_group(left_panel, '文件选择')
            file_group.pack(fill=tk.X, pady=(10, 20))

            input_frame = ctk.CTkFrame(file_frame, fg_color="transparent")
            input_frame.pack(fill=tk.X, pady=5)

            ctk.CTkLabel(input_frame, text='输入文件:').pack(side=tk.LEFT)
            self.input_path_var = tk.StringVar()
            input_entry = ctk.CTkEntry(input_frame, textvariable=self.input_path_var)
            input_entry.pack(side=tk.LEFT, fill=tk.X, expand=True, padx=8)
            ctk.CTkButton(input_frame, text='浏览...', width=90, command=self.select_input_file).pack(side=tk.LEFT)

            output_frame = ctk.CTkFrame(file_frame, fg_color="transparent")
            output_frame.pack(fill=tk.X, pady=5)

            ctk.CTkLabel(output_frame, text='输出文件夹:').pack(side=tk.LEFT)
            self.output_path_var = tk.StringVar()
            output_entry = ctk.CTkEntry(output_frame, textvariable=self.output_path_var)
            output_entry.pack(side=tk.LEFT, fill=tk.X, expand=True, padx=8)
            ctk.CTkButton(output_frame, text='浏览...', width=90, command=self.select_output_folder).pack(side=tk.LEFT)

            options_group, options_frame = self._ctk_group(left_panel, '转换选项')
            options_group.pack(fill=tk.X, pady=(0, 10))

            self.preserve_format_var = tk.BooleanVar(value=True)
            ctk.CTkCheckBox(options_frame, text='尽可能保留原始格式', variable=self.preserve_format_var).pack(anchor=tk.W, pady=4)

            self.extract_images_var = tk.BooleanVar(value=True)
            ctk.CTkCheckBox(options_frame, text='提取并保存图片至images文件夹', variable=self.extract_images_var).pack(anchor=tk.W, pady=4)

            action_frame = ctk.CTkFrame(left_panel)
            action_frame.pack(fill=tk.X, pady=10)

            self.progress = ctk.CTkProgressBar(action_frame)
            self.progress.set(0)
            self.progress.pack(fill=tk.X, pady=(10, 10), padx=12)

            buttons_frame = ctk.CTkFrame(action_frame, fg_color="transparent")
            buttons_frame.pack(fill=tk.X, padx=12, pady=(0, 12))

            self.convert_btn = ctk.CTkButton(
                buttons_frame,
                text='开始转换',
                command=self.start_conversion
            )
            self.convert_btn.pack(side=tk.LEFT, padx=(0, 8))

            self.preview_btn = ctk.CTkButton(
                buttons_frame,
                text='预览结果',
                command=self.preview_markdown
            )
            self.preview_btn.pack(side=tk.LEFT)

            right_panel = ctk.CTkFrame(self.single_frame)
            right_panel.pack(side=tk.RIGHT, anchor=tk.N, fill=tk.BOTH, expand=True, pady=(10, 10))

            preview_group, preview_frame = self._ctk_group(right_panel, 'Markdown 预览')
            preview_group.pack(fill=tk.BOTH, expand=True)

            self.preview_text = ctk.CTkTextbox(
                preview_frame,
                font=('Consolas', 10)
            )
            self.preview_text.pack(fill=tk.BOTH, expand=True)

            preview_controls = ctk.CTkFrame(preview_frame, fg_color="transparent")
            preview_controls.pack(fill=tk.X, pady=(10, 0))

            self.copy_btn = ctk.CTkButton(
                preview_controls,
                text='复制到剪贴板',
                command=self.copy_to_clipboard
            )
            self.copy_btn.pack(side=tk.LEFT, padx=(0, 8))

            self.save_btn = ctk.CTkButton(
                preview_controls,
                text='保存预览',
                command=self.save_preview
            )
            self.save_btn.pack(side=tk.LEFT)
            return
        # 左侧面板 - 文件选择和转换选项
        self.single_frame.grid_columnconfigure(0, weight=1, uniform='single_layout')
        self.single_frame.grid_columnconfigure(1, weight=1, uniform='single_layout')
        self.single_frame.grid_rowconfigure(0, weight=1)
        left_panel = ttk.Frame(self.single_frame)
        left_panel.grid(row=0, column=0, sticky='nsew', padx=(0, 8), pady=(6, 6))  # 2026-04-01 19:25:00 修改：减少上下留白，使界面更紧凑
        
        # 文件选择区域
        file_frame = ttk.LabelFrame(left_panel, text='文件选择', padding=10)
        file_frame.pack(fill=tk.X, pady=(0, 12))  # 2026-04-01 18:45:00 修改：减少分组下方留白，使进度条上方空白更小
        
        # 输入文件
        input_frame = ttk.Frame(file_frame)
        input_frame.pack(fill=tk.X, pady=(0, 8))  # 2026-04-01 17:50:00 修改：行间距更清爽

        ttk.Label(input_frame, text='输入文件:').pack(anchor=tk.W, pady=(0, 4))

        input_entry_row = ttk.Frame(input_frame)
        input_entry_row.pack(fill=tk.X)
        self.input_path_var = tk.StringVar()
        input_entry = ttk.Entry(input_entry_row, textvariable=self.input_path_var, width=40)
        input_entry.pack(fill=tk.X, expand=True)

        input_button_row = ttk.Frame(input_frame)
        input_button_row.pack(fill=tk.X, pady=(6, 0))  # 2026-04-01 17:50:00 修改：按钮区留白更协调
        browse_btn = ttk.Button(input_button_row, text='浏览...', command=self.select_input_file)
        browse_btn.pack(side=tk.RIGHT)

        # 输出文件夹
        output_frame = ttk.Frame(file_frame)
        output_frame.pack(fill=tk.X, pady=(0, 8))  # 2026-04-01 17:50:00 修改：行间距更清爽

        ttk.Label(output_frame, text='输出文件夹:').pack(anchor=tk.W, pady=(0, 4))

        output_entry_row = ttk.Frame(output_frame)
        output_entry_row.pack(fill=tk.X)
        self.output_path_var = tk.StringVar()
        output_entry = ttk.Entry(output_entry_row, textvariable=self.output_path_var, width=40)
        output_entry.pack(fill=tk.X, expand=True)

        output_button_row = ttk.Frame(output_frame)
        output_button_row.pack(fill=tk.X, pady=(6, 0))  # 2026-04-01 17:50:00 修改：按钮区留白更协调
        browse_output_btn = ttk.Button(output_button_row, text='浏览...', command=self.select_output_folder)
        browse_output_btn.pack(side=tk.RIGHT)
        
        # 转换选项
        options_frame = ttk.LabelFrame(left_panel, text='转换选项', padding=10)
        options_frame.pack(fill=tk.X, pady=(0, 4))  # 2026-04-01 19:10:00 修改：进一步减小分组下方留白
        
        # 保留原始格式选项
        self.preserve_format_var = tk.BooleanVar(value=True)
        preserve_format_cb = ttk.Checkbutton(
            options_frame, 
            text='尽可能保留原始格式', 
            variable=self.preserve_format_var
        )
        preserve_format_cb.pack(anchor=tk.W, pady=4)  # 2026-04-01 17:50:00 修改：选项留白更清爽
        
        # 提取图片选项
        self.extract_images_var = tk.BooleanVar(value=True)
        extract_images_cb = ttk.Checkbutton(
            options_frame, 
            text='提取并保存图片至images文件夹', 
            variable=self.extract_images_var
        )
        extract_images_cb.pack(anchor=tk.W, pady=4)  # 2026-04-01 17:50:00 修改：选项留白更清爽
        
        # 转换按钮和进度条
        action_frame = ttk.Frame(left_panel)
        action_frame.pack(side=tk.BOTTOM, fill=tk.X, pady=(0, 10))  # 2026-04-01 20:05:00 修改：左侧操作区贴底，底部留白与右侧LabelFrame内边距对齐
        
        self.progress = ttk.Progressbar(action_frame, mode='determinate')
        self.progress.pack(fill=tk.X, pady=(15, 10))  # 2026-04-01 19:25:00 修改：进度条上下留白统一为10
        
        buttons_frame = ttk.Frame(action_frame)
        buttons_frame.pack(fill=tk.X, pady=(10, 0))  # 2026-04-01 19:55:00 修改：增大按钮上方间距，使左右按钮更易对齐
        
        self.convert_btn = ttk.Button(
            buttons_frame, 
            text='开始转换', 
            command=self.start_conversion,
            style='Primary.TButton'
        )
        self.convert_btn.pack(side=tk.LEFT, padx=(0, 8))  # 2026-04-01 17:50:00 修改：按钮间距更统一
        
        self.preview_btn = ttk.Button(
            buttons_frame, 
            text='预览结果', 
            command=self.preview_markdown
        )
        self.preview_btn.pack(side=tk.LEFT)
        
        # 右侧面板 - 预览区域
        right_panel = ttk.Frame(self.single_frame)
        right_panel.grid(row=0, column=1, sticky='nsew', padx=(8, 0), pady=(6, 6))  # 2026-04-01 19:25:00 修改：减少上下留白，使界面更紧凑
        right_panel.grid_columnconfigure(0, weight=1)
        right_panel.grid_rowconfigure(0, weight=1)
        
        preview_frame = ttk.LabelFrame(right_panel, text='Markdown 预览', padding=10)
        preview_frame.pack(fill=tk.BOTH, side=tk.TOP, anchor=tk.N, expand=True)  # 2026-04-01 18:35:00 修改：恢复右侧预览区自适应高度，并通过底部按钮对齐改善整体协调
        
        self.preview_text = scrolledtext.ScrolledText(
            preview_frame, 
            wrap=tk.WORD, 
            height=17,  # 2026-04-01 18:10:00 修改：预览框高度减少一行，使界面更协调
            font=('Consolas', 10)
        )
        self.preview_text.pack(fill=tk.BOTH, expand=True)
        
        preview_controls = ttk.Frame(preview_frame)
        preview_controls.pack(fill=tk.X, side=tk.BOTTOM, pady=(10, 0))  # 2026-04-01 20:05:00 修改：右侧按钮贴底，确保与左侧贴底按钮在同一基线
        
        self.copy_btn = ttk.Button(
            preview_controls, 
            text='复制到剪贴板', 
            command=self.copy_to_clipboard
        )
        self.copy_btn.pack(side=tk.LEFT, padx=(0, 8))  # 2026-04-01 17:50:00 修改：按钮间距更统一
        
        self.save_btn = ttk.Button(
            preview_controls, 
            text='保存预览', 
            command=self.save_preview
        )
        self.save_btn.pack(side=tk.LEFT)

    def create_batch_file_tab(self):
        """创建批量转换选项卡内容"""
        if self.use_ctk:
            left_panel = ctk.CTkFrame(self.batch_frame)
            left_panel.pack(side=tk.LEFT, fill=tk.BOTH, expand=True, padx=(0, 10))

            files_group, files_frame = self._ctk_group(left_panel, '计划转换文件列表')
            files_group.pack(fill=tk.BOTH, expand=True, pady=(0, 10))

            self.files_listbox = tk.Listbox(
                files_frame,
                selectmode=tk.EXTENDED,
                font=('微软雅黑', 10),
                height=10
            )
            self.files_listbox.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)

            scrollbar = ttk.Scrollbar(files_frame, orient=tk.VERTICAL, command=self.files_listbox.yview)
            scrollbar.pack(side=tk.RIGHT, fill=tk.Y)
            self.files_listbox.config(yscrollcommand=scrollbar.set)

            file_buttons_frame = ctk.CTkFrame(left_panel, fg_color="transparent")
            file_buttons_frame.pack(fill=tk.X, pady=(0, 5))

            ctk.CTkButton(file_buttons_frame, text='添加文件', command=self.add_batch_files).pack(side=tk.LEFT, padx=(0, 8))
            ctk.CTkButton(file_buttons_frame, text='添加文件夹', command=self.add_folder_files).pack(side=tk.LEFT, padx=(0, 8))
            ctk.CTkButton(file_buttons_frame, text='移除选中', command=self.remove_selected_files).pack(side=tk.LEFT, padx=(0, 8))
            ctk.CTkButton(file_buttons_frame, text='清空列表', command=self.clear_file_list).pack(side=tk.LEFT)

            right_panel = ctk.CTkFrame(self.batch_frame)
            right_panel.pack(side=tk.RIGHT, fill=tk.BOTH, expand=True, padx=(0, 0))

            list_group, list_frame = self._ctk_group(right_panel, '转换不成功文件列表')
            list_group.pack(fill=tk.BOTH, anchor=tk.N, expand=True)

            pending_frame = ctk.CTkFrame(list_frame, fg_color="transparent")
            pending_frame.pack(fill=tk.BOTH, expand=True, pady=(0, 5))

            self.pending_files_listbox = tk.Listbox(
                pending_frame,
                height=5,
                font=('微软雅黑', 9),
                selectmode=tk.SINGLE
            )
            scrollbar = ttk.Scrollbar(pending_frame, orient=tk.VERTICAL, command=self.pending_files_listbox.yview)
            self.pending_files_listbox.config(yscrollcommand=scrollbar.set)

            self.pending_files_listbox.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)
            scrollbar.pack(side=tk.RIGHT, fill=tk.Y)

            output_group, output_settings_frame = self._ctk_group(right_panel, '输出设置')
            output_group.pack(fill=tk.X, pady=5)

            output_path_frame = ctk.CTkFrame(output_settings_frame, fg_color="transparent")
            output_path_frame.pack(fill=tk.X, pady=0)

            ctk.CTkLabel(output_path_frame, text='输出文件夹:').pack(anchor=tk.W, pady=(0, 6))

            path_input_frame = ctk.CTkFrame(output_path_frame, fg_color="transparent")
            path_input_frame.pack(fill=tk.X)

            self.batch_output_var = tk.StringVar()
            output_entry = ctk.CTkEntry(path_input_frame, textvariable=self.batch_output_var)
            output_entry.pack(side=tk.LEFT, fill=tk.X, expand=True, padx=(0, 8))
            ctk.CTkButton(path_input_frame, text='浏览...', width=90, command=self.select_batch_output_folder).pack(side=tk.RIGHT)

            options_frame = ctk.CTkFrame(output_settings_frame, fg_color="transparent")
            options_frame.pack(fill=tk.X, pady=(10, 0), anchor=tk.W)

            first_row = ctk.CTkFrame(options_frame, fg_color="transparent")
            first_row.pack(fill=tk.X, pady=4, anchor=tk.W)

            self.batch_preserve_format_var = tk.BooleanVar(value=True)
            ctk.CTkCheckBox(first_row, text='尽可能保留原始格式', variable=self.batch_preserve_format_var).pack(side=tk.LEFT, padx=10)

            self.batch_extract_images_var = tk.BooleanVar(value=True)
            ctk.CTkCheckBox(first_row, text='提取并保存图片至images文件夹', variable=self.batch_extract_images_var).pack(side=tk.LEFT, padx=10)

            second_row = ctk.CTkFrame(options_frame, fg_color="transparent")
            second_row.pack(fill=tk.X, pady=4, anchor=tk.W)

            self.batch_auto_open_folder_var = tk.BooleanVar(value=True)
            ctk.CTkCheckBox(second_row, text="转换完成后打开输出文件夹", variable=self.batch_auto_open_folder_var).pack(side=tk.LEFT, padx=10)

            third_row = ctk.CTkFrame(options_frame, fg_color="transparent")
            third_row.pack(fill=tk.X, pady=4, anchor=tk.W)

            self.batch_overwrite_policy_var = tk.StringVar(value='跳过')  # 2026-03-29 13:41:12 修改：批量同名输出策略（默认跳过防误覆盖）
            ctk.CTkLabel(third_row, text="同名输出:").pack(side=tk.LEFT, padx=(10, 6))
            ctk.CTkOptionMenu(
                third_row,
                variable=self.batch_overwrite_policy_var,
                values=[
                    '跳过',
                    '覆盖',
                    '自动重命名'
                ],
                width=240
            ).pack(side=tk.LEFT)

            batch_action_frame = ctk.CTkFrame(right_panel)
            batch_action_frame.pack(fill=tk.X, pady=(5, 5))

            status_progress_frame = ctk.CTkFrame(batch_action_frame, fg_color="transparent")
            status_progress_frame.pack(fill=tk.X, pady=(10, 10), padx=12)

            self.batch_status_var = tk.StringVar(value='准备就绪')
            self.batch_status_label = ctk.CTkLabel(status_progress_frame, text=self.batch_status_var.get())
            self.batch_status_label.pack(side=tk.LEFT, padx=(0, 10))

            self.batch_progress = ctk.CTkProgressBar(status_progress_frame)
            self.batch_progress.set(0)
            self.batch_progress.pack(side=tk.LEFT, fill=tk.X, expand=True)

            self.batch_convert_btn = ctk.CTkButton(
                batch_action_frame,
                text='开 始 批 量 转 换',
                command=self.start_batch_conversion
            )
            self.batch_convert_btn.pack(side=tk.RIGHT, padx=12, pady=(0, 12))
            return
        # 左侧面板 - 文件列表和操作
        left_panel = ttk.Frame(self.batch_frame)
        left_panel.pack(side=tk.LEFT, fill=tk.BOTH, expand=True, padx=(0, 12))  # 2026-04-01 17:35:00 修改：左右栏间距更舒展
    
        # 文件列表区域
        files_frame = ttk.LabelFrame(left_panel, text='计划转换文件列表', padding=10)
        files_frame.pack(fill=tk.BOTH, expand=True, pady=(0, 10))
    
        # 文件列表
        self.files_listbox = tk.Listbox(
            files_frame, 
            selectmode=tk.EXTENDED, 
            font=('微软雅黑', 10),
            height=10
        )
        self.files_listbox.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)
    
        # 滚动条
        scrollbar = ttk.Scrollbar(files_frame, orient=tk.VERTICAL, command=self.files_listbox.yview)
        scrollbar.pack(side=tk.RIGHT, fill=tk.Y)
        self.files_listbox.config(yscrollcommand=scrollbar.set)
    
        # 文件操作按钮
        file_buttons_frame = ttk.Frame(left_panel)
        file_buttons_frame.pack(fill=tk.X, pady=(0, 6))  # 2026-04-01 19:25:00 修改：减少底部留白，使页面更紧凑
    
        add_files_btn = ttk.Button(
            file_buttons_frame, 
            text='添加文件', 
            command=self.add_batch_files
        )
        add_files_btn.pack(side=tk.LEFT, padx=(0, 5))
    
        add_folder_btn = ttk.Button(
            file_buttons_frame, 
            text='添加文件夹', 
            command=self.add_folder_files
        )
        add_folder_btn.pack(side=tk.LEFT, padx=(0, 5))
    
        remove_btn = ttk.Button(
            file_buttons_frame, 
            text='移除选中', 
            command=self.remove_selected_files
        )
        remove_btn.pack(side=tk.LEFT, padx=(0, 5))
    
        clear_btn = ttk.Button(
            file_buttons_frame, 
            text='清空列表', 
            command=self.clear_file_list
        )
        clear_btn.pack(side=tk.LEFT)
    
        # 右侧面板 - 分上下两部分
        right_panel = ttk.Frame(self.batch_frame)
        right_panel.pack(side=tk.RIGHT, fill=tk.BOTH, expand=True, padx=(0, 0))
    
        # 上部框架：文件列表和状态信息
        list_frame = ttk.LabelFrame(right_panel, text='转换不成功文件列表', padding=10)  # 2026-04-01 17:35:00 修改：分组内边距统一
        list_frame.pack(fill=tk.BOTH, anchor=tk.N, expand=True)
    
        # 状态标签和滚动列表
        status_frame = ttk.Frame(list_frame)
        status_frame.pack(fill=tk.X, anchor=tk.N, pady=(0, 8))  # 2026-04-01 17:35:00 修改：分组间距更清爽
    
        # 待转换文件列表
        pending_frame = ttk.Frame(list_frame)
        pending_frame.pack(fill=tk.BOTH, expand=True, pady=(0, 8))  # 2026-04-01 17:35:00 修改：列表与分组间距更协调
    
        self.pending_files_listbox = tk.Listbox(
            pending_frame, 
            height=5,
            font=('微软雅黑', 9),
            selectmode=tk.SINGLE
        )
        scrollbar = ttk.Scrollbar(pending_frame, orient=tk.VERTICAL, command=self.pending_files_listbox.yview)
        self.pending_files_listbox.config(yscrollcommand=scrollbar.set)
    
        self.pending_files_listbox.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)
        scrollbar.pack(side=tk.RIGHT, fill=tk.Y)
    
        # 下部框架：输出设置
        output_settings_frame = ttk.LabelFrame(right_panel, text='输出设置', padding=10)  # 2026-04-01 17:35:00 修改：分组内边距统一
        output_settings_frame.pack(fill=tk.X, pady=(6, 0))  # 2026-04-01 19:25:00 修改：减少分组间距，使页面更紧凑
    
        # 输出路径设置
        output_path_frame = ttk.Frame(output_settings_frame)
        output_path_frame.pack(fill=tk.X, pady=0)
    
        # 标签放在上方
        ttk.Label(output_path_frame, text='输出文件夹:').pack(anchor=tk.W, pady=(0, 8))  # 2026-04-01 17:35:00 修改：标签与输入区间距更协调
    
        # 输入框和浏览按钮的容器
        path_input_frame = ttk.Frame(output_path_frame)
        path_input_frame.pack(fill=tk.X)
    
        self.batch_output_var = tk.StringVar()
        output_entry = ttk.Entry(path_input_frame, textvariable=self.batch_output_var)
        output_entry.pack(side=tk.LEFT, fill=tk.X, expand=True, padx=(0, 8))  # 2026-04-01 17:35:00 修改：输入区留白更舒展
    
        ttk.Button(
            path_input_frame,
            text='浏览...',
            command=self.select_batch_output_folder
        ).pack(side=tk.RIGHT)
    
        # 转换选项
        options_frame = ttk.Frame(output_settings_frame)
        options_frame.pack(fill=tk.X, pady=(8, 0), anchor=tk.W)  # 2026-04-01 17:35:00 修改：选项区与上方间距更清爽
    
        # 第一行：前两个选项
        first_row = ttk.Frame(options_frame)
        first_row.pack(fill=tk.X, pady=2, anchor=tk.W)
    
        self.batch_preserve_format_var = tk.BooleanVar(value=True)
        ttk.Checkbutton(
            first_row, 
            text='尽可能保留原始格式', 
            variable=self.batch_preserve_format_var
        ).pack(side=tk.LEFT, padx=10)
    
        self.batch_extract_images_var = tk.BooleanVar(value=True)
        ttk.Checkbutton(
            first_row, 
            text='提取并保存图片至images文件夹', 
            variable=self.batch_extract_images_var
        ).pack(side=tk.LEFT, padx=10)
    
        # 第二行：第三个选项
        second_row = ttk.Frame(options_frame)
        second_row.pack(fill=tk.X, pady=2, anchor=tk.W)
    
        self.batch_auto_open_folder_var = tk.BooleanVar(value=True)
        ttk.Checkbutton(
            second_row,
            text="转换完成后打开输出文件夹", 
            variable=self.batch_auto_open_folder_var
        ).pack(side=tk.LEFT, padx=10)

        third_row = ttk.Frame(options_frame)
        third_row.pack(fill=tk.X, pady=2, anchor=tk.W)

        self.batch_overwrite_policy_var = tk.StringVar(value='跳过')  # 2026-03-29 13:41:12 修改：批量同名输出策略（默认跳过防误覆盖）
        ttk.Label(third_row, text="同名输出:").pack(side=tk.LEFT, padx=(10, 6))
        ttk.Combobox(
            third_row,
            textvariable=self.batch_overwrite_policy_var,
            state='readonly',
            values=[
                '跳过',
                '覆盖',
                '自动重命名'
            ],
            width=16
        ).pack(side=tk.LEFT)
    
        # 第四部分：进度条和按钮
        batch_action_frame = ttk.Frame(right_panel)
        batch_action_frame.pack(fill=tk.X, pady=(4, 4))  # 2026-04-01 19:25:00 修改：进一步减少底部留白，使界面更紧凑
    
        # 创建状态和进度条的容器
        status_progress_frame = ttk.Frame(batch_action_frame)
        status_progress_frame.pack(fill=tk.X, pady=(0, 6))  # 2026-04-01 19:25:00 修改：进一步减少进度区下方留白，使界面更紧凑
    
        # 状态标签放在左侧
        self.batch_status_var = tk.StringVar(value='准备就绪')
        batch_status_label = ttk.Label(status_progress_frame, textvariable=self.batch_status_var)
        batch_status_label.pack(side=tk.LEFT, padx=(0, 12))  # 2026-04-01 17:35:00 修改：状态与进度条间距更舒展
    
        # 进度条占据剩余空间
        self.batch_progress = ttk.Progressbar(status_progress_frame, mode='determinate')
        self.batch_progress.pack(side=tk.LEFT, fill=tk.X, expand=True)
    
        # 转换按钮保持右对齐
        self.batch_convert_btn = ttk.Button(
            batch_action_frame, 
            text='开 始 批 量 转 换', 
            command=self.start_batch_conversion,
            style='Primary.TButton'
        )
        self.batch_convert_btn.pack(side=tk.RIGHT, ipadx=12)  # 2026-04-01 17:35:00 修改：按钮水平留白更统一

    def create_settings_tab(self):
        """创建设置选项卡内容"""
        if self.use_ctk:
            general_group, general_frame = self._ctk_group(self.settings_frame, '常规设置')
            general_group.pack(fill=tk.X, pady=(0, 10))

            default_output_frame = ctk.CTkFrame(general_frame, fg_color="transparent")
            default_output_frame.pack(fill=tk.X, pady=5)

            ctk.CTkLabel(default_output_frame, text='默认输出文件夹:').pack(side=tk.LEFT)

            self.default_output_var = tk.StringVar()
            default_output_entry = ctk.CTkEntry(default_output_frame, textvariable=self.default_output_var)
            default_output_entry.pack(side=tk.LEFT, fill=tk.X, expand=True, padx=8)

            ctk.CTkButton(default_output_frame, text='浏览...', width=90, command=self.select_default_output_folder).pack(side=tk.LEFT)

            self.auto_open_folder_var = tk.BooleanVar(value=True)
            ctk.CTkCheckBox(general_frame, text='转换完成后自动打开输出文件夹', variable=self.auto_open_folder_var).pack(anchor=tk.W, pady=4)

            self.auto_preview_var = tk.BooleanVar(value=True)
            ctk.CTkCheckBox(general_frame, text='转换完成后自动预览结果', variable=self.auto_preview_var).pack(anchor=tk.W, pady=4)

            self.embed_images_var = tk.BooleanVar(value=True)
            ctk.CTkCheckBox(
                general_frame,
                text='输出Markdown内嵌图片(Base64，不依赖images文件夹)',
                variable=self.embed_images_var
            ).pack(anchor=tk.W, pady=4)

            self.enable_path_boundary_var = tk.BooleanVar(value=True)  # 2026-03-29 13:41:12 修改：路径边界安全限制默认开启
            ctk.CTkCheckBox(
                general_frame,
                text='启用路径边界安全限制（兼容开关）',
                variable=self.enable_path_boundary_var
            ).pack(anchor=tk.W, pady=4)

            self.image_output_policy_var = tk.StringVar(value='仅原位')  # 2026-03-29 13:41:12 修改：新增图片输出策略开关（默认仅原位）
            ctk.CTkLabel(general_frame, text='图片输出策略:').pack(anchor=tk.W, pady=(10, 2))
            ctk.CTkOptionMenu(
                general_frame,
                variable=self.image_output_policy_var,
                values=[
                    '仅原位',
                    '仅补缺',
                    '全量附录'
                ],
                width=760
            ).pack(anchor=tk.W, pady=(0, 2))

            self.document_output_mode_var = tk.StringVar(value=DOCUMENT_OUTPUT_MODE_LABELS['smart'])
            mode_save_row = ctk.CTkFrame(general_frame, fg_color="transparent")
            mode_save_row.pack(fill=tk.X, pady=(10, 2))
            mode_left = ctk.CTkFrame(mode_save_row, fg_color="transparent")
            mode_left.pack(side=tk.LEFT, fill=tk.X, expand=True)
            mode_right = ctk.CTkFrame(mode_save_row, fg_color="transparent")
            mode_right.pack(side=tk.RIGHT, padx=(8, 0))

            ctk.CTkLabel(mode_left, text='文档输出模式:').pack(anchor=tk.W, pady=(0, 2))
            document_mode_cb = ctk.CTkOptionMenu(
                mode_left,
                variable=self.document_output_mode_var,
                values=list(DOCUMENT_OUTPUT_MODE_LABELS.values()),
                width=760
            )
            document_mode_cb.pack(anchor=tk.W)

            # 2026-09-01 新增：PDF 引擎策略（仅作用于智能转换模式）
            ctk.CTkLabel(mode_left, text='PDF引擎策略(智能转换模式):').pack(anchor=tk.W, pady=(8, 2))
            self.pdf_engine_strategy_var = tk.StringVar(value=PDF_ENGINE_STRATEGY_DEFAULT)
            ctk.CTkOptionMenu(
                mode_left,
                variable=self.pdf_engine_strategy_var,
                values=PDF_ENGINE_STRATEGY_LABELS,
                width=760
            ).pack(anchor=tk.W)

            ctk.CTkButton(mode_right, text='保存设置', command=self.save_settings, width=110).pack(anchor=tk.N, pady=(14, 0))  # 2026-04-01 19:40:00 修改：降低保存按钮上方留白，使页面更紧凑

            advanced_group, advanced_frame = self._ctk_group(self.settings_frame, '高级设置')
            advanced_group.pack(fill=tk.X, pady=(0, 10))

            install_frame = ctk.CTkFrame(advanced_frame, fg_color="transparent")
            install_frame.pack(fill=tk.X, pady=5)

            self.markitdown_status_var = tk.StringVar(value='检查中...')
            self.markitdown_status_label = ctk.CTkLabel(install_frame, text=self.markitdown_status_var.get())
            self.markitdown_status_label.pack(side=tk.LEFT, fill=tk.X, expand=True)

            ctk.CTkButton(install_frame, text='检查功能完整性', command=self.show_features_status).pack(side=tk.LEFT, padx=(0, 8))

            self.install_btn = ctk.CTkButton(install_frame, text='安装/更新 MarkItDown', command=self.install_markitdown)
            self.install_btn.pack(side=tk.LEFT)

            self.load_settings()
            return
        # 常规设置
        general_frame = ttk.LabelFrame(self.settings_frame, text='常规设置', padding=10)
        general_frame.pack(fill=tk.X, pady=(0, 10))  # 2026-04-01 19:25:00 修改：减少分组间距，使页面更紧凑
        
        # 默认输出文件夹
        default_output_frame = ttk.Frame(general_frame)
        default_output_frame.pack(fill=tk.X, pady=(0, 8))  # 2026-04-01 17:50:00 修改：行间距更协调
        
        ttk.Label(default_output_frame, text='默认输出文件夹:').pack(side=tk.LEFT)
        
        self.default_output_var = tk.StringVar()
        default_output_entry = ttk.Entry(default_output_frame, textvariable=self.default_output_var, width=40)
        default_output_entry.pack(side=tk.LEFT, fill=tk.X, expand=True, padx=8)  # 2026-04-01 17:50:00 修改：输入区留白更清爽
        
        default_browse_btn = ttk.Button(
            default_output_frame, 
            text='浏览...', 
            command=self.select_default_output_folder
        )
        default_browse_btn.pack(side=tk.LEFT)
        
        # 自动打开输出文件夹
        self.auto_open_folder_var = tk.BooleanVar(value=True)
        auto_open_folder_cb = ttk.Checkbutton(
            general_frame, 
            text='转换完成后自动打开输出文件夹', 
            variable=self.auto_open_folder_var
        )
        auto_open_folder_cb.pack(anchor=tk.W, pady=4)  # 2026-04-01 17:50:00 修改：选项留白更清爽
        
        # 自动预览转换结果
        self.auto_preview_var = tk.BooleanVar(value=True)
        auto_preview_cb = ttk.Checkbutton(
            general_frame, 
            text='转换完成后自动预览结果', 
            variable=self.auto_preview_var
        )
        auto_preview_cb.pack(anchor=tk.W, pady=4)  # 2026-04-01 17:50:00 修改：选项留白更清爽

        self.embed_images_var = tk.BooleanVar(value=True)
        embed_images_cb = ttk.Checkbutton(
            general_frame,
            text='输出Markdown内嵌图片(Base64，不依赖images文件夹)',
            variable=self.embed_images_var
        )
        embed_images_cb.pack(anchor=tk.W, pady=4)  # 2026-04-01 17:50:00 修改：选项留白更清爽

        self.enable_path_boundary_var = tk.BooleanVar(value=True)  # 2026-03-29 13:41:12 修改：路径边界安全限制默认开启
        path_boundary_cb = ttk.Checkbutton(
            general_frame,
            text='启用路径边界安全限制（兼容开关）',
            variable=self.enable_path_boundary_var
        )
        path_boundary_cb.pack(anchor=tk.W, pady=4)  # 2026-04-01 17:50:00 修改：选项留白更清爽

        self.image_output_policy_var = tk.StringVar(value='仅原位')  # 2026-03-29 13:41:12 修改：新增图片输出策略开关（默认仅原位）
        ttk.Label(general_frame, text='图片输出策略:').pack(anchor=tk.W, pady=(10, 4))  # 2026-04-01 17:50:00 修改：标签留白更协调
        image_policy_cb = ttk.Combobox(
            general_frame,
            textvariable=self.image_output_policy_var,
            state='readonly',
            values=[
                '仅原位',
                '仅补缺',
                '全量附录'
            ],
            width=60
        )
        image_policy_cb.pack(anchor=tk.W, pady=(0, 2))

        self.document_output_mode_var = tk.StringVar(value=DOCUMENT_OUTPUT_MODE_LABELS['smart'])
        mode_save_row = ttk.Frame(general_frame)
        mode_save_row.pack(fill=tk.X, pady=(10, 4))  # 2026-04-01 17:50:00 修改：分组内留白更清爽
        mode_left = ttk.Frame(mode_save_row)
        mode_left.pack(side=tk.LEFT, fill=tk.X, expand=True)
        mode_right = ttk.Frame(mode_save_row)
        mode_right.pack(side=tk.RIGHT, padx=(8, 0))

        ttk.Label(mode_left, text='文档输出模式:').pack(anchor=tk.W, pady=(0, 2))
        document_mode_cb = ttk.Combobox(
            mode_left,
            textvariable=self.document_output_mode_var,
            state='readonly',
            values=list(DOCUMENT_OUTPUT_MODE_LABELS.values()),
            width=60
        )
        document_mode_cb.pack(anchor=tk.W)

        # 2026-09-01 新增：PDF 引擎策略（仅作用于智能转换模式）
        ttk.Label(mode_left, text='PDF引擎策略(智能转换模式):').pack(anchor=tk.W, pady=(8, 2))
        self.pdf_engine_strategy_var = tk.StringVar(value=PDF_ENGINE_STRATEGY_DEFAULT)
        pdf_engine_cb = ttk.Combobox(
            mode_left,
            textvariable=self.pdf_engine_strategy_var,
            state='readonly',
            values=PDF_ENGINE_STRATEGY_LABELS,
            width=60
        )
        pdf_engine_cb.pack(anchor=tk.W)
        
        save_settings_btn = ttk.Button(
            mode_right,
            text='保存设置', 
            command=self.save_settings,
            style='Primary.TButton'
        )
        save_settings_btn.pack(anchor=tk.N, pady=(14, 0))  # 2026-04-01 19:40:00 修改：降低保存按钮上方留白，使页面更紧凑
        
        # 高级设置
        advanced_frame = ttk.LabelFrame(self.settings_frame, text='高级设置', padding=10)
        advanced_frame.pack(fill=tk.X, pady=(0, 10))  # 2026-04-01 19:25:00 修改：减少分组间距，使页面更紧凑
        
        # 安装/更新 MarkItDown 和功能检查
        install_frame = ttk.Frame(advanced_frame)
        install_frame.pack(fill=tk.X, pady=(0, 8))  # 2026-04-01 17:50:00 修改：行间距更协调
        
        self.markitdown_status_var = tk.StringVar(value='检查中...')
        markitdown_status_label = ttk.Label(install_frame, textvariable=self.markitdown_status_var)
        markitdown_status_label.pack(side=tk.LEFT, fill=tk.X, expand=True)
        
        # 检查功能完整性按钮
        check_features_btn = ttk.Button(
            install_frame, 
            text='检查功能完整性', 
            command=self.show_features_status
        )
        check_features_btn.pack(side=tk.LEFT, padx=(0, 8))  # 2026-04-01 17:50:00 修改：按钮间距更统一
        
        # 安装/更新 MarkItDown 按钮
        self.install_btn = ttk.Button(
            install_frame, 
            text='安装/更新 MarkItDown', 
            command=self.install_markitdown
        )
        self.install_btn.pack(side=tk.LEFT)
        
        # 加载设置
        self.load_settings()

    def create_about_tab(self):
        """创建关于选项卡内容"""
        if self.use_ctk:
            about_content = ctk.CTkFrame(self.about_frame)
            about_content.pack(fill=tk.BOTH, expand=True, padx=20, pady=20)

            app_title = ctk.CTkLabel(
                about_content,
            text=f'软件版本 v{APP_VERSION}',  # 2026-09-02 修改：关于页软件版本改为项目自身版本APP_VERSION（原为MarkItDown引擎版本）
            font=('微软雅黑', 15, 'bold')  # 2026-04-01 18:10:00 修改：字号更小一号
            )
            app_title.pack(pady=(0, 10))

            version_label = ctk.CTkLabel(about_content, text=f'{APP_VERSION_DATE} HNZZ', font=('微软雅黑', 10))  # 2026-09-02 修改：版本日期改为APP_VERSION_DATE常量（统一YYYYMMDD）
            version_label.pack()

            description = (
                "AnyToMarkDown 工具用于将各种文档格式转换为 Markdown 格式。\n\n"
                "支持的文件格式包括：\n"
                "• PDF（含文本型/扫描件/版式保持等五种输出模式）\n"
                "• PowerPoint（.ppt/.pptx）\n"
                "• Word（.doc/.docx）\n"
                "• Excel（.xls/.xlsx）\n"
                "• 图片 (EXIF 元数据和 OCR)\n"
                "• 音频 (EXIF 元数据和语音转录)\n"
                "• HTML、HTM、MHTML\n"
                "• 文本格式 (CSV, JSON, XML)\n"
                "• ZIP 文件 (遍历内容)\n\n"
                "增强引擎（2026-09-01 新增）：\n"
                "• AnyDoc 兜底：Office 文档在其他方法均失败时自动启用，保留标题层级与表格结构\n"
                "• pdf-inspector 引擎：纯文本型 PDF 的智能识别与高质量 Markdown 输出（设置页可选自动/图文混排优先/pdf-inspector优先）\n\n"
                "• 本应用基于Python和Tkinter开发，使用了微软的MarkItDown库、Firecrawl AnyDoc、Firecrawl pdf-inspector 以及其他多个与文档转换相关的库。一并表示感谢！安装MarkItDown库请使用以下命令：pip install -e 'packages/markitdown[all]' 详情见该项目页介绍。\n"
                "本软件为个人开发，产权属于开发者本人（本页所示作者名称为昵称，权利人为昵称对应的实际开发者）。作者许可您一项个人的、可撤销的、不可转让的、非独占地和非商业的合法使用本产品的权利，您不享有本产品的所有权。作者基于本协议对您的授权仅为授权您个人以非商业的目的对于本产品进行使用，任何超出个人使用目的的使用行为都必须另行获得作者本人具体的、单独的、书面的授权，协议未明示授权的其他一切权利仍由我方保留，您在行使该等权利前须另行获得我方的书面许可，同时我方如未行使前述任何权利，并不构成对该权利的放弃。严禁任何单位和个人未经授权将软件（或将项目改头换面）用于营利或非法用途，否则将追究法律责任。本软件为个人开发，开发者不承担任何责任，请用户自行承担使用风险。\n"
                "💖🌹感谢您的使用！欢迎提出宝贵意见！😊 "
            )

            desc_text = ctk.CTkTextbox(
                about_content,
                font=('微软雅黑', 10)
            )
            desc_text.pack(fill=tk.BOTH, expand=True, pady=10)
            desc_text.insert(tk.END, description)
            desc_text.configure(state=tk.DISABLED)

            anytomd_links_frame = ctk.CTkFrame(about_content, fg_color="transparent")
            anytomd_links_frame.pack(fill=tk.X, pady=5)

            ctk.CTkLabel(anytomd_links_frame, text='AnyToMarkDown项目地址：').pack(side=tk.LEFT)
            anytomd_github_link = ctk.CTkLabel(anytomd_links_frame, text='AnyToMarkDown GitHub', text_color='blue')
            anytomd_github_link.pack(side=tk.LEFT, padx=8)
            anytomd_github_link.bind('<Button-1>', lambda e: webbrowser.open_new('https://github.com/yihufree/AnyToMarkdown'))

            links_frame = ctk.CTkFrame(about_content, fg_color="transparent")
            links_frame.pack(fill=tk.X, pady=10)

            ctk.CTkLabel(links_frame, text='MarkItDown 项目地址:').pack(side=tk.LEFT)

            pypi_link = ctk.CTkLabel(links_frame, text='Microsoft markitdown PyPI', text_color='blue')
            pypi_link.pack(side=tk.LEFT, padx=8)
            pypi_link.bind('<Button-1>', lambda e: webbrowser.open_new('https://pypi.org/project/markitdown/'))

            ctk.CTkLabel(links_frame, text='|').pack(side=tk.LEFT, padx=8)

            github_link = ctk.CTkLabel(links_frame, text='Microsoft markitdown GitHub', text_color='blue')
            github_link.pack(side=tk.LEFT, padx=8)
            github_link.bind('<Button-1>', lambda e: webbrowser.open_new('https://github.com/microsoft/markitdown'))
            return
        about_content = ttk.Frame(self.about_frame, padding=16)  # 2026-04-01 19:25:00 修改：减少内容区留白，使页面更紧凑
        about_content.pack(fill=tk.BOTH, expand=True)
        
        # 应用标题
        app_title = ttk.Label(
            about_content, 
            text=f'软件版本 v{APP_VERSION}',  # 2026-09-02 修改：关于页软件版本改为项目自身版本APP_VERSION（原为MarkItDown引擎版本）
            font=('微软雅黑', 15, 'bold')  # 2026-04-01 18:10:00 修改：字号更小一号
        )
        app_title.pack(pady=(0, 10))
        
        # 版本信息
        version_label = ttk.Label(
            about_content, 
            text=f'{APP_VERSION_DATE} HNZZ', # 2026-09-02 修改：版本日期改为APP_VERSION_DATE常量（统一YYYYMMDD）
            font=('微软雅黑', 10)
        )
        version_label.pack()
        
        # 分隔线
        separator = ttk.Separator(about_content, orient=tk.HORIZONTAL)
        separator.pack(fill=tk.X, pady=16)  # 2026-04-01 17:50:00 修改：分隔线留白更克制
        
        # 应用说明
        description = (
            "AnyToMarkDown 工具用于将各种文档格式转换为 Markdown 格式。\n\n"
            "支持的文件格式包括：\n"
            "• PDF（含文本型/扫描件/版式保持等五种输出模式）\n"
            "• PowerPoint（.ppt/.pptx）\n"
            "• Word（.doc/.docx）\n"
            "• Excel（.xls/.xlsx）\n"
            "• 图片 (EXIF 元数据和 OCR)\n"
            "• 音频 (EXIF 元数据和语音转录)\n"
            "• HTML、HTM、MHTML\n"
            "• 文本格式 (CSV, JSON, XML)\n"
            "• ZIP 文件 (遍历内容)\n\n"
            "增强引擎（2026-09-01 新增）：\n"
            "• AnyDoc 兜底：Office 文档在其他方法均失败时自动启用，保留标题层级与表格结构\n"
            "• pdf-inspector 引擎：纯文本型 PDF 的智能识别与高质量 Markdown 输出（设置页可选自动/图文混排优先/pdf-inspector优先）\n\n"
            "本应用基于 Python 和 Tkinter 开发，使用了微软的MarkItDown库、Firecrawl AnyDoc、Firecrawl pdf-inspector 以及其他多个与文档转换相关的库。一并表示感谢！\n"
            "本软件为个人开发，产权属于开发者本人（本页所示作者名称为昵称，权利人为昵称对应的实际开发者）。作者许可您一项个人的、可撤销的、不可转让的、非独占地和非商业的合法使用本产品的权利，您不享有本产品的所有权。作者基于本协议对您的授权仅为授权您个人以非商业的目的对于本产品进行使用，任何超出个人使用目的的使用行为都必须另行获得作者本人具体的、单独的、书面的授权，协议未明示授权的其他一切权利仍由我方保留，您在行使该等权利前须另行获得我方的书面许可，同时我方如未行使前述任何权利，并不构成对该权利的放弃。严禁任何单位和个人未经授权将软件（或将项目改头换面）用于营利或非法用途，否则将追究法律责任。本软件为个人开发，开发者不承担任何责任，请用户自行承担使用风险。\n"
            "💖🌹感谢您的使用！欢迎提出宝贵意见！😊 "
        )
        
        desc_text = scrolledtext.ScrolledText(
            about_content, 
            wrap=tk.WORD, 
            width=60, 
            height=12, 
            font=('微软雅黑', 10)
        )
        desc_text.pack(fill=tk.BOTH, expand=True, pady=(10, 10))  # 2026-04-01 19:25:00 修改：减少说明区上下留白，使页面更紧凑
        desc_text.insert(tk.END, description)
        desc_text.config(state=tk.DISABLED)
        
        # 本项目GitHub链接
        anytomd_links_frame = ttk.Frame(about_content)
        anytomd_links_frame.pack(fill=tk.X, pady=(8, 6))  # 2026-04-01 17:50:00 修改：链接区留白更协调
        
        ttk.Label(anytomd_links_frame, text='AnyToMarkDown项目地址：').pack(side=tk.LEFT)
        
        anytomd_github_link = ttk.Label(
            anytomd_links_frame,
            text='AnyToMarkDown GitHub',
            foreground='blue',
            cursor='hand2'
        )
        anytomd_github_link.pack(side=tk.LEFT, padx=5)
        anytomd_github_link.bind('<Button-1>', lambda e: webbrowser.open_new('https://github.com/yihufree/AnyToMarkdown'))
        
        links_frame = ttk.Frame(about_content)
        links_frame.pack(fill=tk.X, pady=(10, 0))  # 2026-04-01 17:50:00 修改：底部留白更协调
        
        ttk.Label(links_frame, text='MarkItDown 项目地址:').pack(side=tk.LEFT)
        
        # PyPI 链接
        pypi_link = ttk.Label(
            links_frame, 
            text='Microsoft markitdown PyPI',
            foreground='blue', 
            cursor='hand2'
        )
        pypi_link.pack(side=tk.LEFT, padx=5)
        pypi_link.bind('<Button-1>', lambda e: webbrowser.open_new('https://pypi.org/project/markitdown/'))
        
        ttk.Label(links_frame, text='|').pack(side=tk.LEFT, padx=5)
        
        # GitHub 链接
        github_link = ttk.Label(
            links_frame, 
            text='Microsoft markitdown GitHub',
            foreground='blue', 
            cursor='hand2'
        )
        github_link.pack(side=tk.LEFT, padx=5)
        github_link.bind('<Button-1>', lambda e: webbrowser.open_new('https://github.com/microsoft/markitdown'))

    def create_statusbar(self):
        """创建状态栏"""
        if self.use_ctk:
            self.statusbar = ctk.CTkFrame(self.root)
            self.statusbar.pack(side=tk.BOTTOM, fill=tk.X, pady=(0, 0))

            self.status_var = tk.StringVar(value='就绪')
            self.status_label = ctk.CTkLabel(self.statusbar, text=self.status_var.get())
            self.status_label.pack(side=tk.LEFT, padx=12, pady=6)

            self.version_label = ctk.CTkLabel(self.statusbar, text=f'v{APP_VERSION}')  # 2026-09-02 修改：状态栏右下角版本号改为项目自身版本APP_VERSION（原为MarkItDown引擎版本）
            self.version_label.pack(side=tk.RIGHT, padx=12, pady=6)
            return
        self.statusbar = ttk.Frame(self.root, relief=tk.SUNKEN, padding=(10, 4))  # 2026-04-01 19:40:00 修改：状态栏留白再收紧一点，使页面更紧凑
        self.statusbar.pack(side=tk.BOTTOM, fill=tk.X, pady=(0, 0))
        
        self.status_var = tk.StringVar(value='就绪')
        status_label = ttk.Label(self.statusbar, textvariable=self.status_var)
        status_label.pack(side=tk.LEFT)
        
        version_label = ttk.Label(self.statusbar, text=f'v{APP_VERSION}')  # 2026-09-02 修改：状态栏右下角版本号改为项目自身版本APP_VERSION（原为MarkItDown引擎版本）
        version_label.pack(side=tk.RIGHT)

    def show_features_status(self):
        """显示功能完整性检查结果"""
        # 更准确的模块与功能映射关系
        required_modules = {
            'python-docx': {
                'name': 'Word文档支持 (DOCX)',
                'desc': '支持 .docx 格式文档的转换',
                'module': 'docx'
            },
            'docx2python': {
                'name': 'Word文档支持 (DOC)',
                'desc': '通过 docx2python 提供 .doc 格式文档的部分支持',
                'module': 'docx2python'
            },
            'pywin32': {
                'name': 'Word文档转换支持',
                'desc': '通过 win32com 支持 .doc 转 .docx 的自动转换',
                'module': 'win32com'
            },
            'pdfminer.six': {
                'name': 'PDF文档基础支持',
                'desc': '支持 PDF 文档的文本内容提取',
                'module': 'pdfminer'
            },
            'pdf2image': {
                'name': 'PDF图像提取支持',
                'desc': '支持从 PDF 提取图像内容',
                'module': 'pdf2image'
            },
            'Pillow': {
                'name': '图像处理支持',
                'desc': '支持图像处理和格式转换',
                'module': 'PIL'
            },
            'openpyxl': {
                'name': 'Excel文档支持',
                'desc': '支持 .xlsx 格式文档的转换',
                'module': 'openpyxl'
            },
            'python-pptx': {
                'name': 'PowerPoint文档支持',
                'desc': '支持 .pptx 格式文档的转换',
                'module': 'pptx'
            },
            'beautifulsoup4': {
                'name': 'HTML文档支持',
                'desc': '支持HTML文档的解析和转换',
                'module': 'bs4'
            },
            'lxml': {
                'name': 'XML解析支持',
                'desc': '提供高级XML解析能力',
                'module': 'lxml'
            },
            'pytesseract': {
                'name': 'OCR文字识别支持',
                'desc': '支持从图像中提取文字',
                'module': 'pytesseract'
            },
            'pydub': {
                'name': '音频处理支持',
                'desc': '支持音频文件处理',
                'module': 'pydub'
            },
            'SpeechRecognition': {
                'name': '语音识别支持',
                'desc': '支持将语音转录为文本',
                'module': 'speech_recognition'
            },
            'firecrawl-anydoc': {  # 2026-09-01 新增：AnyDoc 兜底引擎
                'name': 'AnyDoc 兜底引擎',
                'desc': 'Office 文档其他方法失败时自动启用（保留标题层级与表格结构）',
                'module': 'anydoc'
            },
            'pdf-inspector': {  # 2026-09-01 新增：pdf-inspector 引擎
                'name': 'pdf-inspector 引擎',
                'desc': '纯文本型 PDF 的智能识别与高质量 Markdown 输出',
                'module': 'pdf_inspector'
            }
        }
        
        # 获取markitdown版本和命令行功能
        version_info = ""
        cmd_features = []
        try:
            module = _get_markitdown()
            version_info = f"MarkItDown 版本: {_get_markitdown_version()}" if module is not None else "MarkItDown 未安装"
            help_text = ""
            try:
                help_process = subprocess.run(
                    ["markitdown", "--help"],
                    capture_output=True,
                    text=True,
                    timeout=5,
                    creationflags=subprocess.CREATE_NO_WINDOW
                )
                help_text = help_process.stdout + help_process.stderr
            except Exception:
                help_text = ""
            
            # 分析命令行支持的功能
            if "--extract-images" in help_text or "--images" in help_text:
                cmd_features.append("✓ 命令行图像提取支持")
            else:
                cmd_features.append("✗ 命令行图像提取不支持")
                
            if "pdf" in help_text.lower():
                cmd_features.append("✓ 命令行PDF支持")
            else:
                cmd_features.append("✗ 命令行PDF不支持")
                
            if "docx" in help_text.lower():
                cmd_features.append("✓ 命令行DOCX支持")
            else:
                cmd_features.append("✗ 命令行DOCX不支持 (本软件通过API直接支持DOCX，无需命令行)")
                
            # 分析命令行支持的选项
            options = re.findall(r'--([\w-]+)', help_text)
            cmd_features.append(f"可用选项: {', '.join(options)}")
                
        except Exception as e:
            version_info = f"无法获取MarkItDown版本: {str(e)}"
        
        # 验证模块和功能
        status = []
        
        for pkg_name, pkg_info in required_modules.items():
            module_name = pkg_info['module']
            module_display_name = pkg_info['name']
            module_desc = pkg_info['desc']
            
            try:
                # 尝试导入模块
                __import__(module_name)
                
                # 进行更深入的功能检测
                if module_name == 'docx':
                    # 检测python-docx的功能
                    try:
                        import docx
                        docx.Document
                        status.append(f"✓ {module_display_name} - 已完整支持")
                    except Exception:
                        status.append(f"⚠ {module_display_name} - 模块存在但功能受限")
                
                elif module_name == 'win32com':
                    # 检测pywin32的功能
                    try:
                        import win32com.client
                        status.append(f"✓ {module_display_name} - 已完整支持")
                    except Exception:
                        status.append(f"⚠ {module_display_name} - 模块存在但功能受限")
                
                elif module_name == 'pdf2image':
                    # 检测pdf2image功能是否完整
                    try:
                        import pdf2image
                        # 尝试使用核心功能，这会检查poppler依赖是否可用
                        test_success = False
                        try:
                            # 创建一个非常小的PDF样本用于测试
                            with tempfile.NamedTemporaryFile(suffix='.pdf', delete=False) as pdf_file:
                                pdf_path = pdf_file.name
                                try:
                                    import fitz
                                    doc = fitz.open()
                                    doc.new_page(width=200, height=200)
                                    doc.save(pdf_path)
                                    doc.close()
                                except Exception:
                                    pdf_file.write(b"%PDF-1.0\n%EOF\n")
                        
                            try:
                                pdf2image.pdfinfo_from_path(pdf_path)
                                test_success = True
                            except Exception:
                                poppler_bin = self.find_poppler_bin()
                                if poppler_bin:
                                    pdf2image.pdfinfo_from_path(pdf_path, poppler_path=poppler_bin)
                                    test_success = True
                        except Exception:
                            pass
                        finally:
                            # 清理临时文件
                            try:
                                os.unlink(pdf_path)
                            except:
                                pass
                        
                        if test_success:
                            status.append(f"✓ {module_display_name} - 已完整支持 (包括Poppler)")
                        else:
                            # 再次尝试带有内置poppler路径的测试
                            poppler_path = self.find_poppler_bin()
                            if poppler_path:
                                status.append(f"✓ {module_display_name} - 已完整支持 (使用内置Poppler)")
                            else:
                                status.append(f"⚠ {module_display_name} - 模块存在但缺少Poppler依赖")
                    except Exception:
                        status.append(f"⚠ {module_display_name} - 模块存在但功能受限")
                
                elif module_name == 'pytesseract':
                    # 检测OCR功能
                    try:
                        import pytesseract
                        # 检查tesseract是否可用
                        pytesseract.get_tesseract_version()
                        status.append(f"✓ {module_display_name} - 已完整支持")
                    except Exception:
                        status.append(f"⚠ {module_display_name} - 模块存在但缺少Tesseract OCR")
                
                else:
                    # 一般模块，仅检查导入成功
                    status.append(f"✓ {module_display_name} - 已支持")
                
            except ImportError:
                status.append(f"✗ {module_display_name} - 未支持")
        
        # 文件类型支持检测
        file_type_status = []
        
        # 检测基本文件类型支持
        file_types = {
            '.pdf': '基于pdfminer和pdf2image',
            '.docx': '基于python-docx',
            '.doc': '基于pywin32自动转换',
            '.pptx': '基于python-pptx',
            '.xlsx': '基于openpyxl',
            '.png': '基于PIL',
            '.jpg': '基于PIL',
            '.jpeg': '基于PIL',
            '.html': '基于beautifulsoup4',
            '.htm': '基于beautifulsoup4',
            '.mp3': '基于pydub和SpeechRecognition',
            '.wav': '基于pydub和SpeechRecognition',
            '.ppt': '基于AnyDoc兜底引擎(2026-09-01新增)',
            '.xls': '基于AnyDoc兜底引擎(2026-09-01新增)',
        }
        
        for ext, desc in file_types.items():
            supported = True
            required_modules = []
            
            # 根据文件类型设置所需模块
            if ext == '.pdf':
                required_modules = ['pdfminer', 'pdf2image']
            elif ext == '.docx':
                required_modules = ['docx']
            elif ext == '.doc':
                required_modules = ['win32com', 'docx2python']
            elif ext == '.pptx':
                required_modules = ['pptx']
            elif ext == '.xlsx':
                required_modules = ['openpyxl']
            elif ext in ['.png', '.jpg', '.jpeg']:
                required_modules = ['PIL']
            elif ext in ['.html', '.htm']:
                required_modules = ['bs4', 'lxml']
            elif ext in ['.mp3', '.wav']:
                required_modules = ['pydub', 'speech_recognition']
            elif ext in ['.ppt', '.xls']:
                required_modules = ['anydoc']
            
            # 检查所需模块是否都存在
            for module in required_modules:
                try:
                    __import__(module)
                except ImportError:
                    supported = False
                    break
            
            if supported:
                file_type_status.append(f"✓ {ext} 文件类型支持 ({desc})")
            else:
                file_type_status.append(f"✗ {ext} 文件类型不完全支持 (缺少依赖)")
        
        # 显示结果
        result_window = tk.Toplevel(self.root)
        result_window.title('功能完整性检查')
        result_window.geometry('600x600')
        
        result_text = scrolledtext.ScrolledText(
            result_window, 
            wrap=tk.WORD, 
            font=('微软雅黑', 10)
        )
        result_text.pack(fill=tk.BOTH, expand=True, padx=10, pady=10)
        
        result_text.insert(tk.END, "MarkItDown功能完整性检查结果:\n\n")
        result_text.insert(tk.END, f"{version_info}\n\n")
        
        result_text.insert(tk.END, "命令行工具功能 (注: 本软件通过API直接处理，不依赖命令行受限功能):\n")
        for line in cmd_features:
            result_text.insert(tk.END, line + "\n")
        
        result_text.insert(tk.END, "\n依赖模块状态:\n")
        for line in status:
            result_text.insert(tk.END, line + "\n")
        
        result_text.insert(tk.END, "\n文件类型支持状态:\n")
        for line in file_type_status:
            result_text.insert(tk.END, line + "\n")
        
        # 添加GUI自体功能状态
        result_text.insert(tk.END, "\nGUI功能状态:\n")
        result_text.insert(tk.END, "✓ 单文件转换 - 支持所有文件类型\n")
        result_text.insert(tk.END, "✓ 批量转换 - 支持所有文件类型\n")
        result_text.insert(tk.END, "✓ 图片提取增强 - 通过多种方法支持图片提取\n")
        result_text.insert(tk.END, "✓ 预览功能 - 支持转换后的Markdown预览\n")
        # 2026-09-01 新增：增强引擎功能状态
        try:
            import anydoc
            result_text.insert(tk.END, f"✓ AnyDoc兜底引擎 - 已就绪 (版本 {getattr(anydoc, '__version__', '未知')})\n")
        except Exception:
            result_text.insert(tk.END, "⚠ AnyDoc兜底引擎 - 未安装，Office兜底转换不可用\n")
        try:
            import pdf_inspector
            result_text.insert(tk.END, f"✓ pdf-inspector引擎 - 已就绪 (版本 {getattr(pdf_inspector, '__version__', '未知')})\n")
        except Exception:
            result_text.insert(tk.END, "⚠ pdf-inspector引擎 - 未安装，纯文本PDF引擎策略不可用\n")
        
        # 添加建议
        result_text.insert(tk.END, "\n功能增强建议:\n")
        
        # 检查并提供建议
        suggestions = []
        
        try:
            import pdf2image
            try:
                poppler_path = self.find_poppler_bin()
                if not poppler_path:
                    raise Exception("poppler not found")
            except:
                suggestions.append("- 建议：虽然内置了部分处理，但在某些环境下安装外部Poppler工具可增强PDF图像提取: https://github.com/oschwartz10612/poppler-windows/releases")
        except ImportError:
            suggestions.append("- 安装pdf2image模块以支持PDF图像提取: pip install pdf2image")
        
        try:
            import pytesseract
            try:
                pytesseract.get_tesseract_version()
            except:
                suggestions.append("- 安装Tesseract OCR以支持图像文字识别: https://github.com/UB-Mannheim/tesseract/wiki")
        except ImportError:
            suggestions.append("- 安装pytesseract模块以支持OCR: pip install pytesseract")
        
        # 2026-09-01 新增：增强引擎缺失建议
        try:
            import anydoc
        except ImportError:
            suggestions.append("- 安装AnyDoc兜底引擎以增强Office文档转换: pip install firecrawl-anydoc")
        try:
            import pdf_inspector
        except ImportError:
            suggestions.append("- 安装pdf-inspector引擎以增强纯文本PDF转换: pip install pdf-inspector")
        
        # 添加通用建议
        suggestions.append("- 使用命令 'pip install markitdown[all]' 安装所有依赖")
        suggestions.append("- 确保系统PATH中包含所有必要的外部工具")
        
        for suggestion in suggestions:
            result_text.insert(tk.END, suggestion + "\n")
        
        result_text.config(state=tk.DISABLED)

    def show_install_prompt(self):
        """显示安装提示"""
        response = messagebox.askyesno(
            '安装 MarkItDown',
            'MarkItDown 未安装。是否立即安装？\n\n'
            '这是使用此应用程序所必需的。'
        )
        
        if response:
            self.install_markitdown()

    def install_markitdown(self):
        """在打包后的EXE中，此功能不应该被调用"""
        messagebox.showinfo(
            "提示",
            "此版本已包含所需组件，无需安装。\n"
            "如果遇到问题，请联系软件开发者。"
        )

    def select_input_file(self):
        """选择输入文件"""
        file_path = filedialog.askopenfilename(
            title='选择要转换的文件',
            filetypes=self.file_types
        )
        
        if file_path:
            self.input_path_var.set(file_path)
            self.current_file = file_path
            
            # 如果没有设置输出文件夹，则使用默认输出文件夹或输入文件所在文件夹
            if not self.output_path_var.get():
                if self.default_output_var.get():
                    self.output_path_var.set(self.default_output_var.get())
                else:
                    self.output_path_var.set(os.path.dirname(file_path))
            
            self.status_var.set(f'已选择文件: {os.path.basename(file_path)}')

    def select_output_folder(self):
        """选择输出文件夹"""
        folder_path = filedialog.askdirectory(title='选择输出文件夹')
        
        if folder_path:
            self.output_path_var.set(folder_path)
            self.output_folder = folder_path
            self.status_var.set(f'输出文件夹: {folder_path}')

    def select_batch_output_folder(self):
        """选择批量转换的输出文件夹"""
        folder_path = filedialog.askdirectory(title='选择批量转换输出文件夹')
        
        if folder_path:
            self.batch_output_var.set(folder_path)
            self.status_var.set(f'批量转换输出文件夹: {folder_path}')

    def select_default_output_folder(self):
        """选择默认输出文件夹"""
        folder_path = filedialog.askdirectory(title='选择默认输出文件夹')
        
        if folder_path:
            self.default_output_var.set(folder_path)
            self.status_var.set(f'默认输出文件夹: {folder_path}')

    def add_batch_files(self):
        """添加批量转换文件"""
        file_paths = filedialog.askopenfilenames(
            title='选择要批量转换的文件',
            filetypes=self.file_types
        )
        
        if file_paths:
            for file_path in file_paths:
                if file_path not in self.batch_files:
                    self.batch_files.append(file_path)
                    self.files_listbox.insert(tk.END, os.path.basename(file_path))
            
            self.status_var.set(f'已添加 {len(file_paths)} 个文件到批量转换列表')

    def add_folder_files(self):
        """添加文件夹中的所有支持文件"""
        folder_path = filedialog.askdirectory(title='选择包含要转换文件的文件夹')
        
        if not folder_path:
            return
        
        # 获取所有支持的文件扩展名
        extensions = []
        for _, ext_list in self.file_types[1:]:  # 跳过第一个"所有支持的文件"
            extensions.extend(ext_list.split(';'))
        
        # 去掉 * 号
        extensions = [ext[1:] for ext in extensions]
        
        added_count = 0
        
        for root, _, files in os.walk(folder_path):
            for file in files:
                file_ext = os.path.splitext(file)[1].lower()
                if file_ext in extensions:
                    file_path = os.path.join(root, file)
                    if file_path not in self.batch_files:
                        self.batch_files.append(file_path)
                        self.files_listbox.insert(tk.END, os.path.basename(file))
                        added_count += 1
        
        self.status_var.set(f'已从文件夹添加 {added_count} 个文件到批量转换列表')

    def remove_selected_files(self):
        """移除选中的文件"""
        selected_indices = self.files_listbox.curselection()
        
        if not selected_indices:
            return
        
        # 从后往前删除，避免索引变化
        for i in sorted(selected_indices, reverse=True):
            del self.batch_files[i]
            self.files_listbox.delete(i)
        
        self.status_var.set(f'已移除 {len(selected_indices)} 个文件')

    def clear_file_list(self):
        """清空文件列表"""
        self.batch_files = []
        self.files_listbox.delete(0, tk.END)
        self.status_var.set('已清空文件列表')

    def start_conversion(self):
        """开始单文件转换"""
        if self.conversion_in_progress:
            messagebox.showinfo('提示', '转换正在进行中，请等待完成')
            return
        
        input_file = self.input_path_var.get()
        output_folder = self.output_path_var.get()
        
        if not input_file:
            messagebox.showwarning('警告', '请选择要转换的文件')
            return
        
        if not output_folder:
            messagebox.showwarning('警告', '请选择输出文件夹')
            return
        
        if not os.path.exists(input_file):
            messagebox.showerror('错误', '输入文件不存在')
            return
        
        if not os.path.exists(output_folder):
            try:
                os.makedirs(output_folder)
            except Exception as e:
                messagebox.showerror('错误', f'创建输出文件夹失败: {str(e)}')
                return

        # 以下为新添加文件覆盖检验
        # ... existing code ...
        output_filename = os.path.splitext(os.path.basename(input_file))[0] + '.md'
        output_path = os.path.join(output_folder, output_filename)
        
        # 新增覆盖校验
        if os.path.exists(output_path):
            response = messagebox.askyesno("文件存在", f"{output_path} 已存在，是否覆盖？")
            if not response:
                self.status_var.set("用户取消覆盖操作")
                return
        # ... existing code ...        
        # 新添加文件覆盖检验结束

        # 在新线程中启动转换
        threading.Thread(target=self.convert_file, daemon=True).start()

    def handle_doc_file(self, doc_file):
        """将.doc文件转换为.docx格式"""
        try:
            # 显示转换提示
            self._update_ui(status='正在转换 DOC 到 DOCX...')
            
            # 生成输出路径
            docx_file = os.path.splitext(doc_file)[0] + '.docx'
            
            # 使用更健壮的方法转换文件
            import win32com.client
            import pythoncom
            
            # 初始化COM环境
            pythoncom.CoInitialize()
            
            word = None
            doc = None
            try:
                # 创建Word应用实例
                word = win32com.client.Dispatch("Word.Application")
                word.Visible = False
                
                # 尝试打开文档
                doc = word.Documents.Open(doc_file)
                
                # 保存为.docx
                doc.SaveAs2(docx_file, FileFormat=16)  # 16 代表 .docx 格式
                
                self._update_ui(status='DOC 到 DOCX 转换成功')
                return docx_file
            except Exception as e:
                # 捕获并记录详细错误
                import traceback
                error_details = traceback.format_exc()
                print(f"DOC转换为DOCX详细错误:\n{error_details}")
                
                # 尝试使用备用方法
                return self.handle_doc_file_alternative(doc_file)
            finally:
                # 确保关闭文档和应用，防止进程残留
                try:
                    if doc:
                        doc.Close(SaveChanges=False)
                except:
                    pass
                try:
                    if word:
                        word.Quit()
                except:
                    pass
                # 释放COM资源
                pythoncom.CoUninitialize()
        
        except Exception as e:
            print(f'DOC转换为DOCX失败: {str(e)}')
            return doc_file  # 如果转换失败，返回原始文件

    def handle_doc_file_alternative(self, doc_file):
        """DOC转DOCX的备用方法，使用python-docx"""
        try:
            self._update_ui(status='尝试备用DOC转换方法...')
            
            # 使用python-docx库
            try:
                import docx
                # 生成输出路径
                docx_file = os.path.splitext(doc_file)[0] + '.docx'
                
                # 读取原始文件内容（使用antiword或其他工具）
                try:
                    # 尝试使用antiword提取文本内容
                    content = self.extract_text_from_doc(doc_file)
                    
                    # 创建新文档
                    doc = docx.Document()
                    # 添加文本
                    doc.add_paragraph(content)
                    # 保存文档
                    doc.save(docx_file)
                    
                    self._update_ui(status='使用备用方法转换DOC成功')
                    return docx_file
                except:
                    # 如果提取内容失败，则直接使用原始DOC文件
                    self._update_ui(status='无法使用备用方法，将尝试直接转换DOC')
                    return doc_file
            except ImportError:
                # 如果python-docx不可用
                self._update_ui(status='备用转换库不可用，将尝试直接转换DOC')
                return doc_file
                
        except Exception as e:
            print(f"备用DOC转换失败: {str(e)}")
            return doc_file  # 返回原始文件

    def extract_text_from_doc(self, doc_file):
        """从DOC文件中提取文本，可能需要安装额外工具"""
        # 方法1：尝试使用subprocess调用antiword（如果安装）
        try:
            import subprocess
            kwargs = {"capture_output": True, "text": True}
            if os.name == "nt" and hasattr(subprocess, "CREATE_NO_WINDOW"):
                kwargs["creationflags"] = subprocess.CREATE_NO_WINDOW
            result = subprocess.run(["antiword", doc_file], timeout=20, **kwargs)
            if result.returncode == 0:
                return result.stdout
        except:
            pass
        
        # 方法2：尝试使用textract（如果安装）
        try:
            import textract
            return textract.process(doc_file).decode('utf-8')
        except:
            pass
        
        # 如果以上方法都失败，返回一个简单的错误消息
        return "无法提取文档内容，转换可能不完整。"

    def _parse_mhtml_to_html(self, input_file, images_dir, embed_images=True, mode='layout'):
        """解析MHTML文件，提取HTML、CSS并处理内嵌图片，支持图片内嵌（base64编码）"""
        import base64
        import email
        from email import policy
        from bs4 import BeautifulSoup
        import urllib.parse
        
        try:
            with open(input_file, 'rb') as f:
                msg = email.message_from_binary_file(f, policy=policy.default)
                
            html_content = ""
            css_contents = []
            images = {}
            images_by_location = {}
            images_by_basename = {}
            
            # 遍历MHTML各部分
            for part in msg.walk():
                content_type = part.get_content_type()
                
                if content_type == 'text/html':
                    if html_content:
                        continue  # 只保留第一个（通常是主文档）的 HTML 内容
                    html_bytes = part.get_payload(decode=True)
                    if html_bytes is None:
                        html_content = part.get_content()
                    else:
                        charset = part.get_content_charset()
                        if not charset:
                            head = html_bytes[:4096].decode('latin1', errors='ignore')
                            m = re.search(r'charset\s*=\s*["\']?([A-Za-z0-9_\-]+)', head, flags=re.I)
                            if m:
                                charset = m.group(1)
                        tried = []
                        for enc in [charset, 'utf-8', 'utf-8-sig', 'gb18030', 'gbk', 'big5', 'shift_jis', 'euc-kr', 'iso-8859-1']:
                            if not enc or enc in tried:
                                continue
                            tried.append(enc)
                            try:
                                html_content = html_bytes.decode(enc, errors='replace')
                                break
                            except Exception:
                                continue
                        if not html_content:
                            html_content = html_bytes.decode('utf-8', errors='replace')
                
                elif content_type == 'text/css' and mode == 'layout':
                    # 提取 CSS 样式用于版式保持
                    css_bytes = part.get_payload(decode=True)
                    if css_bytes:
                        try:
                            css_contents.append(css_bytes.decode('utf-8', errors='replace'))
                        except:
                            pass
                            
                elif content_type.startswith('image/'):
                    cid = part.get('Content-ID')
                    content_location = part.get('Content-Location')
                    
                    if cid or content_location:
                        if cid:
                            cid = cid.strip('<>')
                        else:
                            import hashlib
                            cid = hashlib.md5(content_location.encode('utf-8')).hexdigest()[:8]
                            
                        # 尝试获取文件名，如果没有则生成一个
                        filename = part.get_filename()
                        if not filename:
                            if content_location:
                                parsed_url = urllib.parse.urlparse(content_location)
                                filename = os.path.basename(parsed_url.path)
                            if not filename or '.' not in filename:
                                ext = content_type.split('/')[-1]
                                filename = f"image_{cid}.{ext}"
                                
                        # 清理文件名中可能导致路径问题的字符
                        filename = "".join(c for c in filename if c.isalnum() or c in (' ', '.', '_', '-'))
                        if not filename:
                            filename = f"image_{cid}.bin"
                        filepath = os.path.join(images_dir, filename)
                        
                        # 获取图片数据
                        img_data = part.get_payload(decode=True)
                        if not img_data:
                            continue

                        if embed_images:
                            # 直接使用base64编码
                            b64_data = base64.b64encode(img_data).decode('ascii')
                            data_uri = f"data:{content_type};base64,{b64_data}"
                            images[cid] = data_uri
                            if content_location:
                                loc = urllib.parse.unquote(content_location)
                                images_by_location[loc] = data_uri
                                images_by_location[content_location] = data_uri
                                images_by_basename[os.path.basename(loc)] = data_uri
                            images_by_basename[os.path.basename(filename)] = data_uri
                        else:
                            # 保存图片到文件
                            with open(filepath, 'wb') as img_f:
                                img_f.write(img_data)
                                
                            images[cid] = filename
                            if content_location:
                                loc = urllib.parse.unquote(content_location)
                                images_by_location[loc] = filename
                                images_by_location[content_location] = filename
                                images_by_basename[os.path.basename(loc)] = filename
                            images_by_basename[os.path.basename(filename)] = filename
                        
            # 如果没有提取到HTML，返回失败
            if not html_content:
                return None
                
            # 处理 HTML
            soup = BeautifulSoup(html_content, 'html.parser')
            
            # 注入 CSS
            head_tag = soup.find('head')
            if not head_tag:
                head_tag = soup.new_tag('head')
                if soup.html:
                    soup.html.insert(0, head_tag)
                else:
                    soup.insert(0, head_tag)

            # 确保字符集
            meta = head_tag.find('meta', attrs={'charset': True})
            if not meta:
                meta = soup.new_tag('meta', charset='utf-8')
                head_tag.insert(0, meta)

            # 注入提取到的 CSS
            if css_contents and mode == 'layout':
                style_tag = soup.new_tag('style', type='text/css')
                style_tag.string = "\n".join(css_contents)
                head_tag.append(style_tag)
            
            # 如果是文本模式，移除脚本和样式标签（除了我们刚刚注入的）
            if mode == 'text':
                for script_or_style in soup(["script", "style"]):
                    script_or_style.decompose()

            # 替换图片引用
            for img in soup.find_all('img'):
                src = img.get('src', '')
                if src.startswith('cid:'):
                    cid = src[4:]
                    if cid in images:
                        img['src'] = images[cid] if embed_images else f"images/{images[cid]}"
                else:
                    import html as html_lib
                    src_unescaped = html_lib.unescape(src)
                    src_u = urllib.parse.unquote(src)
                    src_u_unescaped = urllib.parse.unquote(src_unescaped)
                    
                    found = False
                    for key in [src, src_unescaped, src_u, src_u_unescaped]:
                        if key in images_by_location:
                            img['src'] = images_by_location[key] if embed_images else f"images/{images_by_location[key]}"
                            found = True
                            break
                    
                    if not found:
                        base = os.path.basename(src_u)
                        if base in images_by_basename:
                            img['src'] = images_by_basename[base] if embed_images else f"images/{images_by_basename[base]}"
                        
            # 将处理后的HTML保存为临时文件
            temp_html_path = tempfile.NamedTemporaryFile(prefix="anytomd_", suffix=".html", delete=False).name
            with open(temp_html_path, 'w', encoding='utf-8') as f:
                f.write(str(soup))
                
            # 添加到临时文件列表以便后续清理
            if not hasattr(self, 'temp_files'):
                self.temp_files = []
            self.temp_files.append(temp_html_path)
                
            return temp_html_path
            
        except Exception as e:
            print(f"解析MHTML失败: {str(e)}")
            return None

    def _set_widget_enabled(self, widget, enabled):
        state = "normal" if enabled else "disabled"
        try:
            widget.configure(state=state)
            return
        except Exception:
            pass
        try:
            widget.config(state=tk.NORMAL if enabled else tk.DISABLED)
        except Exception:
            pass

    def _set_progress_percent(self, widget, progress):
        if widget is None:
            return
        try:
            if hasattr(widget, "set"):
                v = progress / 100.0
                if v < 0:
                    v = 0
                if v > 1:
                    v = 1
                widget.set(v)
            else:
                widget['value'] = progress
        except Exception:
            try:
                widget['value'] = progress
            except Exception:
                pass

    def _update_ui(self, status=None, progress=None, is_batch=False):
        """线程安全地更新 UI"""
        def task():
            if status is not None:
                if is_batch:
                    self.batch_status_var.set(status)
                    if self.use_ctk and hasattr(self, 'batch_status_label'):
                        self.batch_status_label.configure(text=status)
                self.status_var.set(status)
                if self.use_ctk and hasattr(self, 'status_label'):
                    self.status_label.configure(text=status)
            if progress is not None:
                if is_batch:
                    self._set_progress_percent(self.batch_progress, progress)
                else:
                    self._set_progress_percent(self.progress, progress)
        self.root.after(0, task)

    def _set_document_output_mode_label(self, mode):
        if hasattr(self, 'document_output_mode_var'):
            self.document_output_mode_var.set(DOCUMENT_OUTPUT_MODE_LABELS.get(mode, DOCUMENT_OUTPUT_MODE_LABELS['smart']))

    def _map_legacy_output_mode(self, settings=None):
        settings = settings or {}
        pdf_mode = settings.get('pdf_mode')
        mhtml_mode = settings.get('mhtml_mode')
        if pdf_mode == 'layout':
            return 'layout'
        if pdf_mode == 'auto_scanned':
            return 'text_ocr'
        if pdf_mode == 'normal':
            return 'smart'
        if mhtml_mode == 'image':
            return 'image'
        if mhtml_mode == 'text':
            return 'text'
        if settings.get('pdf_scanned_no_ocr') is True:
            return 'text_ocr'
        return 'smart'

    def _get_document_output_mode(self):
        v = self.document_output_mode_var.get().strip() if hasattr(self, 'document_output_mode_var') else ''
        if v in DOCUMENT_OUTPUT_MODE_LABELS:
            return v
        if v in DOCUMENT_OUTPUT_MODE_KEYS:
            return DOCUMENT_OUTPUT_MODE_KEYS[v]
        if v.startswith('保持版式'):
            return 'layout'
        if v.startswith('保留图像'):
            return 'image'
        if v.startswith('保留文字（仅'):
            return 'text'
        if v.startswith('保留文字（含'):
            return 'text_ocr'
        if v.startswith('智能转换'):
            return 'smart'
        return 'smart'

    def _get_pdf_mode(self):
        mode = self._get_document_output_mode()
        if mode == 'layout':
            return 'layout'
        if mode == 'text_ocr':
            return 'auto_scanned'
        return 'normal'

    def _get_pdf_engine_strategy(self):
        """获取 PDF 引擎策略：自动 / 图文混排优先 / pdf-inspector优先"""
        v = self.pdf_engine_strategy_var.get().strip() if hasattr(self, 'pdf_engine_strategy_var') else PDF_ENGINE_STRATEGY_DEFAULT
        if v.startswith('图文混排'):
            return 'layout_first'
        if v.startswith('pdf-inspector'):
            return 'inspector_first'
        return 'auto'
    
    def _get_mhtml_mode(self):
        mode = self._get_document_output_mode()
        if mode == 'image':
            return 'image'
        if mode in ('text', 'text_ocr'):
            return 'text'
        return 'layout'

    def _get_image_output_policy(self):  # 2026-03-29 13:41:12 修改：图片输出策略解析（仅原位/仅补缺/全量附录）
        v = self.image_output_policy_var.get().strip() if hasattr(self, 'image_output_policy_var') else ''
        if v in ('inplace', 'supplement', 'appendix'):
            return v
        if v.startswith('仅原位'):
            return 'inplace'
        if v.startswith('仅补缺'):
            return 'supplement'
        if v.startswith('全量附录'):
            return 'appendix'
        return 'inplace'

    def _get_batch_overwrite_policy(self):  # 2026-03-29 13:41:12 修改：批量同名输出策略解析（跳过/覆盖/自动重命名）
        v = self.batch_overwrite_policy_var.get().strip() if hasattr(self, 'batch_overwrite_policy_var') else ''
        if v in ('skip', 'overwrite', 'rename'):
            return v
        if v.startswith('跳过'):
            return 'skip'
        if v.startswith('覆盖'):
            return 'overwrite'
        if v.startswith('自动重命名'):
            return 'rename'
        return 'skip'

    def _resolve_output_path_for_batch(self, output_folder, base_name):  # 2026-03-29 13:41:12 修改：批量输出路径按策略分配，避免误覆盖
        output_path = os.path.join(output_folder, base_name)
        if not os.path.exists(output_path):
            return output_path, None
        policy = self._get_batch_overwrite_policy()
        if policy == 'overwrite':
            return output_path, None
        if policy == 'skip':
            return None, '同名输出已存在，已跳过'
        stem, ext = os.path.splitext(base_name)
        for i in range(1, 10000):
            candidate = os.path.join(output_folder, f"{stem} ({i}){ext}")
            if not os.path.exists(candidate):
                return candidate, None
        return None, '自动重命名失败'

    def _md_has_any_image_link(self, md_path):
        try:
            with open(md_path, 'r', encoding='utf-8', errors='ignore') as f:
                content = f.read()
            return bool(re.search(r'!\[[^\]]*\]\([^)]+\)', content))
        except Exception:
            return False

    def _get_existing_image_hashes(self, md_path):  # 2026-03-29 14:30:53 修改：为“仅补缺”提供hash去重依据，避免文末重复图片
        import base64
        import hashlib
        import mimetypes
        import re
        import urllib.parse

        hashes = set()
        try:
            with open(md_path, 'r', encoding='utf-8', errors='ignore') as f:
                content = f.read()
        except Exception:
            return hashes

        md_dir = os.path.dirname(md_path)
        for m in re.finditer(r'!\[[^\]]*\]\(([^)]+)\)', content):
            target = (m.group(1) or '').strip()
            if not target:
                continue
            if target.startswith('<') and target.endswith('>'):
                target = target[1:-1].strip()
            target = urllib.parse.unquote(target)
            target = target.split('#', 1)[0].split('?', 1)[0]

            if target.startswith('data:'):
                dm = re.match(r'^data:([^;]+);base64,(.+)$', target, flags=re.IGNORECASE | re.DOTALL)
                if not dm:
                    continue
                try:
                    data = base64.b64decode(dm.group(2).strip(), validate=False)
                    hashes.add(hashlib.sha256(data).hexdigest())
                except Exception:
                    continue
                continue

            if re.match(r'^[a-zA-Z]+://', target):
                continue

            if os.path.isabs(target):
                img_path = os.path.abspath(target)
            else:
                img_path = os.path.abspath(os.path.join(md_dir, target))

            if not os.path.exists(img_path):
                continue

            mime_type = mimetypes.guess_type(img_path)[0] or ''
            if not mime_type.startswith('image/'):
                continue
            try:
                with open(img_path, 'rb') as img_f:
                    data = img_f.read()
                hashes.add(hashlib.sha256(data).hexdigest())
            except Exception:
                continue

        return hashes

    def _convert_with_markitdown_api(self, input_file, output_path):  # 2026-03-29 13:41:12 修改：markitdown命令缺失时用API回退
        module = _get_markitdown()
        if module is None:
            return False
        try:
            converter = module.MarkItDown()
            result = converter.convert(input_file)
            text = getattr(result, 'text_content', None)
            if not text:
                return False
            with open(output_path, 'w', encoding='utf-8') as f:
                f.write(text)
            return True
        except Exception:
            return False

    def _apply_office_image_output_policy(self, original_file_ext, input_file, output_path, images_dir):  # 2026-03-29 13:41:12 修改：Office文末附录图片按策略控制
        policy = self._get_image_output_policy()
        if policy == 'inplace':  # 2026-03-29 14:30:53 修改：仅原位不做文末追加
            return
        try:
            if original_file_ext in ['.docx', '.doc']:
                self.extract_images_from_word(input_file, output_path, images_dir)
            elif original_file_ext == '.pptx':
                self.extract_images_from_pptx(input_file, output_path, images_dir)
        except Exception:
            pass

    # ==================================================================
    # PDF 处理模块【逻辑锁定】（2026-09-02 补全；冻结声明此前仅见诸
    # README / CONTRIBUTING / docs/审核报告与注意事项，现补代码级注释）
    # 未经专门批准，禁止修改下述 PDF 核心解析/渲染成员及其实现。
    # ------------------------------------------------------------------
    # 锁定成员清单（按代码分布分簇，不标注绝对行号以免代码变动后误导）：
    #  A 簇（图片策略 + 文本抽取/目录页/正文重排/代码保护：从
    #    _apply_pdf_image_output_policy 起，至 _convert_html_like_to_pdf 之前的一段连续区）：
    #    _apply_pdf_image_output_policy、_find_pdftotext_exe、
    #    _extract_pdf_page_texts_by_pdftotext / _by_pdfplumber / _by_pypdf2、
    #    _extract_pdf_page_texts、_is_pdf_page_number_line、_is_pdf_dot_leader_line、
    #    _normalize_pdf_page_text、_is_pdf_toc_heading_line、_is_pdf_toc_running_header_line、
    #    _parse_pdf_toc_entry_segment、_extract_pdf_toc_entries_from_line、
    #    _is_pdf_toc_entry_line、_is_probable_pdf_toc_page、_format_pdf_toc_page_text、
    #    _is_pdf_list_like_line、_is_pdf_heading_like_line、_is_pdf_reflow_excluded_line、
    #    _join_pdf_reflow_lines、_is_pdf_output_heading_line、_is_pdf_output_line、
    #    _protect_pdf_output_blocks、_reflow_pdf_paragraph_lines、_merge_pdf_toc_lines、
    #    _is_pdf_example_header_line、_is_pdf_code_line、_protect_pdf_code_blocks、
    #    _escape_pdf_markdown_syntax、_postprocess_pdf_page_text
    #  B 簇（图块渲染与图文混排/OCR 分段）：_render_pdf_image_block、
    #    _extract_pdf_text_sections、_convert_pdf_to_mixed_markdown、
    #    _extract_pdf_text_sections_with_ocr
    #  C 簇（扫描判定与整页转图）：_is_scanned_pdf、_convert_pdf_pages_to_markdown
    #  D 簇（pdf-inspector 引擎转换）：_convert_pdf_with_pdfinspector
    # 另外：_collect_pdf_image_assets（图片收集）同属锁定范围。
    # ------------------------------------------------------------------
    # 不属锁定（外围可改，改后不影响冻结目标）：
    #   参数读取小工具 _get_pdf_mode / _get_pdf_engine_strategy、
    #   “转为 PDF”的辅助 _convert_html_like_to_pdf（HTML→PDF 无头打印）与
    #   _convert_office_document_to_pdf（Office→PDF）、UI/功能检查/引擎检测、
    #   单文件与批量的“调度分支”（仅调用上述成员，不改其内部）。
    # ------------------------------------------------------------------
    # 冻结原因与思路：PDF 图文混排/目录/扫描件等核心效果历经多轮调优
    # （注释日期多集中于 2026-03~04），属已稳定基线；对其改动须逐项
    # 真实样本验证与单独回归（见开发工作记录与《审核报告与注意事项》）。
    # ==================================================================

    def _apply_pdf_image_output_policy(self, input_file, output_path, images_dir):  # 2026-03-29 13:41:12 修改：PDF文末附录图片按策略控制
        policy = self._get_image_output_policy()
        if policy == 'inplace':  # 2026-03-29 14:30:53 修改：仅原位不做文末追加
            return
        self.extract_images_from_pdf(input_file, output_path, images_dir)  # 2026-03-29 14:30:53 修改：由提取函数内部按“仅补缺/全量附录”与hash去重控制追加

    def _register_temp_file(self, file_path):
        if not file_path:
            return
        if not hasattr(self, 'temp_files'):
            self.temp_files = []
        if file_path not in self.temp_files:
            self.temp_files.append(file_path)

    def _write_markdown_content(self, output_path, content):
        with open(output_path, 'w', encoding='utf-8') as f:
            f.write(content if not content or content.endswith('\n') else content + '\n')

    def _write_markdown_sections(self, output_path, sections):
        parts = []
        for title, text in sections:
            text = (text or '').strip()
            if not text:
                continue
            if title:
                parts.append(f"## {title}\n\n{text}")
            else:
                parts.append(text)
        self._write_markdown_content(output_path, "\n\n".join(parts))
        return True

    def _normalize_image_extension(self, ext, default='.png'):
        ext = (ext or '').strip().lower()
        if not ext:
            return default
        if not ext.startswith('.'):
            ext = '.' + ext
        if len(ext) > 10:
            return default
        return ext

    def _normalize_compare_text(self, text):
        return re.sub(r'\s+', '', text or '')

    def _filter_ocr_sections(self, base_sections, ocr_sections):
        merged = list(base_sections)
        seen = {self._normalize_compare_text(text) for _, text in base_sections if text}
        base_joined = ''.join(seen)
        for title, text in ocr_sections:
            norm = self._normalize_compare_text(text)
            if not norm:
                continue
            if norm in seen:
                continue
            if norm and norm in base_joined:
                continue
            seen.add(norm)
            merged.append((title, text))
        return merged

    def _find_pdftotext_exe(self):
        candidates = []
        poppler_bin = self.find_poppler_bin()
        if poppler_bin:
            candidates.append(os.path.join(poppler_bin, 'pdftotext.exe'))
            candidates.append(os.path.join(poppler_bin, 'pdftotext'))
        which_candidate = shutil.which('pdftotext')
        if which_candidate:
            candidates.append(which_candidate)
        seen = set()
        for candidate in candidates:
            if not candidate:
                continue
            candidate_norm = os.path.normcase(os.path.abspath(candidate))
            if candidate_norm in seen:
                continue
            seen.add(candidate_norm)
            if os.path.exists(candidate):
                return candidate
        return None

    def _extract_pdf_page_texts_by_pdftotext(self, input_file):
        pdftotext_exe = self._find_pdftotext_exe()
        if not pdftotext_exe:
            return None
        temp_txt = tempfile.NamedTemporaryFile(prefix='anytomd_pdftotext_', suffix='.txt', delete=False).name
        self._register_temp_file(temp_txt)
        cmd = [pdftotext_exe, '-layout', '-enc', 'UTF-8', input_file, temp_txt]
        kwargs = {
            'capture_output': True,
            'text': True,
            'timeout': 120
        }
        if os.name == 'nt' and hasattr(subprocess, 'CREATE_NO_WINDOW'):
            kwargs['creationflags'] = subprocess.CREATE_NO_WINDOW
        try:
            process = subprocess.run(cmd, **kwargs)
        except Exception:
            return None
        if process.returncode != 0 or not os.path.exists(temp_txt):
            return None
        try:
            with open(temp_txt, 'r', encoding='utf-8', errors='ignore') as f:
                content = f.read()
        except Exception:
            return None
        content = content.replace('\r\n', '\n').replace('\r', '\n')
        pages = content.split('\f')
        while pages and not pages[-1].strip():
            pages.pop()
        return [page.strip('\n') for page in pages]

    def _extract_pdf_page_texts_by_pdfplumber(self, input_file):
        try:
            import pdfplumber
        except Exception:
            return None
        try:
            page_texts = []
            with pdfplumber.open(input_file) as pdf:
                for page in pdf.pages:
                    text = (page.extract_text() or '').replace('\r\n', '\n').replace('\r', '\n').strip('\n')
                    page_texts.append(text)
            return page_texts
        except Exception:
            return None

    def _extract_pdf_page_texts_by_pypdf2(self, input_file):
        try:
            import PyPDF2
        except Exception:
            return None
        try:
            page_texts = []
            with open(input_file, 'rb') as file:
                reader = PyPDF2.PdfReader(file)
                for page in reader.pages:
                    text = (page.extract_text() or '').replace('\r\n', '\n').replace('\r', '\n').strip('\n')
                    page_texts.append(text)
            return page_texts
        except Exception:
            return None

    def _extract_pdf_page_texts(self, input_file):
        for extractor in (
            self._extract_pdf_page_texts_by_pdftotext,
            self._extract_pdf_page_texts_by_pdfplumber,
            self._extract_pdf_page_texts_by_pypdf2
        ):
            page_texts = extractor(input_file)
            if page_texts is not None:
                return page_texts
        return []

    def _is_pdf_page_number_line(self, line):
        stripped = (line or '').strip()
        if not stripped:
            return False
        return bool(re.fullmatch(r'(?:\d+|[ivxlcdm]+)(?:\s*[-–]\s*(?:\d+|[ivxlcdm]+))?', stripped, flags=re.IGNORECASE))

    def _is_pdf_dot_leader_line(self, line):
        stripped = (line or '').strip()
        if not stripped:
            return False
        return bool(re.fullmatch(r'[.·•∙\s]{4,}', stripped))

    def _normalize_pdf_page_text(self, text):
        text = (text or '').replace('\r\n', '\n').replace('\r', '\n').replace('\xa0', ' ')
        lines = [line.rstrip() for line in text.split('\n')]
        normalized_lines = []
        blank_count = 0
        for line in lines:
            if line.strip():
                blank_count = 0
                normalized_lines.append(line)
                continue
            blank_count += 1
            if blank_count <= 1:
                normalized_lines.append('')
        return '\n'.join(normalized_lines).strip()

    def _is_pdf_toc_heading_line(self, line):  # 2026-04-03 19:10:00 修改：新增PDF目录页识别，避免目录被误判为代码
        stripped = (line or '').strip()
        if not stripped:
            return False
        return bool(re.fullmatch(r'(?:table\s+of\s+contents|contents)', stripped, flags=re.IGNORECASE))

    def _is_pdf_toc_running_header_line(self, line):  # 2026-04-03 19:10:00 修改：新增PDF目录页眉识别，减少目录正文混入页眉页码
        stripped = re.sub(r'\s+', ' ', (line or '').strip())
        if not stripped:
            return False
        return bool(re.fullmatch(r'(?:contents\s+[ivxlcdm]+|[ivxlcdm]+\s+contents|contents)', stripped, flags=re.IGNORECASE))

    def _parse_pdf_toc_entry_segment(self, segment):  # 2026-04-03 19:10:00 修改：新增PDF目录项解析，支持目录页专用文本输出
        raw_segment = (segment or '').rstrip()
        stripped = raw_segment.strip()
        if not stripped or self._is_pdf_toc_heading_line(stripped) or self._is_pdf_toc_running_header_line(stripped):
            return None
        match = re.fullmatch(
            r'(?P<title>.+?)(?:\s*(?:[.·•∙]\s*){2,}|\s{2,})(?P<page>(?:\d+|[ivxlcdm]+)(?:\s*[-–]\s*(?:\d+|[ivxlcdm]+))?)',
            stripped,
            flags=re.IGNORECASE
        )
        if not match:
            return None
        title = re.sub(r'\s+', ' ', match.group('title')).strip(' .·•∙')
        page = match.group('page').strip()
        if not title or not page:
            return None
        indent = len(raw_segment) - len(raw_segment.lstrip(' '))
        return indent, title, page

    def _extract_pdf_toc_entries_from_line(self, line):  # 2026-04-03 19:10:00 修改：新增PDF目录行拆分，处理双栏目录与超宽目录行
        raw_line = (line or '').rstrip()
        if not raw_line.strip():
            return []
        direct_entry = self._parse_pdf_toc_entry_segment(raw_line)
        if direct_entry:
            return [direct_entry]
        if not re.search(r'\s{8,}', raw_line):
            return []
        entries = []
        for segment in re.split(r'\s{8,}', raw_line):
            entry = self._parse_pdf_toc_entry_segment(segment)
            if entry:
                entries.append(entry)
        return entries

    def _is_pdf_toc_entry_line(self, line):  # 2026-04-03 19:10:00 修改：新增PDF目录项判断，避免目录子项进入代码块
        return bool(self._extract_pdf_toc_entries_from_line(line))

    def _is_probable_pdf_toc_page(self, text):  # 2026-04-03 19:10:00 修改：新增PDF目录页探测，仅对目录页应用专用输出逻辑
        lines = [line for line in (text or '').split('\n') if line.strip()]
        if not lines:
            return False
        heading_hits = sum(1 for line in lines if self._is_pdf_toc_heading_line(line))
        entry_hits = sum(len(self._extract_pdf_toc_entries_from_line(line)) for line in lines)
        leader_hits = sum(1 for line in lines if re.search(r'(?:[.·•∙]\s*){4,}', line))
        return bool((heading_hits >= 1 and entry_hits >= 3) or entry_hits >= 8 or (heading_hits >= 1 and leader_hits >= 4))

    def _format_pdf_toc_page_text(self, text):  # 2026-04-03 19:10:00 修改：新增PDF目录页专用格式化，避免横向滚动并保留文字输出
        lines = (text or '').split('\n')
        formatted_lines = []
        seen_heading = False
        entry_count = 0
        for line in lines:
            stripped = (line or '').strip()
            if not stripped:
                if formatted_lines and formatted_lines[-1]:
                    formatted_lines.append('')
                continue
            if self._is_pdf_toc_heading_line(stripped):
                if not seen_heading:
                    formatted_lines.append(stripped.title())
                    formatted_lines.append('')
                    seen_heading = True
                continue
            if self._is_pdf_toc_running_header_line(stripped) or self._is_pdf_page_number_line(stripped) or self._is_pdf_dot_leader_line(stripped):
                continue
            entries = self._extract_pdf_toc_entries_from_line(line)
            if entries:
                for indent, title, page in entries:
                    level = max(0, min(indent // 4, 4))
                    formatted_lines.append(f"{'  ' * level}- {title} — {page}")
                    entry_count += 1
                continue
            formatted_lines.append(re.sub(r'\s+', ' ', stripped))
        while formatted_lines and not formatted_lines[-1]:
            formatted_lines.pop()
        if entry_count < 3:
            return text
        return '\n'.join(formatted_lines)

    def _is_pdf_list_like_line(self, line):  # 2026-04-03 20:05:00 修改：新增PDF列表识别，避免正文重组时误并列表项
        stripped = (line or '').strip()
        if not stripped:
            return False
        return bool(re.match(r'^(?:[-*+•]\s+|\d+[.)]\s+|[A-Za-z][.)]\s+)', stripped))

    def _is_pdf_heading_like_line(self, line):  # 2026-04-03 20:05:00 修改：新增PDF标题识别，避免正文重组时误并标题行
        stripped = re.sub(r'\s+', ' ', (line or '').strip())
        if not stripped:
            return False
        if len(stripped) <= 80 and re.fullmatch(r'[A-Z0-9\s:,\-–\'"&]+', stripped):
            return True
        if re.match(r'^(?:Exercise|Chapter|Module|Part|Listing)\b', stripped, flags=re.IGNORECASE):
            return True
        return False

    def _is_pdf_reflow_excluded_line(self, line):  # 2026-04-03 20:05:00 修改：新增PDF正文重组排除规则，降低误合并与误判代码风险
        stripped = (line or '').strip()
        if not stripped:
            return True
        if self._is_pdf_toc_entry_line(line) or self._is_pdf_toc_heading_line(line) or self._is_pdf_toc_running_header_line(line):
            return True
        if self._is_pdf_output_heading_line(line) or self._is_pdf_output_line(line):  # 2026-04-03 20:35:00 修改：输出标题与逐行输出禁止参与正文重组，避免被提前并成一段
            return True
        if self._is_pdf_code_line(line) or self._is_pdf_example_header_line(line):
            return True
        if self._is_pdf_page_number_line(stripped) or self._is_pdf_dot_leader_line(stripped):
            return True
        if self._is_pdf_list_like_line(line) or self._is_pdf_heading_like_line(line):
            return True
        return False

    def _join_pdf_reflow_lines(self, current_line, next_line):  # 2026-04-03 20:05:00 修改：新增PDF断行拼接，修复单词尾部被拆到下一行
        current = current_line.rstrip()
        next_text = next_line.lstrip()
        if not current:
            return next_text
        if not next_text:
            return current
        current_word_match = re.search(r'([A-Za-z]+)$', current)
        next_word_match = re.match(r'([a-z]+)', next_text)
        current_word = current_word_match.group(1) if current_word_match else ''
        next_word = next_word_match.group(1) if next_word_match else ''
        if current.endswith('-'):
            return current[:-1] + next_text
        if current_word and next_word and len(next_word) <= 2 and current_word.lower() not in {
            'a', 'an', 'and', 'are', 'as', 'at', 'be', 'by', 'do', 'go', 'he', 'i', 'if', 'in', 'is', 'it', 'me',
            'my', 'no', 'of', 'on', 'or', 'so', 'to', 'up', 'us', 'we', 'you'
        }:
            return current + next_text
        return f"{current} {next_text}"

    def _is_pdf_output_heading_line(self, line):  # 2026-04-03 20:20:00 修改：新增PDF输出标题识别，保护控制台输出块的逐行显示
        stripped = re.sub(r'\s+', ' ', (line or '').strip())
        if not stripped:
            return False
        return stripped.lower() in {
            'what you should see',
            'output',
            'expected output'
        }

    def _is_pdf_output_line(self, line):  # 2026-04-03 20:20:00 修改：新增PDF输出行识别，避免示例输出被Markdown合并成同一段
        stripped = (line or '').rstrip()
        if not stripped:
            return False
        return bool(re.match(r'^\s*\d+\s{2,}\S+', stripped))

    def _protect_pdf_output_blocks(self, text):  # 2026-04-03 20:20:00 修改：新增PDF输出块保护，保持“所见输出”逐行排版
        lines = (text or '').split('\n')
        result = []
        index = 0
        while index < len(lines):
            line = lines[index]
            result.append(line)
            if not self._is_pdf_output_heading_line(line):
                index += 1
                continue
            index += 1
            while index < len(lines) and not lines[index].strip():
                result.append(lines[index])
                index += 1
            block_lines = []
            while index < len(lines):
                candidate = lines[index]
                if not candidate.strip():
                    break
                if self._is_pdf_output_line(candidate):
                    block_lines.append(candidate.rstrip())
                    index += 1
                    continue
                break
            if block_lines:
                result.append('```text')
                result.extend(block_lines)
                result.append('```')
                continue
        return '\n'.join(result)

    def _reflow_pdf_paragraph_lines(self, text):  # 2026-04-03 20:05:00 修改：新增PDF正文段落重组，减少一行被拆成两行的问题
        lines = (text or '').split('\n')
        reflowed_lines = []
        index = 0
        while index < len(lines):
            line = lines[index]
            stripped = line.strip()
            if not stripped:
                if reflowed_lines and reflowed_lines[-1]:
                    reflowed_lines.append('')
                index += 1
                continue
            if self._is_pdf_reflow_excluded_line(line):
                reflowed_lines.append(stripped if not self._is_pdf_code_line(line) else line.rstrip())
                index += 1
                continue
            current = stripped
            index += 1
            while index < len(lines):
                candidate = lines[index]
                candidate_stripped = candidate.strip()
                if not candidate_stripped:
                    break
                if self._is_pdf_reflow_excluded_line(candidate):
                    break
                if re.search(r'[.!?。！？:：]$', current):
                    break
                current = self._join_pdf_reflow_lines(current, candidate_stripped)
                index += 1
            reflowed_lines.append(current)
        return '\n'.join(reflowed_lines)

    def _merge_pdf_toc_lines(self, text):
        lines = text.split('\n')
        merged_lines = []
        index = 0
        while index < len(lines):
            line = lines[index]
            next_line = lines[index + 1] if index + 1 < len(lines) else None
            next_next_line = lines[index + 2] if index + 2 < len(lines) else None
            stripped = line.strip()
            next_stripped = (next_line or '').strip()
            next_next_stripped = (next_next_line or '').strip()
            if (
                stripped and not self._is_pdf_page_number_line(stripped) and not self._is_pdf_dot_leader_line(stripped)
                and self._is_pdf_dot_leader_line(next_stripped) and self._is_pdf_page_number_line(next_next_stripped)
            ):
                merged_lines.append(f"{line.rstrip()} {next_stripped} {next_next_stripped}")
                index += 3
                continue
            if (
                self._is_pdf_dot_leader_line(stripped) and self._is_pdf_page_number_line(next_stripped)
                and merged_lines and merged_lines[-1].strip()
            ):
                merged_lines[-1] = f"{merged_lines[-1].rstrip()} {stripped} {next_stripped}"
                index += 2
                continue
            if (
                stripped and next_stripped and self._is_pdf_page_number_line(next_stripped)
                and (
                    re.search(r'(?:[.·•∙]\s*){4,}$', stripped)
                    or (len(stripped) <= 120 and not re.search(r'[。！？!?：:；;]$', stripped))
                )
            ):
                merged_lines.append(f"{line.rstrip()} {next_stripped}")
                index += 2
                continue
            merged_lines.append(line)
            index += 1
        return '\n'.join(merged_lines)

    def _is_pdf_example_header_line(self, line):
        stripped = (line or '').strip().lower()
        if not stripped:
            return False
        return stripped in {
            'markdown',
            'html',
            'output',
            'result'
        } or bool(re.fullmatch(r'(?:markdown|html)\s{2,}(?:markdown|html)', stripped))

    def _is_pdf_code_line(self, line):
        stripped = (line or '').strip()
        if not stripped:
            return False
        if self._is_pdf_toc_entry_line(line):
            return False
        indent = len(line) - len(line.lstrip(' '))
        if '\t' in line:
            return True
        if indent >= 4 and re.search(r'(?:print\s*\(|return\b|import\b|from\b|class\b|def\b|for\b|while\b|if\b|elif\b|else:|try:|except\b|with\b|[=(){}\[\];<>])', stripped):
            return True
        patterns = [
            r'^#{1,6}\s+\S+',
            r'^(?:\d+\s+)?(?:from\s+\w+\s+import\s+.+|import\s+\w+.*|def\s+\w+\s*\(|class\s+\w+|for\s+.+:|while\s+.+:|if\s+.+:|elif\s+.+:|else:|try:|except\b.*:|with\s+.+:|return\b|print\s*\(|[A-Za-z_][\w\.]*\s*=)',
            r'^(?:\d+\s+)?["\'].*["\'](?:\s{2,}#.*)?$',
            r'^(?:\d+\s+)?</?[A-Za-z][^>]*>.*$',
            r'^[=-]{3,}$'
        ]
        if any(re.match(pattern, stripped) for pattern in patterns):
            return True
        if '#' in stripped and ('=' in stripped or '\\' in stripped or '"' in stripped or "'" in stripped):
            return True
        if re.search(r'</?[A-Za-z][^>]*>', stripped):
            return True
        return False

    def _protect_pdf_code_blocks(self, text):
        lines = text.split('\n')
        result = []
        index = 0
        while index < len(lines):
            line = lines[index]
            next_line = lines[index + 1] if index + 1 < len(lines) else ''
            if self._is_pdf_code_line(line) or (
                self._is_pdf_example_header_line(line)
                and (self._is_pdf_code_line(next_line) or self._is_pdf_example_header_line(next_line))
            ):
                block_lines = [line]
                index += 1
                while index < len(lines):
                    candidate = lines[index]
                    next_candidate = lines[index + 1] if index + 1 < len(lines) else ''
                    if not candidate.strip():
                        if self._is_pdf_code_line(next_candidate) or self._is_pdf_example_header_line(next_candidate):
                            block_lines.append(candidate)
                            index += 1
                            continue
                        break
                    if self._is_pdf_code_line(candidate) or self._is_pdf_example_header_line(candidate):
                        block_lines.append(candidate)
                        index += 1
                        continue
                    break
                result.append('```text')
                result.extend(block_lines)
                result.append('```')
                continue
            result.append(line)
            index += 1
        return '\n'.join(result)

    def _escape_pdf_markdown_syntax(self, text):
        escaped_lines = []
        in_code_block = False
        for line in text.split('\n'):
            stripped = line.lstrip()
            if stripped.startswith('```'):
                in_code_block = not in_code_block
                escaped_lines.append(line)
                continue
            if in_code_block or not stripped:
                escaped_lines.append(line)
                continue
            leading = line[:len(line) - len(stripped)]
            if re.match(r'^#{1,6}\s', stripped):
                escaped_lines.append(f"{leading}\\{stripped}")
                continue
            if re.match(r'^>\s', stripped):
                escaped_lines.append(f"{leading}\\{stripped}")
                continue
            if re.match(r'^---+$', stripped):
                escaped_lines.append(f"{leading}\\{stripped}")
                continue
            if re.match(r'^\d+\.\s+.+[.·]{4,}\s+(?:\d+|[ivxlcdm]+)$', stripped, flags=re.IGNORECASE):
                escaped_lines.append(f"{leading}{re.sub(r'^(\\d+)\\.', r'\\1\\.', stripped, count=1)}")
                continue
            if re.search(r'</?[A-Za-z][^>]*>', stripped):
                escaped_lines.append(f"{leading}{stripped.replace('&', '&amp;').replace('<', '&lt;').replace('>', '&gt;')}")
                continue
            escaped_lines.append(line)
        return '\n'.join(escaped_lines)

    def _postprocess_pdf_page_text(self, text):
        text = self._normalize_pdf_page_text(text)
        if not text:
            return ''
        text = self._merge_pdf_toc_lines(text)
        if self._is_probable_pdf_toc_page(text):  # 2026-04-03 19:10:00 修改：目录页走专用文本输出，避免被包成代码块
            text = self._format_pdf_toc_page_text(text)
            text = self._escape_pdf_markdown_syntax(text)
            return self._normalize_pdf_page_text(text)
        text = self._reflow_pdf_paragraph_lines(text)  # 2026-04-03 20:05:00 修改：先做正文重组，再做代码保护，降低断行与误入代码块问题
        text = self._protect_pdf_output_blocks(text)  # 2026-04-03 20:20:00 修改：保护编号输出块，避免“所见输出”逐行内容被合并渲染
        text = self._protect_pdf_code_blocks(text)
        text = self._escape_pdf_markdown_syntax(text)
        return self._normalize_pdf_page_text(text)

    def _find_headless_browser(self):
        candidates = []
        for name in ('msedge', 'chrome', 'chromium'):
            path = shutil.which(name)
            if path:
                candidates.append(path)
        for env_name in ('PROGRAMFILES', 'PROGRAMFILES(X86)', 'LOCALAPPDATA'):
            base = os.environ.get(env_name)
            if not base:
                continue
            candidates.extend([
                os.path.join(base, 'Microsoft', 'Edge', 'Application', 'msedge.exe'),
                os.path.join(base, 'Google', 'Chrome', 'Application', 'chrome.exe'),
                os.path.join(base, 'Chromium', 'Application', 'chromium.exe')
            ])
        for candidate in candidates:
            if candidate and os.path.exists(candidate):
                return candidate
        return None

    def _convert_html_like_to_pdf(self, input_file):
        browser = self._find_headless_browser()
        if not browser:
            return None
        temp_pdf = tempfile.NamedTemporaryFile(prefix='anytomd_layout_', suffix='.pdf', delete=False).name
        self._register_temp_file(temp_pdf)
        file_uri = Path(input_file).resolve().as_uri()
        kwargs = {
            'capture_output': True,
            'text': True,
            'timeout': 120
        }
        if os.name == 'nt' and hasattr(subprocess, 'CREATE_NO_WINDOW'):
            kwargs['creationflags'] = subprocess.CREATE_NO_WINDOW
        cmd = [
            browser,
            '--headless',
            '--disable-gpu',
            '--no-pdf-header-footer',
            f'--print-to-pdf={temp_pdf}',
            file_uri
        ]
        process = subprocess.run(cmd, **kwargs)
        if process.returncode == 0 and os.path.exists(temp_pdf) and os.path.getsize(temp_pdf) > 0:
            return temp_pdf
        return None

    def _find_office_executable(self, original_file_ext):
        names = []
        if original_file_ext in ['.doc', '.docx']:
            names = ['WINWORD.EXE', 'winword.exe']
        elif original_file_ext in ['.ppt', '.pptx']:  # 2026-09-02 修改：旧版 .ppt 与 .pptx 均由 PowerPoint 处理（保持版式）
            names = ['POWERPNT.EXE', 'powerpnt.exe']
        if not names:
            return None
        candidates = []
        for env_name in ('PROGRAMFILES', 'PROGRAMFILES(X86)', 'LOCALAPPDATA'):
            base = os.environ.get(env_name)
            if not base:
                continue
            candidates.extend([
                os.path.join(base, 'Microsoft Office'),
                os.path.join(base, 'Microsoft Office', 'root', 'Office16'),
                os.path.join(base, 'Microsoft Office', 'root', 'Office15'),
                os.path.join(base, 'Microsoft Office', 'Office16'),
                os.path.join(base, 'Microsoft Office', 'Office15')
            ])
        for candidate in candidates:
            if not os.path.exists(candidate):
                continue
            if os.path.isfile(candidate):
                if os.path.basename(candidate) in names:
                    return candidate
                continue
            for root, _, files in os.walk(candidate):
                for name in names:
                    if name in files:
                        return os.path.join(root, name)
        return None

    def _convert_office_document_to_pdf(self, input_file, original_file_ext):
        if not self._find_office_executable(original_file_ext):
            return None
        temp_pdf = tempfile.NamedTemporaryFile(prefix='anytomd_layout_', suffix='.pdf', delete=False).name
        self._register_temp_file(temp_pdf)
        try:
            import pythoncom
            import win32com.client

            pythoncom.CoInitialize()
            app = None
            doc = None
            try:
                if original_file_ext in ['.doc', '.docx']:
                    app = win32com.client.DispatchEx("Word.Application")
                    app.Visible = False
                    doc = app.Documents.Open(os.path.abspath(input_file))
                    doc.SaveAs2(os.path.abspath(temp_pdf), FileFormat=17)
                elif original_file_ext in ['.ppt', '.pptx']:  # 2026-09-02 修改：旧版 .ppt 与 .pptx 统一走 PowerPoint COM 另存 PDF（保持版式）
                    app = win32com.client.DispatchEx("PowerPoint.Application")
                    doc = app.Presentations.Open(os.path.abspath(input_file), WithWindow=False, ReadOnly=True)
                    doc.SaveAs(os.path.abspath(temp_pdf), 32)
                else:
                    return None
            finally:
                try:
                    if doc:
                        doc.Close()
                except Exception:
                    pass
                try:
                    if app:
                        app.Quit()
                except Exception:
                    pass
                pythoncom.CoUninitialize()
        except Exception:
            return None
        if os.path.exists(temp_pdf) and os.path.getsize(temp_pdf) > 0:
            return temp_pdf
        return None

    # 【逻辑锁定】PDF 处理模块·图片收集成员，未经专门批准禁止修改（见 A 簇头部锁定说明）。
    def _collect_pdf_image_assets(self, input_file):
        try:
            import fitz
            import hashlib
        except Exception:
            return []
        assets = []
        seen = set()
        try:
            doc = fitz.open(input_file)
            index = 0
            for page_no in range(len(doc)):
                page = doc.load_page(page_no)
                for image_info in page.get_images(full=True):
                    xref = image_info[0]
                    base = doc.extract_image(xref)
                    data = base.get('image')
                    if not data:
                        continue
                    digest = hashlib.sha256(data).hexdigest()
                    if digest in seen:
                        continue
                    seen.add(digest)
                    index += 1
                    assets.append((f'PDF图片 {index}', data, self._normalize_image_extension(base.get('ext'), '.png')))
            doc.close()
        except Exception:
            return []
        return assets

    def _collect_word_image_assets(self, input_file):
        try:
            import docx
            import hashlib
        except Exception:
            return []
        assets = []
        seen = set()
        try:
            doc = docx.Document(input_file)
            index = 0
            for rel in doc.part.rels.values():
                if "image" not in rel.reltype:
                    continue
                data = rel.target_part.blob
                digest = hashlib.sha256(data).hexdigest()
                if digest in seen:
                    continue
                seen.add(digest)
                index += 1
                ext = os.path.splitext(rel.target_ref)[1] if '.' in rel.target_ref else '.png'
                assets.append((f'Word图片 {index}', data, self._normalize_image_extension(ext)))
        except Exception:
            return []
        return assets

    def _collect_pptx_image_assets(self, input_file):
        try:
            from pptx import Presentation
            import hashlib
        except Exception:
            return []
        assets = []
        seen = set()
        index = 0
        try:
            prs = Presentation(input_file)
            for slide in prs.slides:
                for shape in slide.shapes:
                    image = getattr(shape, 'image', None)
                    if image is None:
                        continue
                    data = image.blob
                    digest = hashlib.sha256(data).hexdigest()
                    if digest in seen:
                        continue
                    seen.add(digest)
                    index += 1
                    ext = self._normalize_image_extension(getattr(image, 'ext', '.png'))
                    assets.append((f'幻灯片图片 {index}', data, ext))
        except Exception:
            return []
        return assets

    def _collect_html_image_assets(self, input_file):
        try:
            import base64
            import hashlib
            import mimetypes
            import urllib.parse
            import requests
            from bs4 import BeautifulSoup
        except Exception:
            return []
        assets = []
        seen = set()
        try:
            with open(input_file, 'r', encoding='utf-8', errors='ignore') as f:
                html_content = f.read()
            soup = BeautifulSoup(html_content, 'html.parser')
            index = 0
            for img in soup.find_all('img'):
                src = (img.get('src') or '').strip()
                if not src:
                    continue
                data = None
                ext = '.png'
                if src.startswith('data:image'):
                    m = re.match(r'^data:([^;]+);base64,(.+)$', src, flags=re.IGNORECASE | re.DOTALL)
                    if not m:
                        continue
                    try:
                        data = base64.b64decode(m.group(2).strip(), validate=False)
                    except Exception:
                        continue
                    ext = mimetypes.guess_extension(m.group(1).strip()) or '.png'
                elif src.startswith('http://') or src.startswith('https://'):
                    try:
                        response = requests.get(src, timeout=10)
                        data = response.content
                        ext = os.path.splitext(urllib.parse.urlparse(src).path)[1] or '.png'
                    except Exception:
                        continue
                else:
                    src_path = src
                    if not os.path.isabs(src_path):
                        src_path = os.path.join(os.path.dirname(input_file), src_path)
                    try:
                        with open(src_path, 'rb') as img_file:
                            data = img_file.read()
                        ext = os.path.splitext(src)[1] or '.png'
                    except Exception:
                        continue
                if not data:
                    continue
                digest = hashlib.sha256(data).hexdigest()
                if digest in seen:
                    continue
                seen.add(digest)
                index += 1
                title = (img.get('alt') or '').strip() or f'HTML图片 {index}'
                assets.append((title, data, self._normalize_image_extension(ext)))
        except Exception:
            return []
        return assets

    def _select_html_primary_content(self, soup):
        try:
            body = soup.body or soup
            preferred_selectors = [
                'main',
                'article',
                '[role="main"]',
                '#main',
                '#content',
                '#article',
                '#post',
                '.main',
                '.content',
                '.article',
                '.post',
                '.post-content',
                '.entry-content',
                '.article-content',
                '.markdown-body'
            ]
            for selector in preferred_selectors:
                node = body.select_one(selector)
                if node:
                    text_len = len(node.get_text(" ", strip=True))
                    image_count = len(node.find_all('img'))
                    if text_len >= 200 or image_count >= 1:
                        return node

            candidates = []
            for node in body.find_all(['div', 'section', 'article', 'main'], recursive=True):
                text_len = len(node.get_text(" ", strip=True))
                image_count = len(node.find_all('img'))
                score = text_len + image_count * 500
                if score > 0:
                    candidates.append((score, text_len, image_count, node))
            if not candidates:
                return body
            candidates.sort(key=lambda item: item[0], reverse=True)
            best_score, best_text_len, best_image_count, best_node = candidates[0]
            second_score = candidates[1][0] if len(candidates) > 1 else 0
            if best_text_len >= 300 and best_score >= max(600, second_score * 1.6):
                return best_node
            if best_image_count >= 2 and best_score >= max(800, second_score * 1.5):
                return best_node
            return body
        except Exception:
            return soup.body or soup

    def _resolve_html_image_reference(self, img_src, html_path, output_path, images_dir, file_prefix, image_index):
        try:
            import base64
            import mimetypes
            import urllib.parse
            import requests
        except Exception:
            return None
        if not img_src:
            return None
        img_data = None
        img_ext = '.png'
        try:
            if img_src.startswith('data:'):
                m = re.match(r'^data:([^;]+);base64,(.+)$', img_src, flags=re.IGNORECASE | re.DOTALL)
                if not m:
                    return None
                mime_type = m.group(1).strip()
                img_data = base64.b64decode(m.group(2).strip(), validate=False)
                img_ext = mimetypes.guess_extension(mime_type) or '.png'
            elif img_src.startswith('http://') or img_src.startswith('https://'):
                response = requests.get(img_src, timeout=10)
                if response.status_code != 200:
                    return None
                img_data = response.content
                img_ext = os.path.splitext(urllib.parse.urlparse(img_src).path)[1] or '.png'
            else:
                img_path_abs = img_src
                base_dir = os.path.dirname(html_path)
                if not os.path.isabs(img_src):
                    img_path_abs = os.path.join(base_dir, img_src)
                img_path_abs = os.path.abspath(img_path_abs)
                if hasattr(self, 'enable_path_boundary_var') and self.enable_path_boundary_var.get():
                    if not self._is_path_within_dir(img_path_abs, base_dir):
                        return None
                if not os.path.exists(img_path_abs):
                    return None
                with open(img_path_abs, 'rb') as img_file:
                    img_data = img_file.read()
                img_ext = os.path.splitext(img_path_abs)[1] or '.png'
            if not img_data:
                return None
            img_ext = self._normalize_image_extension(img_ext)
            os.makedirs(images_dir, exist_ok=True)
            img_name = f"{file_prefix}_image_{image_index}{img_ext}"
            img_path = os.path.join(images_dir, img_name)
            with open(img_path, 'wb') as img_file:
                img_file.write(img_data)
            if hasattr(self, 'embed_images_var') and self.embed_images_var.get():
                mime_type = mimetypes.guess_type(img_path)[0] or "image/png"
                b64 = base64.b64encode(img_data).decode('ascii')
                return f"data:{mime_type};base64,{b64}"
            return os.path.relpath(img_path, os.path.dirname(output_path)).replace("\\", "/")
        except Exception:
            return None

    def _convert_html_like_to_mixed_markdown(self, input_file, output_path, images_dir):
        try:
            from bs4 import BeautifulSoup, NavigableString, Tag
        except Exception:
            return False
        try:
            with open(input_file, 'r', encoding='utf-8', errors='ignore') as f:
                html_content = f.read()
            # 【书签处理模块·逻辑锁定】以下“书签判定/分派/兜底”一段属锁定区，未经专门批准禁止修改（见本组头部锁定说明）。
            # 2026-09-02 修改（主方案：事件流 + DL 栈）：标准书签文件含 NETSCAPE-Bookmark-file-1 标记，
            # 直接用 HTMLParser 事件流解析，不依赖 BeautifulSoup 树（树容错对真实书签大文件会造成结构丢失，
            # 实测 html.parser 折叠、lxml 亦仅还原 81/4558 条链接）。解析成功即返回；失败走 HTML 内嵌兜底。
            if 'netscape-bookmark-file-1' in (html_content or '').lower():
                if self._convert_bookmark_html_to_markdown(html_content, output_path):
                    return True
                return self._convert_bookmark_html_to_inline_html(html_content, output_path)
            # 非标准/无 NETSCAPE 标记的书签变体：用树结构判定作为兼容入口（仍走事件流解析输出）
            soup = BeautifulSoup(html_content, 'html.parser')
            if self._is_bookmark_html(soup, html_content):
                if self._convert_bookmark_html_to_markdown(html_content, output_path):
                    return True
                return self._convert_bookmark_html_to_inline_html(html_content, output_path)
            for tag in soup(['script', 'style', 'noscript']):
                tag.decompose()
            root = self._select_html_primary_content(soup)
            file_prefix = os.path.splitext(os.path.basename(output_path))[0]
            image_index = 0
            rendered = []

            def inline_markdown(node):
                nonlocal image_index
                parts = []
                for child in node.children:
                    if isinstance(child, NavigableString):
                        text = str(child)
                        if text:
                            parts.append(text)
                        continue
                    if not isinstance(child, Tag):
                        continue
                    name = child.name.lower()
                    if name == 'br':
                        parts.append('\n')
                    elif name == 'img':
                        src = (child.get('src') or '').strip()
                        if src:
                            image_index += 1
                            image_ref = self._resolve_html_image_reference(src, input_file, output_path, images_dir, file_prefix, image_index)
                            if image_ref:
                                alt_text = (child.get('alt') or f'图片 {image_index}').strip()
                                parts.append(f" ![{alt_text}]({image_ref}) ")
                    elif name in ('strong', 'b'):
                        text = inline_markdown(child).strip()
                        if text:
                            parts.append(f"**{text}**")
                    elif name in ('em', 'i'):
                        text = inline_markdown(child).strip()
                        if text:
                            parts.append(f"*{text}*")
                    elif name == 'a':
                        # Extract visible text of the link without recursion
                        text = child.get_text(' ', strip=True) or (child.get('href') or '')
                        href = (child.get('href') or '').strip()
                        if href:
                            parts.append(f"[{text}]({href})")
                        elif text:
                            parts.append(text)
                    else:
                        parts.append(inline_markdown(child))
                text = ''.join(parts)
                text = re.sub(r'[ \t]+\n', '\n', text)
                text = re.sub(r'\n{3,}', '\n\n', text)
                text = re.sub(r'[ \t]{2,}', ' ', text)
                return text.strip()

            def render_block(node):
                nonlocal image_index
                if not isinstance(node, Tag):
                    return
                name = node.name.lower()
                if name in ('script', 'style', 'noscript'):
                    return
                if name in ('header', 'footer', 'nav', 'form'):
                    return
                if name in ('main', 'article', 'section', 'div', 'aside'):
                    block_children = [child for child in node.children if isinstance(child, Tag) and child.name.lower() in ('h1', 'h2', 'h3', 'h4', 'h5', 'h6', 'p', 'ul', 'ol', 'table', 'blockquote', 'pre', 'figure', 'img', 'section', 'article', 'div')]
                    if block_children:
                        for child in node.children:
                            if isinstance(child, Tag):
                                render_block(child)
                    else:
                        text = inline_markdown(node)
                        if text:
                            rendered.append(text)
                    return
                if name in ('h1', 'h2', 'h3', 'h4', 'h5', 'h6'):
                    level = int(name[1])
                    text = inline_markdown(node)
                    if text:
                        rendered.append(f"{'#' * level} {text}")
                    return
                if name == 'p':
                    text = inline_markdown(node)
                    if text:
                        rendered.append(text)
                    return
                if name in ('ul', 'ol'):
                    items = []
                    for index, li in enumerate(node.find_all('li', recursive=False), 1):
                        text = inline_markdown(li)
                        if text:
                            prefix = f"{index}. " if name == 'ol' else "- "
                            items.append(prefix + text)
                    if items:
                        rendered.append("\n".join(items))
                    return
                if name == 'blockquote':
                    text = inline_markdown(node)
                    if text:
                        rendered.append("\n".join(f"> {line}" for line in text.splitlines() if line.strip()))
                    return
                if name == 'pre':
                    text = node.get_text("\n", strip=False).rstrip()
                    if text:
                        rendered.append(f"```\n{text}\n```")
                    return
                if name == 'table':
                    table_html = str(node).strip()
                    if table_html:
                        rendered.append(table_html)
                    return
                if name == 'figure':
                    for child in node.children:
                        if isinstance(child, Tag):
                            render_block(child)
                    return
                if name == 'img':
                    src = (node.get('src') or '').strip()
                    if src:
                        image_index += 1
                        image_ref = self._resolve_html_image_reference(src, input_file, output_path, images_dir, file_prefix, image_index)
                        if image_ref:
                            alt_text = (node.get('alt') or f'图片 {image_index}').strip()
                            rendered.append(f"![{alt_text}]({image_ref})")
                    return

            start_nodes = [child for child in root.children if isinstance(child, Tag)] if hasattr(root, 'children') else []
            if not start_nodes:
                start_nodes = [root] if isinstance(root, Tag) else []
            for child in start_nodes:
                render_block(child)
            if not rendered and isinstance(root, Tag):
                text = inline_markdown(root)
                if text:
                    rendered.append(text)
            rendered = [item.strip() for item in rendered if item and item.strip()]
            if not rendered:
                return False
            self._write_markdown_content(output_path, "\n\n".join(rendered))
            return True
        except Exception as e:
            print(f"HTML图文混排转换失败: {str(e)}")
            return False

        # ==================================================================
        # 书签处理模块【逻辑锁定】（2026-09-02 冻结，与 PDF 冻结区同等待遇）
        # 未经专门批准，禁止修改本模块（含 _is_bookmark_html /
        # _convert_bookmark_html_to_markdown / _convert_bookmark_html_to_inline_html，
        # 以及 _convert_html_like_to_mixed_markdown 中“书签判定与分派”一段）。
        # ------------------------------------------------------------------
        # 问题根因与解决思路摘要（详见开发工作记录第十三~十七轮及
        # 《20260902_AnyToMD书签与PPT转换问题分析与解决方案报告.md》）：
        #  1) Chrome/Netscape 书签是“非规范但被浏览器容错”的 HTML
        #     （<p>/<dt>/<dd> 大量不闭合）。历史上三处踩坑：
        #     a. BeautifulSoup 树容错会折叠/错乱结构：html.parser 把 <dl> 内容折叠进 <p>，
        #        lxml 对真实大文件也只还原 81/4558 条链接 → 大量书签丢失；
        #     b. markitdown 0.1.7 对深层嵌套 HTML 触发 RecursionError 并回退“纯文本提取”，
        #        链接全变文字（早期版本无此限制，故“以前正常、现在回不去”）；
        #     c. 渲染若按 DL 深度缩进 + 空行分隔，深层(≥3 层)条目会以“空行后 ≥4 空格”
        #        被 Markdown 判为缩进代码块，显示成“原始代码”（真实文件 4522/4558 行中招）。
        #  2) 主方案：HTMLParser 事件流 + DL 栈结构化解析 —— 忠实原始标签顺序、
        #     不重构树；<DD> 从不闭合，需在下一 <dt>/<dd> 或 </dl> 到达时自动结束提交备注。
        #  3) 渲染：层级只用标题 # 数表达；书签一律 0 缩进 bullet（组内连续无空行）；
        #     链接文本内 [ ] 转义为 \[ \]；URL 用尖括号包裹 [t](<url>) 容忍括号/空格；
        #     javascript:/data:/vbscript: 等危险协议降级为“标题 + 行内代码”展示。
        #  4) 兜底：主方案异常时输出 HTML 内嵌视图（剔除 base64 ICON、注释、meta/script/style）。
        # ==================================================================

    def _is_bookmark_html(self, soup, html_content):
        try:
            marker = 'NETSCAPE-Bookmark-file-1'
            if marker.lower() in (html_content or '').lower():
                return True
            h1 = soup.find('h1')
            body = soup.body or soup
            has_dl = body.find('dl') is not None
            has_bookmark_nodes = body.find('a') is not None and body.find(['h3', 'dt']) is not None
            if h1 and 'bookmark' in h1.get_text(' ', strip=True).lower() and has_dl and has_bookmark_nodes:
                return True
        except Exception:
            return False
        return False

    def _convert_bookmark_html_to_markdown(self, html_content, output_path):
        """书签 HTML -> Markdown（2026-09-02 重写：HTMLParser 事件流 + DL 栈结构化解析）。

        背景：Chrome/Netscape 书签属非规范 HTML（<p>/<dt>/<dd> 大量不闭合），BeautifulSoup 树容错
        （html.parser 折叠、lxml 亦仅部分还原）都会造成结构丢失（真实文件实测仅还原 81/4558 条链接），
        且 markitdown 0.1.7 对深层嵌套 HTML 会 RecursionError 回退纯文本。
        事件流忠实原始标签顺序，可完整还原全部文件夹与链接。
        签名由 (soup, output_path) 改为 (html_content, output_path)，调用点已同步。
        """
        try:
            from html.parser import HTMLParser
        except Exception:
            return False

        class _BmParser(HTMLParser):  # 2026-09-02 新增：书签事件流解析器（忠实原始标签顺序，不重构树）
            def __init__(self):
                super().__init__(convert_charrefs=True)
                self.dl_depth = 0   # 当前 <dl> 嵌套深度（文件夹层级）
                self.items = []     # (kind, depth, title, href)，kind ∈ {'H3','A','DD'}
                self._in_h3 = False
                self._in_a = False
                self._in_dd = False
                self._buf_h3 = []
                self._buf_a = []
                self._buf_dd = []
                self._href = ''

            def _finish_dd(self):  # 2026-09-02 新增：<DD> 在 Chrome 导出中从不闭合，需在下一 <dt>/<dd> 或 </dl> 到达时自动结束并提交备注
                if self._in_dd:
                    self._in_dd = False
                    note = re.sub(r'\s+', ' ', ''.join(self._buf_dd)).strip()
                    if note:
                        self.items.append(('DD', self.dl_depth, note, ''))

            def handle_starttag(self, tag, attrs):
                t = tag.lower()
                if t == 'dl':
                    self.dl_depth += 1
                    return
                if t in ('dt', 'h3', 'a', 'dd'):
                    self._finish_dd()
                if t == 'h3':
                    self._in_h3 = True
                    self._buf_h3 = []
                elif t == 'a':
                    self._in_a = True
                    self._buf_a = []
                    self._href = dict(attrs).get('href', '')
                elif t == 'dd':
                    self._in_dd = True
                    self._buf_dd = []

            def handle_data(self, data):
                if self._in_h3:
                    self._buf_h3.append(data)
                if self._in_a:
                    self._buf_a.append(data)
                if self._in_dd:
                    self._buf_dd.append(data)

            def handle_endtag(self, tag):
                t = tag.lower()
                if t == 'dl':
                    self._finish_dd()  # 2026-09-02 修改：<dl> 结束前先提交其中未闭合的 <dd> 备注
                    self.dl_depth = max(0, self.dl_depth - 1)
                    return
                if t == 'h3' and self._in_h3:
                    self._in_h3 = False
                    title = re.sub(r'\s+', ' ', ''.join(self._buf_h3)).strip()
                    if title:
                        self.items.append(('H3', self.dl_depth, title, ''))
                elif t == 'a' and self._in_a:
                    self._in_a = False
                    title = re.sub(r'\s+', ' ', ''.join(self._buf_a)).strip()
                    href = (self._href or '').strip()
                    self.items.append(('A', self.dl_depth, title or href, href))
                elif t == 'dd' and self._in_dd:
                    self._in_dd = False
                    note = re.sub(r'\s+', ' ', ''.join(self._buf_dd)).strip()
                    if note:
                        self.items.append(('DD', self.dl_depth, note, ''))

        try:
            parser = _BmParser()
            parser.feed(html_content or '')
            parser.close()
            items = parser.items
            if not items:
                return False

            # 2026-09-02 修改（渲染结构性修复）：层级只由标题 # 表达；书签一律 0 缩进 bullet，
            # 组内链接行之间不加空行。此前按 DL 深度缩进 + 空行分隔，深层(≥3)条目会以
            # “空行后 ≥4 空格”被 Markdown 判为缩进代码块而显示成原始代码（真实文件 4522/4558 行中招）。
            def _esc_title(text):  # 转义链接文本内的方括号，避免破坏内联链接语法
                return (text or '').replace('\\', '\\\\').replace('[', '\\[').replace(']', '\\]')

            def _code_span(text):  # 行内代码（兼容内容中含单反引号）
                return ('``' + text + '``') if '`' in text else ('`' + text + '`')

            _SAFE_SCHEMES = ('http', 'https', 'ftp', 'ftps', 'mailto', 'tel')
            out_lines = ['# Bookmarks']
            for kind, depth, title, href in items:
                depth = max(0, int(depth))
                if kind == 'H3':
                    if out_lines and out_lines[-1]:
                        out_lines.append('')
                    level = max(2, min(depth + 1, 6))
                    out_lines.append(f"{'#' * level} {title}")
                    continue
                if kind == 'A':
                    if out_lines and out_lines[-1]:
                        prev = out_lines[-1]
                        if prev.startswith('- ') or prev.startswith('#'):
                            pass  # 列表内部连续行：不加空行
                        else:
                            out_lines.append('')
                    href = (href or '').strip().replace('\r', ' ').replace('\n', ' ')
                    if not href:
                        out_lines.append(f"- {_esc_title(title)}")
                        continue
                    scheme = re.match(r'^([a-zA-Z][a-zA-Z0-9+.\-]*):', href)
                    safe = bool(scheme and scheme.group(1).lower() in _SAFE_SCHEMES)
                    if safe:
                        # URL 用尖括号包裹，容忍 URL 内含括号/空格等情况
                        out_lines.append(f"- [{_esc_title(title)}](<{href}>)")
                    else:
                        # 危险/非常规协议降级：标题 + 行内代码展示地址（可复制且不破坏渲染）
                        out_lines.append(f"- {_esc_title(title)}：{_code_span(href)}")
                    continue
                if kind == 'DD':
                    if out_lines and out_lines[-1]:
                        if not out_lines[-1].startswith('- ') and not out_lines[-1].startswith('#'):
                            out_lines.append('')
                    if title:
                        out_lines.append(f"> {title}")
            # 合并连续空行、去首尾空行
            merged = []
            prev_blank = False
            for ln in out_lines:
                if not ln.strip():
                    if not prev_blank:
                        merged.append('')
                    prev_blank = True
                else:
                    merged.append(ln)
                    prev_blank = False
            while merged and not merged[0].strip():
                merged.pop(0)
            while merged and not merged[-1].strip():
                merged.pop()
            if len(merged) <= 1:
                return False
            self._write_markdown_content(output_path, "\n".join(merged))
            return True
        except Exception as e:
            print(f"书签HTML转换失败: {str(e)}")
            return False

    # 【逻辑锁定】书签处理模块成员（HTML 内嵌兜底），未经专门批准禁止修改（见本组头部锁定说明）。
    def _convert_bookmark_html_to_inline_html(self, html_content, output_path):
        """补充兜底（2026-09-02 新增）：事件流主方案失败时，将书签以 HTML 内嵌进 MD。

        Markdown 允许内嵌原始 HTML。此处剔除无用且巨大的 base64 ICON 图标数据与注释/脚本，
        保留 <DL>/<DT>/<A> 书签结构原样内嵌，使支持 HTML 渲染的 Markdown 编辑器可显示
        为“定义列表 + 可点击链接”的原生外观。"""
        try:
            content = html_content or ''
            cleaned = re.sub(r'\s+ICON\s*=\s*"data:[^"]*"', '', content, flags=re.IGNORECASE)
            cleaned = re.sub(r'<!--.*?-->', '', cleaned, flags=re.DOTALL)
            cleaned = re.sub(r'<meta\b[^>]*/?>', '', cleaned, flags=re.IGNORECASE)
            cleaned = re.sub(r'<(script|style)\b[^>]*>.*?</\1>', '', cleaned, flags=re.IGNORECASE | re.DOTALL)
            header = ('# Bookmarks（HTML 视图）\n\n'
                      '> 以下为书签的原始 HTML 内嵌视图。请在使用支持 HTML 渲染的 Markdown 编辑器中查看，链接可直接点击。\n\n')
            self._write_markdown_content(output_path, header + cleaned.strip())
            return True
        except Exception as e:
            print(f"书签HTML内嵌兜底失败: {str(e)}")
            return False

    def _write_image_assets_markdown(self, assets, output_path, images_dir, file_prefix):
        os.makedirs(images_dir, exist_ok=True)
        parts = []
        for index, (title, data, ext) in enumerate(assets, 1):
            img_name = f"{file_prefix}_image_{index}{self._normalize_image_extension(ext)}"
            img_path = os.path.join(images_dir, img_name)
            with open(img_path, 'wb') as img_file:
                img_file.write(data)
            rel_path = os.path.relpath(img_path, os.path.dirname(output_path)).replace("\\", "/")
            parts.append(f"![{title}]({rel_path})")
        self._write_markdown_content(output_path, "\n\n".join(parts))
        return True

    # 【逻辑锁定】PDF 处理模块·B 簇成员，未经专门批准禁止修改（见 A 簇头部锁定说明）。
    def _render_pdf_image_block(self, page, block):  # 2026-04-03 20:05:00 修改：新增PDF图片区域渲染，按页面实际显示效果高保真导出图片
        raw_image_bytes = block.get("image")
        raw_image_ext = self._normalize_image_extension(block.get("ext"), ".png")
        try:
            import fitz
            bbox = block.get("bbox") or [0, 0, 0, 0]
            clip_rect = fitz.Rect(bbox)
            if clip_rect.width < 1 or clip_rect.height < 1:
                return raw_image_bytes, raw_image_ext
            scale = max(2.0, min(4.0, 1200.0 / max(clip_rect.width, 1.0)))
            pixmap = page.get_pixmap(matrix=fitz.Matrix(scale, scale), clip=clip_rect, alpha=False)
            return pixmap.tobytes("png"), ".png"
        except Exception:
            return raw_image_bytes, raw_image_ext

    def _extract_pdf_text_sections(self, input_file):
        sections = []
        for page_num, text in enumerate(self._extract_pdf_page_texts(input_file), 1):
            text = self._postprocess_pdf_page_text(text)
            if text:
                sections.append((f'第 {page_num} 页', text))
        return sections

    # 【逻辑锁定】PDF 处理模块·B 簇成员（PDF 图文混排引擎主入口），未经专门批准禁止修改（见 A 簇头部锁定说明）。
    def _convert_pdf_to_mixed_markdown(self, input_file, output_path, images_dir):
        try:
            import base64
            import fitz
            import mimetypes
        except Exception:
            return False
        os.makedirs(images_dir, exist_ok=True)
        embed_images = hasattr(self, 'embed_images_var') and self.embed_images_var.get()
        pdf_prefix = os.path.splitext(os.path.basename(input_file))[0]
        markdown_parts = []
        image_index = 0
        try:
            page_texts = self._extract_pdf_page_texts(input_file)
            doc = fitz.open(input_file)
            for page_num, page in enumerate(doc, 1):
                page_parts = []
                page_text = page_texts[page_num - 1] if page_num - 1 < len(page_texts) else ''
                page_text = self._postprocess_pdf_page_text(page_text)
                if page_text:
                    page_parts.append(page_text)

                blocks = page.get_text("dict").get("blocks", [])
                for block in blocks:
                    block_type = block.get("type")
                    if block_type != 1:
                        continue
                    bbox = block.get("bbox") or [0, 0, 0, 0]
                    width = max(0, (bbox[2] or 0) - (bbox[0] or 0))
                    height = max(0, (bbox[3] or 0) - (bbox[1] or 0))
                    if width < 48 or height < 48:
                        continue
                    image_bytes, image_ext = self._render_pdf_image_block(page, block)  # 2026-04-03 20:05:00 修改：图片改为按页面区域渲染，避免颜色失真与看不清
                    if not image_bytes:
                        continue
                    image_index += 1
                    image_name = f"{pdf_prefix}_page_{page_num}_image_{image_index}{image_ext}"
                    image_path = os.path.join(images_dir, image_name)
                    with open(image_path, "wb") as image_file:
                        image_file.write(image_bytes)
                    if embed_images:
                        mime_type = mimetypes.guess_type(image_path)[0] or "image/png"
                        image_ref = f"data:{mime_type};base64,{base64.b64encode(image_bytes).decode('ascii')}"
                    else:
                        image_ref = os.path.relpath(image_path, os.path.dirname(output_path)).replace("\\", "/")
                    page_parts.append(f"![第 {page_num} 页图片 {image_index}]({image_ref})")
                if page_parts:
                    markdown_parts.append(f"## 第 {page_num} 页\n\n" + "\n\n".join(page_parts))
            doc.close()
        except Exception as e:
            print(f"PDF图文混排转换失败: {str(e)}")
            return False
        if not markdown_parts:
            return False
        self._write_markdown_content(output_path, "\n\n".join(markdown_parts))
        return True

    def _extract_word_text_sections(self, input_file):
        try:
            import docx
            doc = docx.Document(input_file)
            lines = [para.text.strip() for para in doc.paragraphs if para.text.strip()]
            return [('正文', "\n\n".join(lines))] if lines else []
        except Exception:
            return []

    def _extract_pptx_text_sections(self, input_file):
        try:
            from pptx import Presentation
        except Exception:
            return []
        sections = []
        try:
            prs = Presentation(input_file)
            for slide_index, slide in enumerate(prs.slides, 1):
                lines = []
                for shape in slide.shapes:
                    text = getattr(shape, 'text', '')
                    if text and text.strip():
                        lines.append(text.strip())
                if lines:
                    sections.append((f'第 {slide_index} 张幻灯片', "\n\n".join(lines)))
        except Exception:
            return []
        return sections

    def _extract_html_text_sections(self, input_file):
        try:
            from bs4 import BeautifulSoup
        except Exception:
            return []
        try:
            with open(input_file, 'r', encoding='utf-8', errors='ignore') as f:
                html_content = f.read()
            soup = BeautifulSoup(html_content, 'html.parser')
            for tag in soup(['script', 'style', 'noscript']):
                tag.decompose()
            text = soup.get_text("\n")
            text = re.sub(r'\n{3,}', '\n\n', text).strip()
            return [('正文', text)] if text else []
        except Exception:
            return []

    def _ocr_image_bytes(self, image_bytes):
        try:
            import pytesseract
            from PIL import Image
            from io import BytesIO
            pytesseract.get_tesseract_version()
            image = Image.open(BytesIO(image_bytes))
            if image.mode not in ('L', 'RGB'):
                image = image.convert('RGB')
            for lang in ('chi_sim+eng', 'eng'):
                try:
                    text = pytesseract.image_to_string(image, lang=lang)
                    if text and text.strip():
                        return text.strip()
                except Exception:
                    continue
        except Exception:
            return ''
        return ''

    def _extract_pdf_text_sections_with_ocr(self, input_file):
        raw_page_texts = self._extract_pdf_page_texts(input_file)
        processed_page_texts = [self._postprocess_pdf_page_text(text) for text in raw_page_texts]
        try:
            from pdf2image import convert_from_path
        except Exception:
            return [
                (f'第 {page_num} 页', text)
                for page_num, text in enumerate(processed_page_texts, 1)
                if text
            ]
        poppler_path = self.find_poppler_bin()
        try:
            images = convert_from_path(input_file, dpi=180, poppler_path=poppler_path)
        except Exception:
            return [
                (f'第 {page_num} 页', text)
                for page_num, text in enumerate(processed_page_texts, 1)
                if text
            ]
        sections = []
        for page_index, image in enumerate(images, 1):
            base_text_raw = raw_page_texts[page_index - 1] if page_index - 1 < len(raw_page_texts) else ''
            base_text_processed = processed_page_texts[page_index - 1] if page_index - 1 < len(processed_page_texts) else ''
            if len(self._normalize_compare_text(base_text_raw)) >= 30:
                if base_text_processed:
                    sections.append((f'第 {page_index} 页', base_text_processed))
                continue
            temp_buffer = tempfile.NamedTemporaryFile(prefix='anytomd_ocr_', suffix='.png', delete=False)
            temp_buffer.close()
            self._register_temp_file(temp_buffer.name)
            try:
                image.save(temp_buffer.name, 'PNG')
                with open(temp_buffer.name, 'rb') as img_file:
                    ocr_text = self._ocr_image_bytes(img_file.read())
            except Exception:
                ocr_text = ''
            final_text = self._postprocess_pdf_page_text(ocr_text or base_text_processed or base_text_raw)
            if final_text:
                sections.append((f'第 {page_index} 页', final_text))
        return sections

    def _collect_image_ocr_sections(self, assets, title_prefix):
        sections = []
        for index, (_, data, _) in enumerate(assets, 1):
            text = self._ocr_image_bytes(data)
            if text:
                sections.append((f'{title_prefix} {index}', text))
        return sections

    def _convert_document_to_layout_markdown(self, original_input_file, input_file, output_path, images_dir, original_file_ext):
        pdf_source = None
        if original_file_ext == '.pdf':
            pdf_source = original_input_file
        elif original_file_ext in ['.doc', '.docx', '.ppt', '.pptx']:  # 2026-09-02 修改：旧版 .ppt 一并支持 Office 版式保持
            pdf_source = self._convert_office_document_to_pdf(input_file, original_file_ext)
        elif original_file_ext in ['.html', '.htm', '.mhtml', '.mht']:
            pdf_source = self._convert_html_like_to_pdf(input_file)
        if not pdf_source:
            return False
        embed_images = hasattr(self, 'embed_images_var') and self.embed_images_var.get()
        return self._convert_pdf_pages_to_markdown(pdf_source, output_path, images_dir, embed_images=embed_images)

    def _convert_document_to_image_markdown(self, original_input_file, input_file, output_path, images_dir, original_file_ext):
        if original_file_ext == '.pdf':
            assets = self._collect_pdf_image_assets(original_input_file)
            if not assets and self._is_scanned_pdf(original_input_file):
                embed_images = hasattr(self, 'embed_images_var') and self.embed_images_var.get()
                return self._convert_pdf_pages_to_markdown(original_input_file, output_path, images_dir, embed_images=embed_images)
        elif original_file_ext in ['.doc', '.docx']:
            assets = self._collect_word_image_assets(input_file)
        elif original_file_ext == '.pptx':
            assets = self._collect_pptx_image_assets(input_file)
        elif original_file_ext in ['.html', '.htm', '.mhtml', '.mht']:
            assets = self._collect_html_image_assets(input_file)
        else:
            assets = []
        return self._write_image_assets_markdown(
            assets,
            output_path,
            images_dir,
            os.path.splitext(os.path.basename(original_input_file))[0]
        )

    def _convert_document_to_text_markdown(self, original_input_file, input_file, output_path, original_file_ext):
        if original_file_ext == '.pdf':
            sections = self._extract_pdf_text_sections(original_input_file)
        elif original_file_ext in ['.doc', '.docx']:
            if original_file_ext == '.doc' and str(input_file).lower().endswith('.doc'):
                text = self.extract_text_from_doc(input_file)
                sections = [('正文', text)] if text else []
            else:
                sections = self._extract_word_text_sections(input_file)
        elif original_file_ext == '.pptx':
            sections = self._extract_pptx_text_sections(input_file)
        elif original_file_ext in ['.html', '.htm', '.mhtml', '.mht']:
            sections = self._extract_html_text_sections(input_file)
        else:
            sections = []
        if sections:
            return self._write_markdown_sections(output_path, sections)
        return self._convert_with_markitdown_api(input_file, output_path)

    def _convert_document_to_text_with_ocr_markdown(self, original_input_file, input_file, output_path, original_file_ext):
        if original_file_ext == '.pdf':
            sections = self._extract_pdf_text_sections_with_ocr(original_input_file)
            if sections:
                return self._write_markdown_sections(output_path, sections)
            return self._convert_document_to_text_markdown(original_input_file, input_file, output_path, original_file_ext)
        if original_file_ext in ['.doc', '.docx']:
            if original_file_ext == '.doc' and str(input_file).lower().endswith('.doc'):
                text = self.extract_text_from_doc(input_file)
                base_sections = [('正文', text)] if text else []
                ocr_sections = []
            else:
                base_sections = self._extract_word_text_sections(input_file)
                ocr_sections = self._collect_image_ocr_sections(self._collect_word_image_assets(input_file), '识别图片')
        elif original_file_ext == '.pptx':
            base_sections = self._extract_pptx_text_sections(input_file)
            ocr_sections = self._collect_image_ocr_sections(self._collect_pptx_image_assets(input_file), '识别图片')
        elif original_file_ext in ['.html', '.htm', '.mhtml', '.mht']:
            base_sections = self._extract_html_text_sections(input_file)
            ocr_sections = self._collect_image_ocr_sections(self._collect_html_image_assets(input_file), '识别图片')
        else:
            base_sections = []
            ocr_sections = []
        sections = self._filter_ocr_sections(base_sections, ocr_sections)
        if sections:
            return self._write_markdown_sections(output_path, sections)
        return self._convert_document_to_text_markdown(original_input_file, input_file, output_path, original_file_ext)

    def _try_ppt_smart_via_pdf(self, original_file_ext, input_file, output_path, images_dir):
        """2026-09-02 新增：PPT 智能（图文并存）增强——“重映射 + 复用”。

        若本机装有 Office（PowerPoint），先把 .ppt/.pptx 渲染为 PDF，再复用现有 PDF 图文混排引擎
        _convert_pdf_to_mixed_markdown 输出“页级文字 + 图块按序混排”，从而绕开 markitdown 0.1.7
        “图片全前置、背景图每页重复、图片断链”的问题；Office 缺失或渲染失败返回 False，由调用方回退原路径。
        """
        if original_file_ext not in ('.ppt', '.pptx'):
            return False
        try:
            pdf_source = self._convert_office_document_to_pdf(input_file, original_file_ext)
            if not pdf_source or not os.path.exists(pdf_source):
                return False
            ok = self._convert_pdf_to_mixed_markdown(pdf_source, output_path, images_dir)
            return bool(ok)
        except Exception as e:
            print(f"PPT 智能渲染增强失败(将回退原路径): {e}")
            return False

    def _convert_document_by_mode(self, original_input_file, input_file, output_path, images_dir, original_file_ext):
        mode = self._get_document_output_mode()
        if mode == 'layout':
            return self._convert_document_to_layout_markdown(original_input_file, input_file, output_path, images_dir, original_file_ext)
        if mode == 'image':
            return self._convert_document_to_image_markdown(original_input_file, input_file, output_path, images_dir, original_file_ext)
        if mode == 'text':
            return self._convert_document_to_text_markdown(original_input_file, input_file, output_path, original_file_ext)
        if mode == 'text_ocr':
            return self._convert_document_to_text_with_ocr_markdown(original_input_file, input_file, output_path, original_file_ext)
        return False

    def _show_conversion_result(self, success, message, output_path=None, output_folder=None):
        """在主线程中显示转换结果"""
        if success:
            messagebox.showinfo('成功', message)
            if output_path:
                self.last_output_path = output_path
                try:
                    # 检查文件大小，避免大文件卡死 UI
                    file_size = os.path.getsize(output_path)
                    if file_size > 1024 * 1024: # 大于 1MB
                        markdown_content = "文件过大 (%.2f MB)，无法在预览窗口直接显示。请在输出目录中查看。" % (file_size / (1024 * 1024))
                    else:
                        with open(output_path, 'r', encoding='utf-8') as f:
                            markdown_content = f.read()
                    
                    self.preview_text.delete(1.0, tk.END)
                    self.preview_text.insert(tk.END, markdown_content)
                except Exception as e:
                    print(f"预览更新失败: {e}")
            
            if output_folder and self.auto_open_folder_var.get():
                self.open_output_folder(output_folder)
        else:
            messagebox.showerror('错误', message)

    def convert_file(self):
        """转换单个文件"""
        try:
            self.conversion_in_progress = True
            self.root.after(0, lambda: self._set_widget_enabled(self.convert_btn, False))
            self._update_ui(status='正在转换...', progress=0)
            
            input_file = self.input_path_var.get()
            original_input_file = input_file
            output_folder = self.output_path_var.get()
            
            # 初始化临时文件列表（如果还没有）
            if not hasattr(self, 'temp_files'):
                self.temp_files = []

            # 创建图片目录
            images_dir = os.path.join(output_folder, "images")
            os.makedirs(images_dir, exist_ok=True)
            
            # 记录原始文件类型用于后续处理
            original_file_ext = os.path.splitext(input_file)[1].lower()
            
            # 处理.doc文件
            if original_file_ext == '.doc':
                input_file = self.handle_doc_file(input_file)
            
            # 生成输出路径
            os.makedirs(output_folder, exist_ok=True)
            output_filename = os.path.splitext(os.path.basename(self.input_path_var.get()))[0] + '.md'
            output_path = os.path.join(output_folder, output_filename)
            document_mode = self._get_document_output_mode()
            document_mode_exts = ['.pdf', '.docx', '.doc', '.ppt', '.pptx', '.html', '.htm', '.mhtml', '.mht']  # 2026-09-02 修改：补充旧版 .ppt，使其可用“保持版式”等模式
            if document_mode != 'smart' and original_file_ext in document_mode_exts:
                if original_file_ext in ['.mhtml', '.mht']:
                    self._update_ui(status='正在解析HTML/MHTML文件...', progress=10)
                    temp_html = self._parse_mhtml_to_html(input_file, images_dir, embed_images=True, mode=self._get_mhtml_mode())
                    if not temp_html:
                        self._update_ui(status='MHTML解析失败')
                        self.root.after(0, lambda: messagebox.showerror('转换失败', 'MHTML解析失败'))
                        return
                    input_file = temp_html
                conversion_success = self._convert_document_by_mode(
                    original_input_file,
                    input_file,
                    output_path,
                    images_dir,
                    original_file_ext
                )
                if conversion_success:
                    self._update_ui(status='转换完成', progress=100)
                    msg = f'文件已成功转换为 Markdown 格式并保存到:\n{output_path}'
                    self.root.after(0, lambda: self._show_conversion_result(True, msg, output_path, output_folder))
                    return
                if document_mode == 'layout':
                    self._update_ui(status='版式转换失败，自动回退智能转换...', progress=35)
                    input_file = original_input_file
                else:
                    self._update_ui(status='转换失败')
                    self.root.after(0, lambda: messagebox.showerror('转换失败', '文档输出模式转换失败'))
                    return
            
            # 处理.mhtml/.mht/html/htm文件
            if original_file_ext in ['.mhtml', '.mht', '.html', '.htm']:
                self._update_ui(status='正在解析HTML/MHTML文件...', progress=10)
                if original_file_ext in ['.mhtml', '.mht']:
                    mhtml_mode = self._get_mhtml_mode()
                    # MHTML 强制使用内嵌图片，以保证其作为独立包的完整性
                    embed_images = True 
                    
                    if mhtml_mode == 'layout':
                        self._update_ui(status='MHTML版式保持：提取CSS并保持原样式布局...', progress=30)
                    elif mhtml_mode == 'text':
                        self._update_ui(status='MHTML文本优先：清理样式提取纯文本...', progress=30)
                    elif mhtml_mode == 'image':
                        self._update_ui(status='MHTML图片优先：重点处理高保真图片...', progress=30)
                    
                    # 将 mhtml_mode 传递给解析器
                    temp_html = self._parse_mhtml_to_html(input_file, images_dir, embed_images=embed_images, mode=mhtml_mode)
                    if temp_html:
                        input_file = temp_html
                    else:
                        self._update_ui(status='MHTML解析失败')
                        self.root.after(0, lambda: messagebox.showerror('转换失败', 'MHTML解析失败'))
                        return
                
                conversion_success = self._convert_html_like_to_mixed_markdown(input_file, output_path, images_dir)
                if not conversion_success:
                    conversion_success = self.convert_html_to_markdown(input_file, output_path)
                if conversion_success:
                    self._update_ui(status='基本转换完成，处理图片...', progress=50)
                    if self.extract_images_var.get() and not (hasattr(self, 'embed_images_var') and self.embed_images_var.get()):
                        # 只有当没有选择内嵌图片时，才处理图片保存
                        self._update_ui(status='处理并保留图片原位...')
                        self._process_html_images(input_file, output_path, images_dir)
                    
                    # 如果选择了内嵌图片，确保本地图片也被内嵌
                    if hasattr(self, 'embed_images_var') and self.embed_images_var.get():
                        # 跳过_process_html_images，直接内嵌本地图片
                        self._inline_local_images_in_md(output_path)
                    
                    self._update_ui(status='转换完成', progress=100)
                    
                    msg = f'文件已成功转换为 Markdown 格式并保存到:\n{output_path}'
                    self.root.after(0, lambda: self._show_conversion_result(True, msg, output_path, output_folder))
                else:
                    self._update_ui(status='转换失败')
                    self.root.after(0, lambda: messagebox.showerror('转换失败', 'HTML转换失败：HTML到Markdown转换未成功'))
                return
            
            if original_file_ext == '.pdf':
                pdf_mode = self._get_pdf_mode()
                embed_images = hasattr(self, 'embed_images_var') and self.embed_images_var.get()

                if pdf_mode == 'layout':
                    self._update_ui(status='PDF版式保持：按页转换为图片...', progress=30)
                    conversion_success = self._convert_pdf_pages_to_markdown(
                        self.input_path_var.get(),
                        output_path,
                        images_dir,
                        embed_images=embed_images
                    )
                    if conversion_success:
                        self._update_ui(status='转换完成', progress=100)
                        msg = f'文件已成功转换为 Markdown 格式并保存到:\n{output_path}'
                        self.root.after(0, lambda: self._show_conversion_result(True, msg, output_path, output_folder))
                    else:
                        self._update_ui(status='转换失败')
                        self.root.after(0, lambda: messagebox.showerror('转换失败', 'PDF版式转换失败'))
                    return

                if pdf_mode == 'auto_scanned':
                    is_scanned = self._is_scanned_pdf(self.input_path_var.get())
                    if is_scanned:
                        self._update_ui(status='检测到扫描件PDF，转换为图片...', progress=30)
                        conversion_success = self._convert_pdf_pages_to_markdown(
                            self.input_path_var.get(),
                            output_path,
                            images_dir,
                            embed_images=embed_images
                        )
                        if conversion_success:
                            self._update_ui(status='转换完成', progress=100)
                            msg = f'文件已成功转换为 Markdown 格式并保存到:\n{output_path}'
                            self.root.after(0, lambda: self._show_conversion_result(True, msg, output_path, output_folder))
                        else:
                            self._update_ui(status='转换失败')
                            self.root.after(0, lambda: messagebox.showerror('转换失败', 'PDF扫描件转换失败'))
                        return

                if pdf_mode == 'normal':
                    self._update_ui(status='PDF智能转换：按原文图文顺序重建内容...', progress=25)
                    
                    # 2026-09-01 新增：PDF 引擎策略（仅作用于智能转换模式）
                    engine_strategy = self._get_pdf_engine_strategy()
                    conversion_success = False
                    if engine_strategy == 'auto':
                        # 自动：detect_pdf 分类，text_based 用 pdf-inspector，否则回退图文混排
                        conversion_success = self._convert_pdf_with_pdfinspector(
                            self.input_path_var.get(),
                            output_path,
                            force=False
                        )
                    elif engine_strategy == 'inspector_first':
                        # pdf-inspector 优先：不预判类型直接尝试，失败回退图文混排
                        conversion_success = self._convert_pdf_with_pdfinspector(
                            self.input_path_var.get(),
                            output_path,
                            force=True
                        )
                    if conversion_success:
                        self._update_ui(status='转换完成', progress=100)
                        msg = f'文件已成功转换为 Markdown 格式并保存到:\n{output_path}'
                        self.root.after(0, lambda: self._show_conversion_result(True, msg, output_path, output_folder))
                        return
                    # 回退：现有图文混排路径（layout_first 或 pdf-inspector 失败）
                    conversion_success = self._convert_pdf_to_mixed_markdown(
                        self.input_path_var.get(),
                        output_path,
                        images_dir
                    )
                    if conversion_success:
                        self._update_ui(status='转换完成', progress=100)
                        msg = f'文件已成功转换为 Markdown 格式并保存到:\n{output_path}'
                        self.root.after(0, lambda: self._show_conversion_result(True, msg, output_path, output_folder))
                        return

            # 2026-09-02 新增：PPT 智能（图文并存）增强——Office→PDF→复用 PDF 图文混排引擎
            if original_file_ext in ('.ppt', '.pptx'):
                self._update_ui(status='PPT智能转换：Office渲染PDF后按页图文混排...', progress=35)
                if self._try_ppt_smart_via_pdf(original_file_ext, input_file, output_path, images_dir):
                    self._update_ui(status='转换完成', progress=100)
                    msg = f'文件已成功转换为 Markdown 格式并保存到:\n{output_path}'
                    self.root.after(0, lambda: self._show_conversion_result(True, msg, output_path, output_folder))
                    return
                self._update_ui(status='Office渲染不可用，回退默认转换引擎...', progress=35)

            cmd = ["markitdown", input_file, "-o", output_path]
            
            # 如果选择提取图片，我们需要让MarkItDown保留Base64数据，以便后续提取和替换
            if self.extract_images_var.get() and original_file_ext in ['.docx', '.doc', '.pptx']:
                cmd.append("--keep-data-uris")
                
            print(f"执行命令: {' '.join(cmd)}")
            process = None
            try:
                process = subprocess.run(
                    cmd,
                    capture_output=True,
                    text=True,
                    timeout=_MARKITDOWN_CLI_TIMEOUT,
                    creationflags=subprocess.CREATE_NO_WINDOW
                )
            except FileNotFoundError:  # 2026-03-29 13:41:12 修改：markitdown命令不可用时自动API回退
                ok = self._convert_with_markitdown_api(input_file, output_path)
                if ok:
                    class _Proc:
                        returncode = 0
                        stderr = ""
                    process = _Proc()
                else:
                    raise
            
            # 检查基本转换是否成功
            if process.returncode == 0:
                self._update_ui(status='基本转换完成，处理图片...', progress=50)
                
                # 如果用户选择了提取图片选项，执行自定义图片提取
                if self.extract_images_var.get():
                    self._update_ui(status='处理并保留图片原位...')
                    
                    if original_file_ext in ['.docx', '.doc', '.pptx']:
                        self._process_base64_images_in_md(output_path, images_dir)
                        self._apply_office_image_output_policy(original_file_ext, input_file, output_path, images_dir)  # 2026-03-29 14:05:06 修改：使用实际参与转换的文件路径，避免DOC->DOCX时提图指向原DOC
                    elif original_file_ext == '.pdf':
                        self._apply_pdf_image_output_policy(self.input_path_var.get(), output_path, images_dir)
                    elif original_file_ext in ['.html', '.htm', '.mhtml', '.mht']:
                        # HTML/MHTML已经在解析或转换时处理好了相对路径，我们只需要确保本地图片文件被拷贝
                        self._process_html_images(input_file if original_file_ext in ['.mhtml', '.mht'] else self.input_path_var.get(), output_path, images_dir)

                    if hasattr(self, 'embed_images_var') and self.embed_images_var.get():
                        self._inline_local_images_in_md(output_path)
                
                self._update_ui(status='转换完成', progress=100)
                
                msg = f'文件已成功转换为 Markdown 格式并保存到:\n{output_path}'
                self.root.after(0, lambda: self._show_conversion_result(True, msg, output_path, output_folder))
            else:
                stderr_output = process.stderr
                
                # AnyDoc 兜底：以上所有方法均失败时，对 Office 格式启用 AnyDoc（新增处理方法）
                if original_file_ext in ['.doc', '.docx', '.ppt', '.pptx', '.xls', '.xlsx']:
                    if self.convert_office_with_anydoc(original_input_file, output_path):
                        self._update_ui(status='转换完成', progress=100)
                        msg = f'文件已成功转换为 Markdown 格式并保存到:\n{output_path}'
                        self.root.after(0, lambda: self._show_conversion_result(True, msg, output_path, output_folder))
                        return
                
                self._update_ui(status='转换失败')
                
                # 针对 PDF 依赖缺失的特殊处理
                if "MissingDependencyException" in stderr_output and "pdf" in stderr_output.lower():
                    error_msg = "检测到 MarkItDown 缺少 PDF 转换组件。\n\n请在终端运行以下命令修复：\npip install \"markitdown[pdf]\""
                else:
                    error_msg = f'转换过程中出错:\n{stderr_output}'
                
                self.root.after(0, lambda: messagebox.showerror('转换失败', error_msg))
        except Exception as e:
            self._update_ui(status=f"转换出错: {str(e)}")
            self.root.after(0, lambda: messagebox.showerror('错误', f'转换过程中出错:\n{str(e)}'))
        finally:
            self.conversion_in_progress = False
            # 将 UI 状态重置调度回主线程
            def reset_ui_state():
                self._set_widget_enabled(self.convert_btn, True)
                self._set_widget_enabled(self.preview_btn, True)
                # 延迟重置进度条和状态，让用户看到完成状态
                def reset_progress():
                    self._set_progress_percent(self.progress, 0)
                    if self.status_var.get() in ['转换完成', '转换成功']:
                        self._update_ui(status='就绪')
                self.root.after(3000, reset_progress)
            
            self.root.after(0, reset_ui_state)
            
            # 清理临时文件
            if hasattr(self, 'temp_files'):
                for temp_file in self.temp_files:
                    try:
                        if os.path.exists(temp_file):
                            os.remove(temp_file)
                    except Exception as e:
                        print(f"清理临时文件 {temp_file} 失败: {e}")
                self.temp_files = []

    def _process_base64_images_in_md(self, md_path, images_dir):
        """解析Markdown文件中的base64图片，保存为文件并将MD中的base64替换为相对路径"""
        import base64
        import re
        import mimetypes
        
        try:
            with open(md_path, 'r', encoding='utf-8') as f:
                content = f.read()
                
            # 正则表达式匹配 base64 图片: ![alt](data:image/png;base64,xxxx)
            pattern = r'!\[(.*?)\]\(data:(image/[a-zA-Z0-9]+);base64,([^)]+)\)'
            
            img_count = 0
            doc_prefix = os.path.splitext(os.path.basename(md_path))[0]
            
            def replace_image(match):
                nonlocal img_count
                alt_text = match.group(1)
                mime_type = match.group(2)
                b64_data = match.group(3)
                
                img_count += 1
                
                # 确定扩展名
                ext = mimetypes.guess_extension(mime_type) or '.png'
                
                # 清理alt_text，使其适合作为文件名的一部分（可选）
                safe_alt = "".join(c for c in alt_text if c.isalnum() or c in (' ', '_', '-'))[:20]
                if safe_alt:
                    img_name = f"{doc_prefix}_{safe_alt}_{img_count}{ext}"
                else:
                    img_name = f"{doc_prefix}_image_{img_count}{ext}"
                    
                img_path = os.path.join(images_dir, img_name)
                
                try:
                    img_data = base64.b64decode(b64_data, validate=False)
                    with open(img_path, 'wb') as img_f:
                        img_f.write(img_data)
                        
                    rel_path = f"images/{img_name}"
                    if hasattr(self, 'embed_images_var') and self.embed_images_var.get():
                        return match.group(0)
                    return f"![{alt_text}]({rel_path})"
                except Exception as e:
                    print(f"保存图片 {img_name} 失败: {str(e)}")
                    return match.group(0) # 失败则保留原样
                    
            # 替换内容
            new_content = re.sub(pattern, replace_image, content)
            
            # 写回文件
            if img_count > 0:
                with open(md_path, 'w', encoding='utf-8') as f:
                    f.write(new_content)
                self._update_ui(status=f'已提取并保留 {img_count} 张原位图片')
                print(f"已处理 {img_count} 张 base64 图片")
                
        except Exception as e:
            print(f"处理 base64 图片失败: {str(e)}")
            self._update_ui(status=f'图片处理失败: {str(e)}')

    def _process_html_images(self, html_path, md_path, images_dir):
        """处理HTML文档中的图片，将其下载/拷贝到本地并替换MD中的路径"""
        import re
        import urllib.parse
        import requests
        import base64
        import mimetypes
        
        try:
            with open(md_path, 'r', encoding='utf-8') as f:
                md_content = f.read()
                
            # 正则匹配普通图片链接 ![alt](src) 和 HTML 格式的 <img src="src" ...>
            # 排除已经是 images/ 开头的链接
            md_img_pattern = r'!\[(.*?)\]\((?!images/)(.*?)\)'
            html_img_pattern = r'<img\s+[^>]*?src=["\'](?!images/)(.*?)["\'][^>]*?>'
            
            img_count = 0
            doc_prefix = os.path.splitext(os.path.basename(md_path))[0]
            
            def get_and_save_image(img_src, alt_text="image"):
                nonlocal img_count
                img_count += 1
                img_data = None
                img_ext = '.png'
                
                try:
                    if img_src.startswith('data:'):
                        m = re.match(r'^data:([^;]+);base64,(.+)$', img_src, flags=re.IGNORECASE | re.DOTALL)
                        if not m:
                            return None
                        mime_type = m.group(1).strip()
                        b64_part = m.group(2).strip()
                        img_data = base64.b64decode(b64_part, validate=False)
                        img_ext = mimetypes.guess_extension(mime_type) or '.png'
                    elif img_src.startswith('http://') or img_src.startswith('https://'):
                    # 处理网络图片
                        response = requests.get(img_src, timeout=10)
                        if response.status_code == 200:
                            img_data = response.content
                            parsed_url = urllib.parse.urlparse(img_src)
                            ext = os.path.splitext(parsed_url.path)[1]
                            if ext and len(ext) <= 5:
                                img_ext = ext
                    # 处理本地图片
                    else:
                        img_path_abs = img_src
                        base_dir = os.path.dirname(html_path)
                        if not os.path.isabs(img_src):
                            img_path_abs = os.path.join(base_dir, img_src)
                        img_path_abs = os.path.abspath(img_path_abs)
                        if hasattr(self, 'enable_path_boundary_var') and self.enable_path_boundary_var.get():
                            if not self._is_path_within_dir(img_path_abs, base_dir):
                                return None
                        if os.path.exists(img_path_abs):
                            with open(img_path_abs, 'rb') as img_file:
                                img_data = img_file.read()
                            img_ext = os.path.splitext(img_src)[1]
                            
                    if img_data:
                        if not img_ext or len(img_ext) > 10:
                            img_ext = '.png'

                        img_name = f"{doc_prefix}_image_{img_count}{img_ext}"
                        img_save_path = os.path.join(images_dir, img_name)
                        with open(img_save_path, 'wb') as f:
                            f.write(img_data)

                        if hasattr(self, 'embed_images_var') and self.embed_images_var.get():
                            mime_type = mimetypes.guess_type(img_save_path)[0] or "image/png"
                            b64 = base64.b64encode(img_data).decode('ascii')
                            return f"data:{mime_type};base64,{b64}"

                        return f"images/{img_name}"
                except Exception as e:
                    print(f"处理图片 {img_src} 时出错: {str(e)}")
                return None

            def replace_md_image(match):
                alt_text = match.group(1)
                img_src = match.group(2)
                new_rel_path = get_and_save_image(img_src, alt_text)
                if new_rel_path:
                    return f"![{alt_text}]({new_rel_path})"
                return match.group(0)

            def replace_html_img_tag(match):
                img_src = match.group(1)
                
                # 尝试提取 alt 属性
                alt_match = re.search(r'alt=["\'](.*?)["\']', match.group(0))
                alt_text = alt_match.group(1) if alt_match else "image"
                
                new_rel_path = get_and_save_image(img_src, alt_text)
                if new_rel_path:
                    return f"![{alt_text}]({new_rel_path})"
                return match.group(0)
                
            # 替换内容
            md_content = re.sub(md_img_pattern, replace_md_image, md_content)
            md_content = re.sub(html_img_pattern, replace_html_img_tag, md_content)
            
            # 写回文件
            if img_count > 0:
                with open(md_path, 'w', encoding='utf-8') as f:
                    f.write(md_content)
                self._update_ui(status=f'已提取并保留 {img_count} 张 HTML 原位图片')
                
        except Exception as e:
            print(f"处理 HTML 图片失败: {str(e)}")
            self._update_ui(status=f'处理 HTML 图片失败: {str(e)}')

    # 添加图片提取方法
    def extract_images_from_word(self, input_file, output_path, images_dir):
        """从Word文档提取图片并添加到Markdown文件中"""
        try:
            import hashlib  # 2026-03-29 14:30:53 修改：仅补缺模式使用hash去重，避免文末重复图片
            # 从docx中提取图片
            import docx
            doc = docx.Document(input_file)
            
            # 文档名前缀
            doc_prefix = os.path.splitext(os.path.basename(input_file))[0]
            
            # 保存提取的图片
            img_count = 0
            policy = self._get_image_output_policy()  # 2026-03-29 14:30:53 修改：根据图片输出策略决定是否追加/补缺
            if policy == 'inplace':  # 2026-03-29 14:30:53 修改：仅原位不做文末追加
                return
            existing_hashes = self._get_existing_image_hashes(output_path) if policy == 'supplement' else set()  # 2026-03-29 14:30:53 修改：补缺模式按hash去重
            
            # 遍历文档中的关系
            for rel in doc.part.rels.values():
                if "image" in rel.reltype:
                    try:
                        img_count += 1
                        img_ext = os.path.splitext(rel.target_ref)[1] if '.' in rel.target_ref else '.png'
                        img_name = f"{doc_prefix}_image_{img_count}{img_ext}"
                        img_path = os.path.join(images_dir, img_name)
                        
                        # 保存图片
                        with open(img_path, "wb") as f:
                            img_bytes = rel.target_part.blob  # 2026-03-29 14:30:53 修改：用于hash去重与保存
                            f.write(img_bytes)
                        img_hash = hashlib.sha256(img_bytes).hexdigest()  # 2026-03-29 14:30:53 修改：补缺模式按hash去重
                        
                        # 构建相对路径用于Markdown
                        rel_path = os.path.relpath(img_path, os.path.dirname(output_path))
                        rel_path = rel_path.replace("\\", "/")  # 确保路径分隔符正确
                        should_append = policy == 'appendix' or (policy == 'supplement' and img_hash not in existing_hashes)  # 2026-03-29 14:30:53 修改：补缺模式仅追加缺失图片
                        if should_append:
                            with open(output_path, "a", encoding="utf-8") as md_file:
                                md_file.write(f"\n\n![图片 {img_count}]({rel_path})\n")
                            existing_hashes.add(img_hash)  # 2026-03-29 14:30:53 修改：避免同次追加出现重复
                        
                    except Exception as e:
                        print(f"处理图片 {img_count} 时出错: {str(e)}")
            
            print(f"已从Word文档中提取 {img_count} 张图片")
            self._update_ui(status=f'已提取 {img_count} 张图片')
            
        except Exception as e:
            print(f"从Word提取图片失败: {str(e)}")
            self._update_ui(status=f'图片提取失败: {str(e)}')

    def extract_images_from_pdf(self, input_file, output_path, images_dir):
        """从PDF提取图片并添加到Markdown文件中"""
        try:
            # 尝试使用pdf2image将PDF页面转为图片
            from pdf2image import convert_from_path
            import hashlib  # 2026-03-29 14:30:53 修改：仅补缺模式使用hash去重，避免文末重复图片
            import base64
            import mimetypes
            policy = self._get_image_output_policy()  # 2026-03-29 14:30:53 修改：根据图片输出策略决定是否追加/补缺
            if policy == 'inplace':  # 2026-03-29 14:30:53 修改：仅原位不做文末追加
                return
            existing_hashes = self._get_existing_image_hashes(output_path) if policy == 'supplement' else set()  # 2026-03-29 14:30:53 修改：补缺模式按hash去重
            
            # 文档名前缀
            pdf_prefix = os.path.splitext(os.path.basename(input_file))[0]
            
            # 获取 poppler 路径并打印
            poppler_path = self.find_poppler_bin()
            print(f"PDF 提取图片使用 Poppler 路径: {poppler_path}")
            
            # 转换PDF页面为图片
            try:
                # 尝试转换每一页
                # 修复：添加 poppler_path 参数以确保在非系统路径环境下能找到二进制文件
                images = convert_from_path(input_file, dpi=150, poppler_path=poppler_path) 
                
                for i, image in enumerate(images):
                    # 保存图片
                    img_name = f"{pdf_prefix}_page_{i+1}.png"
                    img_path = os.path.join(images_dir, img_name)
                    image.save(img_path, "PNG")
                    
                    # 构建相对路径用于Markdown
                    rel_path = os.path.relpath(img_path, os.path.dirname(output_path))
                    rel_path = rel_path.replace("\\", "/")  # 确保路径分隔符正确

                    with open(img_path, 'rb') as img_f:
                        img_bytes = img_f.read()
                    img_hash = hashlib.sha256(img_bytes).hexdigest()  # 2026-03-29 14:30:53 修改：补缺模式按hash去重
                    should_append = policy == 'appendix' or (policy == 'supplement' and img_hash not in existing_hashes)  # 2026-03-29 14:30:53 修改：补缺模式仅追加缺失图片
                    if not should_append:
                        continue

                    with open(output_path, "a", encoding="utf-8") as md_file:
                        if hasattr(self, 'embed_images_var') and self.embed_images_var.get():
                            mime_type = mimetypes.guess_type(img_path)[0] or "image/png"
                            b64 = base64.b64encode(img_bytes).decode('ascii')
                            md_file.write(f"\n\n![页面 {i+1}](data:{mime_type};base64,{b64})\n")
                        else:
                            md_file.write(f"\n\n![页面 {i+1}]({rel_path})\n")
                    existing_hashes.add(img_hash)  # 2026-03-29 14:30:53 修改：避免同次追加出现重复
                
                print(f"已从PDF提取 {len(images)} 页图片")
                self._update_ui(status=f'已提取 {len(images)} 页图片')
                
            except Exception as e:
                print(f"PDF转图片失败: {str(e)}")
                self._update_ui(status=f'PDF转图片失败: {str(e)}')
            
        except ImportError:
            print("pdf2image模块不可用，无法提取PDF图片")
            self._update_ui(status='pdf2image模块不可用，无法提取PDF图片')

    # 【逻辑锁定】PDF 处理模块·C 簇成员（扫描件判定），未经专门批准禁止修改（见 A 簇头部锁定说明）。
    def _is_scanned_pdf(self, input_file):
        try:
            import pdfplumber
            total_pages = 0
            total_text_len = 0
            with pdfplumber.open(input_file) as pdf:
                total_pages = len(pdf.pages)
                for page in pdf.pages[:10]:
                    text = page.extract_text() or ""
                    total_text_len += len(text.strip())
            if total_pages == 0:
                return False
            avg_text = total_text_len / min(total_pages, 10)
            return avg_text < 30
        except Exception:
            try:
                import PyPDF2
                with open(input_file, 'rb') as f:
                    reader = PyPDF2.PdfReader(f)
                    pages = reader.pages[:10]
                    total = 0
                    for p in pages:
                        t = p.extract_text() or ""
                        total += len(t.strip())
                    if not pages:
                        return False
                    return (total / len(pages)) < 30
            except Exception:
                return False

    # 【逻辑锁定】PDF 处理模块·C 簇成员（整页转图），未经专门批准禁止修改（见 A 簇头部锁定说明）。
    def _convert_pdf_pages_to_markdown(self, input_file, output_path, images_dir, embed_images=True):
        try:
            from pdf2image import convert_from_path
            import base64
            import mimetypes

            os.makedirs(images_dir, exist_ok=True)
            poppler_path = self.find_poppler_bin()
            print(f"使用 Poppler 路径: {poppler_path}")

            pdf_prefix = os.path.splitext(os.path.basename(input_file))[0]
            
            # 使用 poppler_path 参数
            images = convert_from_path(input_file, dpi=150, poppler_path=poppler_path)
            parts = []
            for i, image in enumerate(images):
                img_name = f"{pdf_prefix}_page_{i+1}.png"
                img_path = os.path.join(images_dir, img_name)
                image.save(img_path, "PNG")

                if embed_images:
                    with open(img_path, 'rb') as img_f:
                        img_bytes = img_f.read()
                    mime_type = mimetypes.guess_type(img_path)[0] or "image/png"
                    b64 = base64.b64encode(img_bytes).decode('ascii')
                    parts.append(f"![页面 {i+1}](data:{mime_type};base64,{b64})")
                else:
                    rel_path = os.path.relpath(img_path, os.path.dirname(output_path)).replace("\\", "/")
                    parts.append(f"![页面 {i+1}]({rel_path})")

            with open(output_path, 'w', encoding='utf-8') as f:
                f.write("\n\n".join(parts) + "\n")
            return True
        except Exception as e:
            print(f"PDF转图片失败: {str(e)}")
            return False

    def _inline_local_images_in_md(self, md_path):
        import base64
        import mimetypes
        import re
        import urllib.parse

        try:
            with open(md_path, 'r', encoding='utf-8') as f:
                content = f.read()

            md_dir = os.path.dirname(md_path)

            def repl(m):
                alt = m.group(1)
                target = m.group(2).strip()
                if not target or target.startswith('data:'):
                    return m.group(0)
                if re.match(r'^[a-zA-Z]+://', target):
                    return m.group(0)
                target = urllib.parse.unquote(target)
                target = target.split('#', 1)[0].split('?', 1)[0]
                if target.startswith('<') and target.endswith('>'):
                    target = target[1:-1]
                if os.path.isabs(target):
                    img_path = os.path.abspath(target)
                else:
                    img_path = os.path.abspath(os.path.join(md_dir, target))
                if hasattr(self, 'enable_path_boundary_var') and self.enable_path_boundary_var.get():
                    if not self._is_path_within_dir(img_path, md_dir):
                        return m.group(0)
                if not os.path.exists(img_path):
                    return m.group(0)
                mime_type = mimetypes.guess_type(img_path)[0]
                if not mime_type or not mime_type.startswith('image/'):
                    return m.group(0)
                with open(img_path, 'rb') as img_f:
                    data = img_f.read()
                b64 = base64.b64encode(data).decode('ascii')
                return f"![{alt}](data:{mime_type};base64,{b64})"

            pattern = r'!\[([^\]]*)\]\(([^)]+)\)'
            new_content = re.sub(pattern, repl, content)

            if new_content != content:
                with open(md_path, 'w', encoding='utf-8') as f:
                    f.write(new_content)
        except Exception as e:
            print(f"内嵌图片失败: {str(e)}")

    def extract_images_from_pptx(self, input_file, output_path, images_dir):
        """从PowerPoint提取图片并添加到Markdown文件中"""
        try:
            import hashlib  # 2026-03-29 14:30:53 修改：仅补缺模式使用hash去重，避免文末重复图片
            # 使用pptx库处理PowerPoint文件
            from pptx import Presentation
            
            # 文档名前缀
            ppt_prefix = os.path.splitext(os.path.basename(input_file))[0]
            
            # 打开演示文稿
            prs = Presentation(input_file)
            
            img_count = 0
            policy = self._get_image_output_policy()  # 2026-03-29 14:30:53 修改：根据图片输出策略决定是否追加/补缺
            if policy == 'inplace':  # 2026-03-29 14:30:53 修改：仅原位不做文末追加
                return
            existing_hashes = self._get_existing_image_hashes(output_path) if policy == 'supplement' else set()  # 2026-03-29 14:30:53 修改：补缺模式按hash去重
            # 遍历所有幻灯片
            for i, slide in enumerate(prs.slides):
                # 遍历幻灯片中的所有形状
                for shape in slide.shapes:
                    # 检查是否是图片
                    if shape.shape_type == 13:  # MSO_SHAPE_TYPE.PICTURE
                        try:
                            img_count += 1
                            img_name = f"{ppt_prefix}_slide_{i+1}_image_{img_count}.png"
                            img_path = os.path.join(images_dir, img_name)
                            
                            # 获取图片数据并保存
                            image = shape.image
                            img_bytes = image.blob  # 2026-03-29 14:30:53 修改：用于hash去重与保存
                            with open(img_path, "wb") as f:
                                f.write(img_bytes)
                            img_hash = hashlib.sha256(img_bytes).hexdigest()  # 2026-03-29 14:30:53 修改：补缺模式按hash去重
                            
                            # 构建相对路径用于Markdown
                            rel_path = os.path.relpath(img_path, os.path.dirname(output_path))
                            rel_path = rel_path.replace("\\", "/")
                            should_append = policy == 'appendix' or (policy == 'supplement' and img_hash not in existing_hashes)  # 2026-03-29 14:30:53 修改：补缺模式仅追加缺失图片
                            if should_append:
                                with open(output_path, "a", encoding="utf-8") as md_file:
                                    md_file.write(f"\n\n![幻灯片 {i+1} 图片 {img_count}]({rel_path})\n")
                                existing_hashes.add(img_hash)  # 2026-03-29 14:30:53 修改：避免同次追加出现重复
                        except Exception as e:
                            print(f"处理幻灯片 {i+1} 中的图片时出错: {str(e)}")
                            
            print(f"已从PowerPoint中提取 {img_count} 张图片")
            self.status_var.set(f'已提取 {img_count} 张图片')
                        
        except Exception as e:
            print(f"从PowerPoint提取图片失败: {str(e)}")
            self.status_var.set(f'图片提取失败: {str(e)}')

    def extract_images_from_html(self, input_file, output_path, images_dir):
        """从HTML文件提取图片并添加到Markdown文件中"""
        try:
            # 使用BeautifulSoup解析HTML
            from bs4 import BeautifulSoup
            import requests
            import urllib.parse
            import hashlib  # 2026-03-29 14:30:53 修改：仅补缺模式使用hash去重，避免文末重复图片
            policy = self._get_image_output_policy()  # 2026-03-29 14:30:53 修改：根据图片输出策略决定是否追加/补缺
            if policy == 'inplace':  # 2026-03-29 14:30:53 修改：仅原位不做文末追加
                return
            existing_hashes = self._get_existing_image_hashes(output_path) if policy == 'supplement' else set()  # 2026-03-29 14:30:53 修改：补缺模式按hash去重
            
            # 文档名前缀
            html_prefix = os.path.splitext(os.path.basename(input_file))[0]
            
            # 读取HTML文件
            with open(input_file, 'r', encoding='utf-8') as f:
                html_content = f.read()
            
            # 解析HTML
            soup = BeautifulSoup(html_content, 'html.parser')
            
            # 查找所有图片标签
            img_tags = soup.find_all('img')
            
            img_count = 0
            for img in img_tags:
                try:
                    img_src = img.get('src')
                    if not img_src:
                        continue
                    
                    img_count += 1
                    
                    # 处理图片路径
                    if img_src.startswith('http://') or img_src.startswith('https://'):
                        # 网络图片
                        try:
                            response = requests.get(img_src, timeout=10)
                            img_data = response.content
                            # 从URL中提取文件扩展名
                            img_ext = os.path.splitext(urllib.parse.urlparse(img_src).path)[1]
                            if not img_ext or len(img_ext) > 5:  # 如果扩展名不存在或太长
                                img_ext = '.png'  # 默认为PNG
                        except:
                            print(f"无法下载网络图片: {img_src}")
                            continue
                    else:
                        # 本地图片
                        if img_src.startswith('data:image'):
                            # 数据URI
                            continue  # 暂时跳过数据URI图片
                        else:
                            # 文件路径可能是相对路径
                            img_path_abs = img_src
                            if not os.path.isabs(img_src):
                                img_path_abs = os.path.join(os.path.dirname(input_file), img_src)
                            
                            try:
                                with open(img_path_abs, 'rb') as img_file:
                                    img_data = img_file.read()
                                # 获取扩展名
                                img_ext = os.path.splitext(img_src)[1]
                            except:
                                print(f"无法读取本地图片: {img_path_abs}")
                                continue
                    
                    # 保存图片
                    img_name = f"{html_prefix}_image_{img_count}{img_ext}"
                    img_path = os.path.join(images_dir, img_name)
                    
                    with open(img_path, 'wb') as f:
                        f.write(img_data)
                    img_hash = hashlib.sha256(img_data).hexdigest()  # 2026-03-29 14:30:53 修改：补缺模式按hash去重
                    
                    # 构建相对路径用于Markdown
                    rel_path = os.path.relpath(img_path, os.path.dirname(output_path))
                    rel_path = rel_path.replace("\\", "/")
                    should_append = policy == 'appendix' or (policy == 'supplement' and img_hash not in existing_hashes)  # 2026-03-29 14:30:53 修改：补缺模式仅追加缺失图片
                    if should_append:
                        with open(output_path, "a", encoding="utf-8") as md_file:
                            alt_text = img.get('alt', f'图片 {img_count}')
                            md_file.write(f"\n\n![{alt_text}]({rel_path})\n")
                        existing_hashes.add(img_hash)  # 2026-03-29 14:30:53 修改：避免同次追加出现重复
                    
                except Exception as e:
                    print(f"处理HTML图片 {img_count} 时出错: {str(e)}")
            
            print(f"已从HTML中提取 {img_count} 张图片")
            self.status_var.set(f'已提取 {img_count} 张图片')
            
        except Exception as e:
            print(f"从HTML提取图片失败: {str(e)}")
            self.status_var.set(f'图片提取失败: {str(e)}')

    def preview_markdown(self):
        try:
            output_path = getattr(self, 'last_output_path', None)  # 2026-03-29 13:41:12 修改：修复预览功能，优先预览最近输出
            if not output_path:
                input_file = self.input_path_var.get() if hasattr(self, 'input_path_var') else ''
                output_folder = self.output_path_var.get() if hasattr(self, 'output_path_var') else ''
                if input_file and output_folder:
                    output_filename = os.path.splitext(os.path.basename(input_file))[0] + '.md'
                    output_path = os.path.join(output_folder, output_filename)
            if not output_path or not os.path.exists(output_path):
                messagebox.showinfo('提示', '未找到可预览的 Markdown 文件，请先转换。')
                return
            file_size = os.path.getsize(output_path)
            if file_size > 1024 * 1024:
                markdown_content = "文件过大 (%.2f MB)，无法在预览窗口直接显示。请在输出目录中查看。" % (file_size / (1024 * 1024))
            else:
                with open(output_path, 'r', encoding='utf-8', errors='ignore') as f:
                    markdown_content = f.read()
            self.preview_text.delete(1.0, tk.END)
            self.preview_text.insert(tk.END, markdown_content)
            self.status_var.set('预览已更新')
        except Exception as e:
            messagebox.showerror('错误', f'预览失败:\n{str(e)}')

    def copy_to_clipboard(self):
        """复制预览内容到剪贴板"""
        content = self.preview_text.get(1.0, tk.END)
        
        if not content.strip():
            messagebox.showinfo('提示', '预览内容为空')
            return
        
        self.root.clipboard_clear()
        self.root.clipboard_append(content)
        self.status_var.set('已复制到剪贴板')

    def save_preview(self):
        """保存预览内容"""
        content = self.preview_text.get(1.0, tk.END)
        
        if not content.strip():
            messagebox.showinfo('提示', '预览内容为空')
            return
        
        file_path = filedialog.asksaveasfilename(
            title='保存预览',
            defaultextension='.md',
            filetypes=[('Markdown 文件', '*.md'), ('文本文件', '*.txt'), ('所有文件', '*.*')]
        )
        
        if not file_path:
            return
        
        try:
            with open(file_path, 'w', encoding='utf-8') as f:
                f.write(content)
            
            self.status_var.set(f'预览已保存到: {file_path}')
            messagebox.showinfo('成功', f'预览内容已保存到:\n{file_path}')
        except Exception as e:
            messagebox.showerror('保存失败', f'保存预览内容时出错:\n{str(e)}')

    def open_output_folder(self, folder_path):
        """打开输出文件夹"""
        try:
            if os.name == 'nt':  # Windows
                os.startfile(folder_path)
            elif os.name == 'posix':  # macOS 和 Linux
                subprocess.run(['xdg-open', folder_path], timeout=10)  # 2026-03-29 13:41:12 修改：外部打开命令增加超时，防卡死
        except Exception as e:
            messagebox.showerror('错误', f'无法打开输出文件夹:\n{str(e)}')

    def load_settings(self):
        """加载设置"""
        settings_file = os.path.join(os.path.expanduser('~'), '.markitdown_gui_settings.json')
        
        if os.path.exists(settings_file):
            try:
                with open(settings_file, 'r', encoding='utf-8') as f:
                    settings = json.load(f)
                
                # 应用设置
                if 'default_output_folder' in settings:
                    self.default_output_var.set(settings['default_output_folder'])
                
                if 'auto_open_folder' in settings:
                    self.auto_open_folder_var.set(settings['auto_open_folder'])
                
                if 'auto_preview' in settings:
                    self.auto_preview_var.set(settings['auto_preview'])

                if 'embed_images' in settings and hasattr(self, 'embed_images_var'):
                    self.embed_images_var.set(settings['embed_images'])

                if 'enable_path_boundary_guard' in settings and hasattr(self, 'enable_path_boundary_var'):
                    self.enable_path_boundary_var.set(settings['enable_path_boundary_guard'])

                if 'image_output_policy' in settings and hasattr(self, 'image_output_policy_var'):
                    policy = settings.get('image_output_policy')
                    if policy == 'appendix':
                        self.image_output_policy_var.set('全量附录')
                    elif policy == 'supplement':
                        self.image_output_policy_var.set('仅补缺')
                    else:
                        self.image_output_policy_var.set('仅原位')

                if 'batch_overwrite_policy' in settings and hasattr(self, 'batch_overwrite_policy_var'):  # 2026-03-29 13:41:12 修改：批量覆盖策略持久化读取
                    policy = settings.get('batch_overwrite_policy')
                    if policy == 'overwrite':
                        self.batch_overwrite_policy_var.set('覆盖')
                    elif policy == 'rename':
                        self.batch_overwrite_policy_var.set('自动重命名')
                    else:
                        self.batch_overwrite_policy_var.set('跳过')

                if 'document_output_mode' in settings and hasattr(self, 'document_output_mode_var'):
                    self._set_document_output_mode_label(settings.get('document_output_mode'))
                elif hasattr(self, 'document_output_mode_var'):
                    self._set_document_output_mode_label(self._map_legacy_output_mode(settings))

                if 'pdf_engine_strategy' in settings and hasattr(self, 'pdf_engine_strategy_var'):
                    strategy = settings.get('pdf_engine_strategy')
                    if strategy == 'layout_first':
                        self.pdf_engine_strategy_var.set('图文混排优先')
                    elif strategy == 'inspector_first':
                        self.pdf_engine_strategy_var.set('pdf-inspector优先')
                    else:
                        self.pdf_engine_strategy_var.set(PDF_ENGINE_STRATEGY_DEFAULT)
            except Exception as e:
                print(f"加载设置时出错: {str(e)}")

    def save_settings(self):
        """保存设置"""
        settings = {
            'default_output_folder': self.default_output_var.get(),
            'auto_open_folder': self.auto_open_folder_var.get(),
            'auto_preview': self.auto_preview_var.get(),
            'embed_images': self.embed_images_var.get() if hasattr(self, 'embed_images_var') else False,
            'enable_path_boundary_guard': self.enable_path_boundary_var.get() if hasattr(self, 'enable_path_boundary_var') else False,
            'image_output_policy': self._get_image_output_policy(),
            'batch_overwrite_policy': self._get_batch_overwrite_policy(),  # 2026-03-29 13:41:12 修改：批量覆盖策略持久化写入
            'document_output_mode': self._get_document_output_mode(),
            'pdf_engine_strategy': self._get_pdf_engine_strategy()  # 2026-09-01 新增：PDF 引擎策略持久化
        }
        
        settings_file = os.path.join(os.path.expanduser('~'), '.markitdown_gui_settings.json')
        
        try:
            with open(settings_file, 'w', encoding='utf-8') as f:
                json.dump(settings, f, indent=2)
            
            self.status_var.set('设置已保存')
            messagebox.showinfo('成功', '设置已保存')
        except Exception as e:
            messagebox.showerror('保存失败', f'保存设置时出错:\n{str(e)}')

    def get_markitdown_info(self):
        """获取markitdown版本和功能信息"""
        info = {
            'version': 'unknown',
            'extract_images_param': None,
            'images_dir_param': None,
            'pdf_extract_param': None,
            'html_mode_param': None,
            'preserve_links_param': None,
            'all_params': [],
            'features_text': '基本转换'
        }
        
        try:
            # 获取版本
            try:
                version_process = subprocess.run(
                    ["markitdown", "--version"],
                    capture_output=True,
                    text=True,
                    timeout=5,
                    creationflags=subprocess.CREATE_NO_WINDOW
                )
                version_text = version_process.stdout.strip()
                if version_text:
                    info['version'] = version_text
            except (subprocess.TimeoutExpired, Exception):
                pass
            
            # 获取帮助信息
            help_text = ""
            try:
                help_process = subprocess.run(
                    ["markitdown", "--help"],
                    capture_output=True,
                    text=True,
                    timeout=5,
                    creationflags=subprocess.CREATE_NO_WINDOW
                )
                help_text = help_process.stdout + help_process.stderr
            except (subprocess.TimeoutExpired, Exception):
                pass
            
            if help_text:
                # 提取所有参数
                params = re.findall(r'--([\w-]+)', help_text)
                info['all_params'] = params
                
                # 检查特定参数
                if "--extract-images" in help_text:
                    info['extract_images_param'] = "--extract-images"
                elif "--images" in help_text:
                    info['extract_images_param'] = "--images"
                
                if "--images-dir" in help_text:
                    info['images_dir_param'] = "--images-dir"
                
                if "--pdf-extract-images" in help_text:
                    info['pdf_extract_param'] = "--pdf-extract-images"
                
                # 更新功能描述文本
                features = []
                if info['extract_images_param']:
                    features.append("图片提取")
                if "pdf" in help_text.lower():
                    features.append("PDF")
                if "docx" in help_text.lower() or "doc" in help_text.lower():
                    features.append("Word")
                if "pptx" in help_text.lower():
                    features.append("PPT")
                
                if features:
                    info['features_text'] = "支持: " + "、".join(features)
                
            return info
        except Exception as e:
            print(f"获取markitdown信息失败: {str(e)}")
            if info.get('version') == 'unknown':
                info['version'] = _get_markitdown_version()
            return info

    def run_initial_feature_check(self):
        """运行初始功能检查，确保关键功能可用"""
        # 检查图片提取功能
        self.image_extraction_supported = bool(self.markitdown_info.get('extract_images_param'))
        if not self.image_extraction_supported:
            print("MarkItDown 命令行不支持图片提取，将使用自定义图片提取")
        
        # 检查PDF功能
        self.pdf_supported = False
        try:
            import pdfminer
            self.pdf_supported = True
            print("支持PDF处理")
        except ImportError:
            print("不支持PDF处理 (缺少pdfminer)")
        
        # 检查PDF图片提取
        self.pdf_image_extraction_supported = False
        try:
            import pdf2image
            # 仅检查模块是否存在，不再运行 noisy 的 pdfinfo_from_bytes
            self.pdf_image_extraction_supported = True
            print("支持PDF图片提取 (pdf2image)")
        except ImportError:
            print("不支持PDF图片提取 (缺少pdf2image)")
        
        # 检查Word文档支持
        self.word_supported = False
        try:
            import docx
            self.word_supported = True
            print("支持Word文档 (DOCX)")
        except ImportError:
            print("不支持Word文档 (缺少python-docx)")
        
        # 检查DOC转换支持
        self.doc_conversion_supported = False
        try:
            import win32com.client
            self.doc_conversion_supported = True
            print("支持Word DOC转换 (pywin32)")
        except ImportError:
            print("不支持Word DOC转换 (缺少pywin32)")

    def start_batch_conversion(self):
        """开始批量文件转换"""
        if self.conversion_in_progress:
            messagebox.showinfo('提示', '转换正在进行中，请等待完成')
            return
            
        # 获取输入文件和输出目录
        input_files = self.batch_files
        output_folder = self.batch_output_var.get()
        
        # 检查是否有文件要转换
        if not input_files:
            messagebox.showwarning('警告', '请先添加需要转换的文件')
            return
        
        # 检查输出目录
        if not output_folder:
            messagebox.showwarning('警告', '请选择输出目录')
            return
        
        # 确保输出目录存在
        os.makedirs(output_folder, exist_ok=True)
        
        # 在新线程中启动批量转换
        threading.Thread(target=self._batch_conversion_worker, daemon=True).start()

    def _batch_conversion_worker(self):
        """批量转换工作线程"""
        try:
            self.conversion_in_progress = True
            self.root.after(0, lambda: self._set_widget_enabled(self.batch_convert_btn, False))
            self._update_ui(status='准备批量转换...', progress=0, is_batch=True)
            
            # 清空不成功文件列表
            self.root.after(0, lambda: self.pending_files_listbox.delete(0, tk.END))
            
            input_files = self.batch_files
            output_folder = self.batch_output_var.get()
            
            # 创建图片目录
            images_dir = os.path.join(output_folder, "images")
            os.makedirs(images_dir, exist_ok=True)
            
            # 转换每个文件
            total_files = len(input_files)
            successful_conversions = 0
            failed_conversions = []
            skipped_conversions = []
            
            for i, input_file in enumerate(input_files):
                try:
                    original_input_file = input_file
                    # 更新进度
                    progress_percent = (i / total_files) * 100
                    self._update_ui(
                        status=f'正在转换 {i+1}/{total_files}: {os.path.basename(input_file)}',
                        progress=progress_percent,
                        is_batch=True
                    )
                    
                    # 获取文件扩展名
                    original_file_ext = os.path.splitext(input_file)[1].lower()
                    
                    # 处理.mhtml/.mht文件
                    if original_file_ext in ['.mhtml', '.mht']:
                        mhtml_mode = self._get_mhtml_mode()
                        temp_html = self._parse_mhtml_to_html(input_file, images_dir, embed_images=True, mode=mhtml_mode)
                        if temp_html:
                            input_file = temp_html
                        else:
                            failed_conversions.append((input_file, "MHTML解析失败"))
                            continue
                            
                    # 处理.doc文件
                    elif original_file_ext == '.doc':
                        original_file = input_file
                        input_file = self.handle_doc_file(input_file)
                        # 如果转换失败，尝试直接转换
                        if input_file == original_file:
                            output_path = self.convert_doc_directly(input_file, output_folder)
                            if output_path:
                                successful_conversions += 1
                                continue
                    
                    # 生成输出路径
                    output_filename = os.path.splitext(os.path.basename(original_input_file))[0] + '.md'
                    output_path, conflict_reason = self._resolve_output_path_for_batch(output_folder, output_filename)
                    if not output_path:
                        skipped_conversions.append((original_input_file, conflict_reason or "已跳过"))
                        continue
                    document_mode = self._get_document_output_mode()
                    document_mode_exts = ['.pdf', '.docx', '.doc', '.ppt', '.pptx', '.html', '.htm', '.mhtml', '.mht']  # 2026-09-02 修改：补充旧版 .ppt，使其可用“保持版式”等模式
                    if document_mode != 'smart' and original_file_ext in document_mode_exts:
                        conversion_success = self._convert_document_by_mode(
                            original_input_file,
                            input_file,
                            output_path,
                            images_dir,
                            original_file_ext
                        )
                        if conversion_success:
                            successful_conversions += 1
                            continue
                        if document_mode == 'layout':
                            input_file = original_input_file
                        else:
                            failed_conversions.append((original_input_file, "文档输出模式转换失败"))
                            self.root.after(0, lambda f=original_input_file: self.pending_files_listbox.insert(tk.END, os.path.basename(f)))
                            continue
                    
                    # 转换标志
                    conversion_success = False
                    converted_as_scanned_pdf = False
                    
                    if original_file_ext in ['.mhtml', '.mht', '.html', '.htm']:
                        conversion_success = self._convert_html_like_to_mixed_markdown(input_file, output_path, images_dir)
                        if not conversion_success:
                            conversion_success = self.convert_html_to_markdown(input_file, output_path)
                    else:
                        if original_file_ext == '.pdf':
                            pdf_mode = self._get_pdf_mode()
                            embed_images = hasattr(self, 'embed_images_var') and self.embed_images_var.get()

                            if pdf_mode == 'layout':
                                conversion_success = self._convert_pdf_pages_to_markdown(
                                    original_input_file,
                                    output_path,
                                    images_dir,
                                    embed_images=embed_images
                                )
                                converted_as_scanned_pdf = conversion_success
                            elif pdf_mode == 'auto_scanned':
                                if self._is_scanned_pdf(original_input_file):
                                    conversion_success = self._convert_pdf_pages_to_markdown(
                                        original_input_file,
                                        output_path,
                                        images_dir,
                                        embed_images=embed_images
                                    )
                                    converted_as_scanned_pdf = conversion_success
                            elif pdf_mode == 'normal':
                                # 2026-09-01 新增：PDF 引擎策略（仅作用于智能转换模式）
                                engine_strategy = self._get_pdf_engine_strategy()
                                conversion_success = False
                                if engine_strategy == 'auto':
                                    # 自动：detect_pdf 分类，text_based 用 pdf-inspector，否则回退图文混排
                                    conversion_success = self._convert_pdf_with_pdfinspector(
                                        original_input_file,
                                        output_path,
                                        force=False
                                    )
                                elif engine_strategy == 'inspector_first':
                                    # pdf-inspector 优先：不预判类型直接尝试，失败回退图文混排
                                    conversion_success = self._convert_pdf_with_pdfinspector(
                                        original_input_file,
                                        output_path,
                                        force=True
                                    )
                                if not conversion_success:
                                    # 回退：现有图文混排路径（layout_first 或 pdf-inspector 失败）
                                    conversion_success = self._convert_pdf_to_mixed_markdown(
                                        original_input_file,
                                        output_path,
                                        images_dir
                                    )
                        
                        # 2026-09-02 新增：PPT 智能（图文并存）增强（批量）——Office→PDF→复用 PDF 图文混排引擎
                        if not conversion_success and original_file_ext in ('.ppt', '.pptx'):
                            try:
                                if self._try_ppt_smart_via_pdf(original_file_ext, input_file, output_path, images_dir):
                                    conversion_success = True
                            except Exception:
                                conversion_success = False
                        if not conversion_success:
                            try:
                                cmd = ["markitdown", input_file, "-o", output_path]
                                
                                if hasattr(self, 'batch_extract_images_var') and self.batch_extract_images_var.get() and original_file_ext in ['.docx', '.doc', '.pptx']:
                                    cmd.append("--keep-data-uris")
                                    
                                print(f"执行命令: {' '.join(cmd)}")
                                process = None
                                try:
                                    process = subprocess.run(
                                        cmd,
                                        capture_output=True,
                                        text=True,
                                        timeout=_MARKITDOWN_CLI_TIMEOUT,
                                        creationflags=subprocess.CREATE_NO_WINDOW
                                    )
                                except FileNotFoundError:  # 2026-03-29 13:41:12 修改：markitdown命令不可用时自动API回退
                                    ok = self._convert_with_markitdown_api(input_file, output_path)
                                    if ok:
                                        class _Proc:
                                            returncode = 0
                                            stderr = ""
                                        process = _Proc()
                                    else:
                                        raise
                                
                                if process.returncode == 0:
                                    conversion_success = True
                                else:
                                    print(f"MarkItDown转换失败，错误: {process.stderr}")
                            except Exception as e:
                                print(f"MarkItDown转换异常: {str(e)}")
                    
                    # 如果常规转换失败，尝试备用方法
                    if not conversion_success:
                        if original_file_ext in ['.docx', '.doc']:
                            conversion_success = self.convert_word_to_markdown(input_file, output_path)
                        elif original_file_ext == '.pdf':
                            conversion_success = self.convert_pdf_to_markdown(input_file, output_path)
                        elif original_file_ext in ['.html', '.htm', '.mhtml', '.mht']:
                            conversion_success = self._convert_html_like_to_mixed_markdown(input_file, output_path, images_dir)
                            if not conversion_success:
                                conversion_success = self.convert_html_to_markdown(input_file, output_path)
                    
                    # AnyDoc 兜底：以上所有方法均失败时，对 Office 格式启用 AnyDoc（新增处理方法）
                    if not conversion_success and original_file_ext in ['.doc', '.docx', '.ppt', '.pptx', '.xls', '.xlsx']:
                        conversion_success = self.convert_office_with_anydoc(original_input_file, output_path)
                    
                    # 如果最终转换成功
                    if conversion_success:
                        # 提取图片并处理原位图片
                        if original_file_ext in ['.mhtml', '.mht', '.html', '.htm'] and hasattr(self, 'batch_extract_images_var') and self.batch_extract_images_var.get():
                            self._process_html_images(input_file, output_path, images_dir)
                        elif hasattr(self, 'batch_extract_images_var') and self.batch_extract_images_var.get():
                            if original_file_ext in ['.docx', '.doc', '.pptx']:
                                self._process_base64_images_in_md(output_path, images_dir)
                                self._apply_office_image_output_policy(original_file_ext, input_file, output_path, images_dir)  # 2026-03-29 14:05:06 修改：使用实际参与转换的文件路径，避免DOC->DOCX时提图指向原DOC
                            elif original_file_ext == '.pdf':
                                if not converted_as_scanned_pdf:
                                    self._apply_pdf_image_output_policy(original_input_file, output_path, images_dir)

                        if hasattr(self, 'embed_images_var') and self.embed_images_var.get():
                            self._inline_local_images_in_md(output_path)
                        
                        successful_conversions += 1
                    else:
                        failed_conversions.append((input_file, "所有转换方法均失败"))
                        self.root.after(0, lambda f=input_file: self.pending_files_listbox.insert(tk.END, os.path.basename(f)))
                    
                except Exception as e:
                    print(f"转换 {input_file} 出错: {str(e)}")
                    failed_conversions.append((input_file, str(e)))
                    self.root.after(0, lambda f=input_file: self.pending_files_listbox.insert(tk.END, os.path.basename(f)))
            
            # 更新最终进度和状态
            self._update_ui(progress=100, is_batch=True)
            
            # 生成结果消息
            if failed_conversions or skipped_conversions:
                failed_msg = "\n".join([f"{os.path.basename(f)}: {err}" for f, err in failed_conversions])
                skipped_msg = "\n".join([f"{os.path.basename(f)}: {err}" for f, err in skipped_conversions])
                self._update_ui(status=f'完成 {successful_conversions}/{total_files} 个文件转换', is_batch=True)
                self.root.after(0, lambda: messagebox.showwarning('批量转换结果', 
                    f'成功转换: {successful_conversions}/{total_files}\n\n'
                    f'失败文件:\n{failed_msg}\n\n'
                    f'跳过文件:\n{skipped_msg}'))
            else:
                self._update_ui(status=f'成功完成所有 {total_files} 个文件转换', is_batch=True)
                self.root.after(0, lambda: messagebox.showinfo('成功', 
                    f'所有 {total_files} 个文件已成功转换到:\n{output_folder}'))
                
                # 自动打开输出文件夹
                if hasattr(self, 'batch_auto_open_folder_var') and self.batch_auto_open_folder_var.get():
                    self.root.after(0, lambda: self.open_output_folder(output_folder))
        
        except Exception as e:
            self._update_ui(status=f"批量转换出错: {str(e)}", is_batch=True)
            self.root.after(0, lambda: messagebox.showerror('错误', f'批量转换过程中出错:\n{str(e)}'))
        
        finally:
            self.conversion_in_progress = False
            self.root.after(0, lambda: self._set_widget_enabled(self.batch_convert_btn, True))
            
            # 延迟重置进度条和状态
            def reset_batch_ui():
                self._set_progress_percent(self.batch_progress, 0)
                if self.batch_status_var.get().startswith('成功完成') or '批量转换完成' in self.batch_status_var.get():
                    self.batch_status_var.set('准备就绪')
                    if self.use_ctk and hasattr(self, 'batch_status_label'):
                        self.batch_status_label.configure(text=self.batch_status_var.get())
            
            self.root.after(5000, reset_batch_ui)
            
            # 清理临时文件
            if hasattr(self, 'temp_files'):
                for temp_file in self.temp_files:
                    try:
                        if os.path.exists(temp_file):
                            os.remove(temp_file)
                    except Exception as e:
                        print(f"清理临时文件 {temp_file} 失败: {e}")
                self.temp_files = []

    def convert_doc_directly(self, input_file, output_folder):
        """直接将DOC转换为Markdown，不经过DOCX中间格式"""
        try:
            # 生成输出路径
            output_filename = os.path.splitext(os.path.basename(input_file))[0] + '.md'
            output_path = os.path.join(output_folder, output_filename)
            
            # 尝试使用pandoc进行转换
            try:
                self._update_ui(status='尝试使用pandoc转换...')
                
                cmd = ["pandoc", input_file, "-o", output_path]
                process = subprocess.run(
                    cmd,
                    capture_output=True,
                    text=True,
                    timeout=30,
                    creationflags=subprocess.CREATE_NO_WINDOW
                )
                
                if process.returncode == 0:
                    return output_path
            except:
                pass
            
            # 如果pandoc失败，尝试使用其他方法
            try:
                import textract
                text = textract.process(input_file).decode('utf-8')
                
                # 简单格式化文本
                with open(output_path, 'w', encoding='utf-8') as f:
                    f.write(text)
                
                return output_path
            except:
                pass
            
            return None
        except Exception as e:
            print(f"直接转换DOC失败: {str(e)}")
            return None

    def convert_office_with_anydoc(self, input_file, output_path):
        """AnyDoc 兜底：在其他方法均失败后，用 AnyDoc 将 Office 文档转换为 Markdown（新增处理方法，不改原有逻辑）"""
        try:
            import anydoc
        except Exception:
            return False
        try:
            markdown = anydoc.to_markdown(input_file)
            if not markdown or not str(markdown).strip():
                return False
            with open(output_path, 'w', encoding='utf-8') as f:
                f.write(str(markdown))
            return True
        except Exception as e:
            print(f"AnyDoc兜底转换失败: {str(e)}")
            return False

    def convert_word_to_markdown(self, input_file, output_path):
        """Word到Markdown的备用转换方法"""
        try:
            # 尝试使用pandoc
            try:
                cmd = ["pandoc", input_file, "-o", output_path]
                process = subprocess.run(
                    cmd,
                    capture_output=True,
                    text=True,
                    timeout=30,
                    creationflags=subprocess.CREATE_NO_WINDOW
                )
                
                if process.returncode == 0:
                    return True
            except:
                pass
            
            # 尝试使用python-docx和markdown库
            try:
                import docx
                from docx.shared import Pt
                
                # 读取Word文档
                doc = docx.Document(input_file)
                
                # 提取文本并转换为Markdown
                markdown_text = ""
                
                # 处理段落
                for para in doc.paragraphs:
                    if not para.text.strip():
                        continue
                    
                    # 检查段落样式
                    style = para.style.name.lower()
                    text = para.text.strip()
                    
                    # 标题处理
                    if 'heading' in style:
                        heading_level = int(style[-1]) if style[-1].isdigit() else 1
                        markdown_text += '#' * heading_level + ' ' + text + '\n\n'
                    else:
                        # 检查是否有粗体、斜体等格式
                        formatted_text = ""
                        for run in para.runs:
                            run_text = run.text
                            if run.bold and run.italic:
                                formatted_text += f"***{run_text}***"
                            elif run.bold:
                                formatted_text += f"**{run_text}**"
                            elif run.italic:
                                formatted_text += f"*{run_text}*"
                            else:
                                formatted_text += run_text
                        
                        markdown_text += formatted_text + '\n\n'
                
                # 保存Markdown文件
                with open(output_path, 'w', encoding='utf-8') as f:
                    f.write(markdown_text)
                
                return True
            except Exception as e:
                print(f"python-docx转换失败: {str(e)}")
            
            return False
        except Exception as e:
            print(f"Word备用转换失败: {str(e)}")
            return False

    def convert_pdf_to_markdown(self, input_file, output_path):
        """PDF到Markdown的备用转换方法"""
        try:
            images_dir = os.path.join(os.path.dirname(output_path), "images")
            if self._convert_pdf_to_mixed_markdown(input_file, output_path, images_dir):
                return True
            # 尝试使用pandoc
            try:
                cmd = ["pandoc", input_file, "-o", output_path]
                process = subprocess.run(
                    cmd,
                    capture_output=True,
                    text=True,
                    timeout=30,
                    creationflags=subprocess.CREATE_NO_WINDOW
                )
                
                if process.returncode == 0:
                    return True
            except:
                pass
            
            sections = self._extract_pdf_text_sections(input_file)
            if sections:
                return self._write_markdown_sections(output_path, sections)
            return False
        except Exception as e:
            print(f"PDF备用转换失败: {str(e)}")
            return False

    # 【逻辑锁定】PDF 处理模块·D 簇成员（pdf-inspector 引擎转换），未经专门批准禁止修改（见 A 簇头部锁定说明）。
    def _convert_pdf_with_pdfinspector(self, input_file, output_path, force=False):
        """pdf-inspector 兜底：将文本型 PDF 转换为 Markdown（新增处理方法，不改原有逻辑）
        
        规则（仅作用于智能转换 normal 模式）：
        - try: import pdf_inspector，失败则回退返回 False
        - force=False（自动模式）：detect_pdf() 分类，仅 text_based → 用 pdf-inspector 输出 Markdown；否则回退返回 False
        - force=True（pdf-inspector优先模式）：不预判类型，直接尝试 pdf-inspector 输出
        - pdf-inspector 未安装或转换失败时返回 False，由调用方回退现有图文混排路径
        """
        try:
            import pdf_inspector
        except Exception:
            return False
        
        try:
            if not force:
                # 快速分类：仅 text_based 使用 pdf-inspector，其余回退现有路径
                try:
                    detection = pdf_inspector.detect_pdf(input_file)
                    pdf_type = getattr(detection, 'pdf_type', None)
                except Exception:
                    return False
                if pdf_type != 'text_based':
                    return False
            
            result = pdf_inspector.process_pdf(input_file)
            markdown = getattr(result, 'markdown', None)
            if not markdown or not str(markdown).strip():
                return False
            
            with open(output_path, 'w', encoding='utf-8') as f:
                f.write(str(markdown))
            return True
        except Exception as e:
            print(f"pdf-inspector转换失败: {str(e)}")
            return False

    def convert_html_to_markdown(self, input_file, output_path):
        """HTML到Markdown的备用转换方法"""
        try:
            # 1. 优先尝试使用 markitdown 命令行工具
            try:
                cmd = ["markitdown", input_file, "-o", output_path]
                # 添加 --keep-data-uris 选项以保留base64图片
                cmd.append("--keep-data-uris")
                try:
                    process = subprocess.run(
                        cmd,
                        capture_output=True,
                        text=True,
                        timeout=60,
                        creationflags=subprocess.CREATE_NO_WINDOW
                    )
                except FileNotFoundError:
                    process = None
                    ok = self._convert_with_markitdown_api(input_file, output_path)
                    if ok and os.path.exists(output_path):
                        if os.path.getsize(output_path) > 0:
                            return True
                if process is not None and process.returncode == 0 and os.path.exists(output_path):
                    # 检查是否真的有内容（排除空文件）
                    if os.path.getsize(output_path) > 0:
                        return True
            except:
                pass

            # 2. 尝试使用 pandoc
            try:
                cmd = ["pandoc", input_file, "-o", output_path]
                process = subprocess.run(
                    cmd,
                    capture_output=True,
                    text=True,
                    timeout=30,
                    creationflags=subprocess.CREATE_NO_WINDOW
                )
                if process.returncode == 0:
                    return True
            except:
                pass
            
            # 3. 尝试使用 html2text
            try:
                import html2text
                with open(input_file, 'r', encoding='utf-8', errors='ignore') as f:
                    html_content = f.read()
                
                converter = html2text.HTML2Text()
                converter.ignore_links = False
                converter.ignore_images = False
                converter.body_width = 0 # 不换行
                markdown_text = converter.handle(html_content)
                
                if markdown_text.strip():
                    with open(output_path, 'w', encoding='utf-8') as f:
                        f.write(markdown_text)
                    return True
            except:
                pass

            # 4. 尝试使用 BeautifulSoup (增强版，包含图片提取)
            try:
                from bs4 import BeautifulSoup
                with open(input_file, 'r', encoding='utf-8', errors='ignore') as f:
                    html_content = f.read()
                
                soup = BeautifulSoup(html_content, 'html.parser')
                
                # 简单提取内容并尝试保留图片
                parts = []
                
                # 处理所有主要标签
                for tag in soup.find_all(['h1', 'h2', 'h3', 'h4', 'h5', 'h6', 'p', 'ul', 'ol', 'img']):
                    if tag.name.startswith('h'):
                        level = int(tag.name[1])
                        parts.append('#' * level + ' ' + tag.get_text().strip())
                    elif tag.name == 'p':
                        parts.append(tag.get_text().strip())
                    elif tag.name in ['ul', 'ol']:
                        for li in tag.find_all('li'):
                            prefix = '* ' if tag.name == 'ul' else '1. '
                            parts.append(prefix + li.get_text().strip())
                    elif tag.name == 'img':
                        src = tag.get('src', '')
                        alt = tag.get('alt', 'image')
                        if src:
                            parts.append(f'![{alt}]({src})')
                
                markdown_text = "\n\n".join(parts)
                if markdown_text.strip():
                    with open(output_path, 'w', encoding='utf-8') as f:
                        f.write(markdown_text)
                    return True
            except:
                pass
            
            return False
        except Exception as e:
            print(f"HTML综合转换失败: {str(e)}")
            return False


    # 20250325新增def __del__(self)，功能是资源清理改进
    def __del__(self):
        """确保程序退出时清理资源"""
        try:
            # 清理临时文件
            if hasattr(self, 'temp_files'):
                for temp_file in self.temp_files:
                    if os.path.exists(temp_file):
                        os.remove(temp_file)
            # 清理临时目录
            if hasattr(self, 'temp_dirs'):
                for temp_dir in self.temp_dirs:
                    if os.path.exists(temp_dir):
                        shutil.rmtree(temp_dir, ignore_errors=True)
        except Exception as e:
            print(f"清理资源时出错: {e}")

    # 20250325新增def __del__(self)结束

def main():
    root = tk.Tk()
    app = MarkItDownApp(root)
    root.mainloop()

if __name__ == '__main__':
    main()
